"""Сборка колоды из шаблона: выбор образца под каждый слайд и раскладка содержимого по его слотам.

Всё детерминированно: одинаковый контент + шаблон = одинаковый файл. Модель здесь не участвует —
она пишет содержимое раньше (планировщик, писатель) и исправляет найденное позже (аудит).
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation

from prism.color import delta_e, relative_luminance
from prism.layout import fit
from prism.layout import pptx_ops as ops
from prism.parsing.model import BBox
from prism.parsing.patterns import Pattern, RepeatGroup, Slot, SlotKind
from prism.parsing.spec import PatternSpec, TemplateSpec
from prism.planning.schemas import Item, SlideContent, SlideKind

# если в шаблоне нет образца нужного типа — чем его заменить (по убыванию пригодности)
FALLBACK: dict[SlideKind, list[SlideKind]] = {
    SlideKind.title: [SlideKind.section, SlideKind.closing],
    SlideKind.section: [SlideKind.title, SlideKind.closing],
    SlideKind.closing: [SlideKind.section, SlideKind.title],
    SlideKind.agenda: [SlideKind.process, SlideKind.cards, SlideKind.bullets],
    SlideKind.cards: [SlideKind.process, SlideKind.kpi, SlideKind.two_column, SlideKind.bullets],
    SlideKind.process: [SlideKind.timeline, SlideKind.cards],
    SlideKind.timeline: [SlideKind.process, SlideKind.cards],
    SlideKind.kpi: [SlideKind.cards, SlideKind.process],
    SlideKind.bullets: [SlideKind.two_column, SlideKind.cards],
    SlideKind.two_column: [SlideKind.comparison, SlideKind.cards, SlideKind.bullets],
    SlideKind.comparison: [SlideKind.two_column, SlideKind.cards],
    SlideKind.team: [SlideKind.cards],
    SlideKind.quote: [SlideKind.section, SlideKind.bullets],
    SlideKind.image: [SlideKind.two_column, SlideKind.bullets, SlideKind.section],
    SlideKind.chart: [SlideKind.two_column, SlideKind.bullets, SlideKind.cards],
    SlideKind.table: [SlideKind.bullets, SlideKind.two_column, SlideKind.cards],
}
_TEXTUAL = {SlotKind.heading, SlotKind.text, SlotKind.caption, SlotKind.number, SlotKind.title}


@dataclass
class SlideReport:
    index: int                       # номер слайда в итоговой колоде (с 1)
    pattern: str
    kind: str
    warnings: list[str] = field(default_factory=list)


# ── выбор образца ──────────────────────────────────────────────────────────


def _capacity(slot: Slot) -> float:
    """Грубо: сколько символов влезет в слот его кеглем (фигура с автоподбором может подрасти вдвое)."""
    size = (slot.styles[0].size_pt if slot.styles and slot.styles[0].size_pt else 14) or 14
    w_pt, h_pt = slot.bbox.w / 12700, slot.bbox.h / 12700
    if slot.autofit == "shape":
        h_pt *= 2
    return max((w_pt / (size * 0.55)) * (h_pt / (size * 1.2)), 1.0)


def _overflow_penalty(p: Pattern, content: SlideContent) -> float:
    """Штраф, если реальный текст пунктов заведомо не влезет в текстовые слоты карточек образца."""
    g = _main_group(p)
    texts = [i.text or "" for i in content.items if i.text]
    if g is None or not texts:
        return 0.0
    slots = [sl for sl in g.items[0] if sl.kind in (SlotKind.text, SlotKind.caption)]
    if not slots:
        slots = [sl for sl in g.items[0] if sl.kind == SlotKind.heading]
    if not slots:
        return 0.0
    capacity = max(_capacity(sl) for sl in slots)
    ratio = (sum(len(t) for t in texts) / len(texts)) / capacity
    return min(4.0, max(0.0, ratio - 1.1) * 2.5)


def _main_group(p: Pattern) -> RepeatGroup | None:
    return p.groups[0] if p.groups else None


def _score(ps: PatternSpec, content: SlideContent, used: Counter, exact_kind: bool,
           slide_area: int = 0) -> float:
    p = ps.pattern
    n = len(content.items)
    g = _main_group(p)
    if any(i.value for i in content.items):
        # крупные одиночные числа («герой» слайда) тоже вмещают показатели
        n = max(n - sum(1 for s in p.slots if s.kind == SlotKind.number), 0 if g is None else 1)
    score = 0.0 if exact_kind else 4.0
    if p.title is None:
        # некуда поставить заголовок — например, обложка с «запечённым» в картинку текстом
        score += 20.0
    if n:
        if g is None:
            text_singles = [s for s in p.slots if s.kind in (SlotKind.text, SlotKind.caption)]
            score += 1.5 if text_singles else 6.0            # пункты уйдут списком в текстовый блок
        else:
            diff = g.count - n
            if g.fixed_count and diff:
                score += 8.0                                 # карточки нарисованы в макете — лишние торчали бы
            elif diff >= 0:
                score += 0.6 * diff                          # лишние карточки удалим
            else:
                score += 3.0 * -diff                         # недостающие — придётся добавлять
            # покрытие полей: у карточки образца должно быть куда поставить каждое поле содержимого
            slot_kinds = {k for it in g.items for k in (sl.kind for sl in it)}
            two_para = any(sl.kind == SlotKind.heading and len(sl.styles) >= 2 for it in g.items for sl in it)
            if any(i.value for i in content.items) and SlotKind.number not in slot_kinds:
                score += 4.5 if content.kind == SlideKind.kpi else 2.5   # KPI без крупных чисел теряет смысл
            # вместимость слота-числа: «01» в узкой рамке не вместит «1 150 ₽»
            longest = max((len(i.value) for i in content.items if i.value), default=0)
            numbers = [sl for it in g.items for sl in it if sl.kind == SlotKind.number]
            if longest and numbers:
                n0 = numbers[0]
                size = n0.styles[0].size_pt if n0.styles and n0.styles[0].size_pt else 18
                capacity = n0.bbox.w / (size * 0.62 * 12700)
                if longest > capacity:
                    score += min(4.0, (longest - capacity) * 0.8)
            heading_matters = content.kind not in (SlideKind.kpi, SlideKind.timeline, SlideKind.bullets)
            if heading_matters and any(i.heading for i in content.items) and SlotKind.heading not in slot_kinds:
                score += 2.0
            if any(i.text for i in content.items) and not (slot_kinds & {SlotKind.text, SlotKind.caption}) \
                    and not two_para:
                score += 2.0 if any(i.heading for i in content.items) else 0.5
    elif g is not None:
        score += 3.0                                         # группа останется пустой
    kinds = {s.kind for s in p.slots}
    for needed, slot_kind in ((content.chart, SlotKind.chart), (content.table, SlotKind.table)):
        if needed and slot_kind not in kinds:
            # без своего слота график встанет на место картинки или карточек — пустой образец лучше
            score += 1.0 + (2.0 if p.groups else 0.0) + (0.0 if SlotKind.image in kinds else 0.5)
            # и места должно хватить: график в четверть слайда нечитаем
            regions = [s.bbox.area for s in p.slots if s.kind in (SlotKind.image, SlotKind.text, SlotKind.caption)]
            regions += [pic.bbox.area for pic in p.pictures if not pic.from_layout and not pic.bleed]
            if p.groups:
                bs = p.groups[0].item_bboxes
                regions.append((max(b.right for b in bs) - min(b.x for b in bs))
                               * (max(b.bottom for b in bs) - min(b.y for b in bs)))
            if max(regions, default=0) < 0.25 * slide_area:
                score += 3.0
    has_image_slot = SlotKind.image in kinds or (g and SlotKind.image in set(g.item_kinds()))
    if content.image and not has_image_slot:
        score += 2.0
    if not content.image and has_image_slot:
        score += 1.5                                         # фото-заглушку придётся убрать
    score += _overflow_penalty(p, content)
    return score + 1.2 * used[p.id]                          # разнообразие: не повторять один образец подряд


def choose_pattern(spec: TemplateSpec, content: SlideContent, used: Counter,
                   variant: int = 1, avoid: str | None = None) -> PatternSpec:
    """Оцениваются все образцы разом: свой тип — без штрафа, запасной — по порядку пригодности, прочие —
    с большим штрафом. Так негодный образец «своего» типа (обложка без текстовых слотов) проигрывает
    нормальному запасному. variant/avoid — смещения для вариантов вёрстки (см. VARIANTS)."""
    fallback = FALLBACK.get(content.kind, [])
    slide_area = spec.tokens.slide_w * spec.tokens.slide_h

    def total(ps: PatternSpec) -> float:
        if ps.kind == content.kind:
            tier = 0.0
        elif ps.kind in fallback:
            tier = 1.0 + fallback.index(ps.kind)
        else:
            tier = 10.0
        score = tier + _score(ps, content, used, exact_kind=True, slide_area=slide_area)
        return score + _variant_bias(ps, variant, avoid)

    return min(spec.patterns, key=total)


def choose_with_score(spec: TemplateSpec, content: SlideContent, used: Counter, variant: int = 1,
                      avoid: str | None = None) -> tuple[PatternSpec, float]:
    """Как choose_pattern, но ещё и насколько хорошо образец подошёл (без смещения варианта)."""
    ps = choose_pattern(spec, content, used, variant, avoid)
    fallback = FALLBACK.get(content.kind, [])
    tier = 0.0 if ps.kind == content.kind else (1.0 + fallback.index(ps.kind) if ps.kind in fallback else 10.0)
    return ps, tier + _score(ps, content, used, exact_kind=True, slide_area=spec.tokens.slide_w * spec.tokens.slide_h)


# ── варианты вёрстки ───────────────────────────────────────────────────────
# Ось различий: одно содержимое, разный выбор образцов и форма графиков — все три по правилам шаблона.
#   1 «Сбалансированный» — лучший образец для каждого слайда;
#   2 «Визуальный»       — образцы с иллюстрациями, иконками, крупными числами; меньше текста;
#   3 «Альтернативный»   — другой образец, чем в первом варианте, если есть сопоставимый по качеству.

VARIANTS = {1: "Сбалансированный", 2: "Визуальный", 3: "Альтернативный"}
# форма графика по вариантам: столбцы -> полосы -> линия (для рядов во времени), круг <-> кольцо
_CHART_FORM = {
    2: {"column": "bar", "line": "column", "pie": "doughnut", "doughnut": "pie", "bar": "column"},
    3: {"column": "line", "bar": "line", "pie": "doughnut", "doughnut": "pie", "line": "column"},
}
_TIME = re.compile(r"^(19|20)\d\d|^Q[1-4]|^[IV]+\s|квартал|янв|фев|мар|апр|май|июн|июл|авг|сен|окт|ноя|дек", re.IGNORECASE)


def _visual_weight(ps: PatternSpec) -> int:
    p = ps.pattern
    icons = sum(1 for it in (p.groups[0].items if p.groups else []) for sl in it if sl.kind in (SlotKind.icon, SlotKind.image))
    pics = sum(1 for pic in p.pictures if not pic.from_layout)
    numbers = sum(1 for sl in p.slots if sl.kind == SlotKind.number) + sum(
        1 for it in (p.groups[0].items[:1] if p.groups else []) for sl in it if sl.kind == SlotKind.number)
    return min(icons, 3) + min(pics, 2) + min(numbers, 2) + sum(1 for sl in p.slots if sl.kind == SlotKind.image)


def _variant_bias(ps: PatternSpec, variant: int, avoid: str | None) -> float:
    # смещения мягкие: вариант отличается там, где есть сопоставимая замена, но не ценой потери смысла
    if variant == 2:
        return -0.6 * min(_visual_weight(ps), 4)
    if variant == 3 and avoid is not None and ps.pattern.id == avoid:
        return 2.0                               # тот же образец, что в первом варианте, — только если нет замены
    return 0.0


def chart_for_variant(chart, variant: int):
    """Альтернативная форма того же графика; линия — только для рядов во времени."""
    if chart is None or variant not in _CHART_FORM:
        return chart
    new_type = _CHART_FORM[variant].get(chart.type, chart.type)
    if new_type == "line" and not all(_TIME.search(c) for c in chart.categories):
        new_type = "bar"
    if new_type in ("pie", "doughnut") and (len(chart.series) != 1):
        return chart
    return chart.model_copy(update={"type": new_type})


def select_patterns(spec: TemplateSpec, contents: list[SlideContent], variant: int = 1,
                    avoid: list[str] | None = None) -> list[PatternSpec]:
    used: Counter = Counter()
    picks = []
    for i, content in enumerate(contents):
        ps = choose_pattern(spec, content, used, variant, avoid[i] if avoid else None)
        used[ps.pattern.id] += 1
        picks.append(ps)
    return picks


# ── раскладка ──────────────────────────────────────────────────────────────


def _numbering(sample: str, i: int) -> str:
    sample = sample.strip()
    if re.fullmatch(r"0\d", sample):
        return f"{i:02d}"
    return str(i)


def _nbsp(value: str | None) -> str | None:
    """«3 200», «1 150 ₽» не должны переноситься между разрядами и перед единицей."""
    return value.replace(" ", "\u00a0") if value else value


def _assign_item(slots: list[Slot], item: Item, index: int) -> dict[int, list[str]]:
    """Раскладывает поля элемента (value, heading, text) по текстовым слотам карточки.

    Число — в слот-число; заголовок+пояснение — в фигуру с двумя абзацами-образцами; остальное —
    по свободным слотам, чтобы пояснение не терялось, если отдельного текстового слота нет.
    """
    free = {"value": _nbsp(item.value), "heading": item.heading, "text": item.text}
    out: dict[int, list[str]] = {}

    def take(*names: str) -> str | None:
        for n in names:
            if free.get(n):
                value, free[n] = free[n], None
                return value
        return None

    textual = [s for s in slots if s.kind in _TEXTUAL]
    for s in textual:                                    # 1. числа
        if s.kind == SlotKind.number:
            v = take("value")
            out[s.shape_id] = [v] if v else ([_numbering(s.sample, index)] if s.sample else [])
    for s in textual:                                    # 2. заголовок и пояснение в одной фигуре
        if s.kind == SlotKind.heading and len(s.styles) >= 2 and free["text"] and (free["heading"] or free["value"]):
            out[s.shape_id] = [take("heading", "value"), take("text")]
    has_text_slot = any(s.kind in (SlotKind.text, SlotKind.caption) for s in textual)
    for s in textual:                                    # 3. остальные заголовки
        if s.kind == SlotKind.heading and s.shape_id not in out:
            v = take("heading", "value") or (None if has_text_slot else take("text"))
            out[s.shape_id] = [v] if v else []
    text_slots = [s for s in textual if s.kind in (SlotKind.text, SlotKind.caption) and s.shape_id not in out]
    for s in text_slots:                                 # 4. тексты и подписи
        if len(text_slots) == 1 and free["value"] and (free["text"] or free["heading"]):
            # слота под число нет — число первым абзацем, чтобы показатель не потерялся
            out[s.shape_id] = [take("value"), take("text", "heading")]
            continue
        if len(text_slots) == 1 and free["heading"] and free["text"]:
            # слот один, а у пункта и заголовок, и пояснение — оба абзацами, чтобы заголовок не пропал
            out[s.shape_id] = [take("heading"), take("text")]
            continue
        v = take("text", "heading", "value")      # число без своего слота не должно теряться
        out[s.shape_id] = [v] if v else []
    return out


def _chrome(spec: TemplateSpec) -> tuple[set[str], set[str]]:
    """Бегущий колонтитул и статичные номера страниц: текст в нижней/верхней полосе, повторённый на многих
    образцах. Возвращает (образцы колонтитула, образцы номеров) — их заменяем, а не удаляем."""
    h = spec.tokens.slide_h
    band = [s for ps in spec.patterns for s in ps.pattern.slots
            if s.sample and (s.bbox.y > h * 0.86 or s.bbox.bottom < h * 0.1)]
    texts = Counter(s.sample for s in band if s.kind != SlotKind.number)
    numbers = [s for s in band if s.kind == SlotKind.number and re.fullmatch(r"\d{1,2}", s.sample.strip())]
    need = max(3, len(spec.patterns) // 3)
    running = {t for t, n in texts.items() if n >= need}
    pages = {s.sample for s in numbers} if len(numbers) >= need else set()
    return running, pages


class _Composer:
    def __init__(self, spec: TemplateSpec, images: dict[str, bytes], footer: str = ""):
        self.spec = spec
        self.images = images
        self.running, self.pages = _chrome(spec)
        self.footer = footer
        self.index = 0
        t = spec.tokens
        self.body_font = t.fonts.body
        self.body_size = next((s.size_pt for s in t.type_scale if s.role == "body"), 14.0)
        accents = [c.hex for c in t.palette if "accent" in c.roles]
        first = t.accent() or (accents[0] if accents else "0077FF")
        self.pattern: Pattern | None = None
        self.chart_colors = [first] + [c for c in accents if delta_e(c, first) > 15][:4] + ["8F8F8F", "C4C4C4"]

    # ── текстовые слоты ──

    def _fill(self, slide, slot: Slot, texts: list[str], report: SlideReport | None = None,
              limit_bottom: int | None = None) -> bool:
        el = ops.find_shape(slide, slot.shape_id)
        if el is None:
            return False
        if not ops.set_paragraphs(el, texts):
            ops.remove_shape(slide, slot.shape_id)
            return False
        self._fit(el, slot, report, limit_bottom)
        return True

    def _steps(self, slot: Slot, size: float) -> list[float]:
        """Допустимые множители кегля: только ступени шкалы шаблона, не ниже пола для роли слота."""
        scale = self.spec.tokens.type_scale
        if slot.kind == SlotKind.title:
            # «дисплейный» образец (144 pt в «Q&A») можно ужать до обычного заголовка по шкале
            title_size = next((s.size_pt for s in scale if s.role == "title"), size)
            floor = max(min(size * 0.55, title_size), self.body_size * 1.15)
        elif slot.kind == SlotKind.number:
            floor = max(size * 0.3, self.body_size * 1.5)   # крупное число можно ужать сильнее текста
        elif slot.kind == SlotKind.heading:
            floor = self.body_size
        else:
            floor = min((s.size_pt for s in scale if s.usage >= 0.01), default=size * fit.MIN_SCALE)
        floor = floor if slot.kind in (SlotKind.title, SlotKind.number) else max(floor, size * fit.MIN_SCALE)
        sizes = sorted({size} | {st for s in scale for st in (s.size_pt,) if floor <= st < size}, reverse=True)
        return [x / size for x in sizes]

    def _free_width(self, slot: Slot) -> int:
        """Ширина до ближайшего блока справа: в шаблоне рамка заголовка часто шире, чем ему можно, —
        короткий образец «Заголовок» до карточек не доставал, а настоящий текст доставал бы."""
        p = self.pattern
        if p is None:
            return slot.bbox.w
        b = slot.bbox
        others = [s.bbox for s in p.slots if s.shape_id != slot.shape_id]
        others += [ib for g in p.groups for ib in g.item_bboxes] + [pic.bbox for pic in p.pictures]
        if p.title and p.title.shape_id != slot.shape_id:
            others.append(p.title.bbox)
        edges = [o.x for o in others
                 if o.y < b.bottom and o.bottom > b.y                    # по высоте пересекаются
                 and b.x + 0.25 * b.w < o.x < b.right]                    # и блок стоит справа внутри рамки
        if not edges:
            return b.w
        return max(min(edges) - b.x - self.spec.tokens.slide_w // 100, int(b.w * 0.4))

    def _fit(self, el, slot: Slot, report: SlideReport | None, limit_bottom: int | None) -> None:
        paras = fit.paragraph_specs(el, slot.styles)
        if not paras:
            return
        t = self.spec.tokens
        width = slot.bbox.w
        if ops.no_wrap(el):                  # без переноса строка тянется до правого поля
            width = t.slide_w - t.margins.right - slot.bbox.x
        elif limit_bottom is None:           # одиночные блоки и заголовок — с учётом соседей справа
            width = self._free_width(slot)
            if width < slot.bbox.w or slot.kind == SlotKind.title:
                # у заголовка рамка уже «безопасная» (без логотипа и декора) — ставим её фигуре
                ops.set_box(el, BBox(x=slot.bbox.x, y=slot.bbox.y, w=width, h=slot.bbox.h))
        height = slot.bbox.h
        # заголовок держим в своей рамке (она бывает прижата книзу — рост вверх наехал бы на логотип);
        # остальное растёт вниз — до низа карточки или нижнего поля слайда
        if slot.autofit == "shape" and slot.kind != SlotKind.title:
            bottom = limit_bottom or (t.slide_h - t.margins.bottom)
            height = max(height, bottom - slot.bbox.y)
        result = fit.fit(paras, width, height, slot.insets, self._steps(slot, max(p.size_pt for p in paras)))
        if result.scale != 1.0:
            fit.apply_scale(el, paras, result.scale)
        if not result.fits and report is not None:
            text = " / ".join(p.text for p in paras)
            report.warnings.append(f"текст не помещается ({slot.kind.value}): «{text[:60]}»")
        if slot.kind == SlotKind.title:
            top_inset = (slot.insets or fit.DEFAULT_INSETS)[1]
            self.title_bottom = slot.bbox.y + top_inset + int(result.height_pt * fit.EMU_PER_PT)
            self._fit_plate(slot, paras, width, result)

    def _fit_plate(self, slot: Slot, paras, width: int, result) -> None:
        """«Таблетка» под заголовком была по длине образца «ЗАГОЛОВОК» — подгоняем под настоящий текст."""
        plate = self.pattern.title_plate if self.pattern else None
        el = ops.find_shape(self.slide, plate.shape_id) if plate else None
        if el is None:
            return
        l_in, t_in, r_in, _ = slot.insets or fit.DEFAULT_INSETS
        width_pt = (width - l_in - r_in) / fit.EMU_PER_PT
        text_w = int(fit.widest_line_pt(paras, width_pt, result.scale) * fit.EMU_PER_PT)
        pad_x = max(slot.bbox.x + l_in - plate.bbox.x, 0)
        pad_y = max(slot.bbox.y + t_in - plate.bbox.y, 0)
        t = self.spec.tokens
        new_w = min(max(2 * pad_x + text_w, int(plate.bbox.w * 0.5)), t.slide_w - t.margins.right - plate.bbox.x)
        text_h = int(result.height_pt * fit.EMU_PER_PT)
        new_h = max(plate.bbox.h, 2 * pad_y + text_h) if result.lines > 1 else plate.bbox.h
        ops.set_box(el, BBox(x=plate.bbox.x, y=plate.bbox.y, w=new_w, h=new_h))
        if new_h > plate.bbox.h:
            self.title_bottom = max(self.title_bottom, plate.bbox.y + new_h)

    def _push_below_title(self) -> None:
        """Заголовок вырос — всё, что стояло прямо под ним, сдвигаем вниз на величину переполнения."""
        p = self.pattern
        if not p or not p.title:
            return
        t = self.spec.tokens
        delta = self.title_bottom + t.slide_h // 60 - p.title.bbox.bottom
        if delta <= 0:
            return
        tb = p.title.bbox
        below = [s for s in p.slots if s.bbox.y >= tb.bottom - t.slide_h // 50]
        below += [sl for g in p.groups for it in g.items for sl in it if sl.bbox.y >= tb.bottom - t.slide_h // 50]
        first_top = min((s.bbox.y for s in below), default=None)
        if first_top is None or first_top > self.title_bottom + t.slide_h // 60:
            return                                     # под заголовком и так есть зазор
        for s in below:
            if s.bbox.x < tb.right and s.bbox.right > tb.x and s.bbox.bottom + delta <= t.slide_h - t.margins.bottom:
                el = ops.find_shape(self.slide, s.shape_id)
                if el is not None:
                    ops.offset_shape(el, 0, delta)

    def compose_synth(self, slide, canvas: PatternSpec, content: SlideContent, report: SlideReport, roles,
                      force_dark: bool) -> None:
        """Своя композиция на холсте шаблона: заголовок и подзаголовок — в слоты рамки, остальное — рисуем."""
        from prism.layout import synth

        p = canvas.pattern
        self.pattern, self.slide = p, slide
        self.title_bottom = p.title.bbox.bottom if p.title else self.spec.tokens.margins.top
        subtitle = synth._subtitle(p)
        head = {p.title.shape_id} | ({subtitle.shape_id} if subtitle else set()) | (
            {p.title_plate.shape_id} if p.title_plate else set())
        chrome = self._chrome_slots(p)
        synth.clear_canvas(slide, p, self.spec.tokens, head | {s.shape_id for s in chrome}, head)
        self._fill_chrome(slide, p)
        t = self.spec.tokens
        on_dark = force_dark or bool(p.background and relative_luminance(p.background) < 0.3)
        if force_dark:                                   # акцентный фон из палитры поверх рамки основного фона
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = synth._rgb(roles.emph_bg)

        self._fill(slide, p.title, [content.title], report)
        top = self.title_bottom
        lead = content.subtitle if content.kind not in (SlideKind.section, SlideKind.closing, SlideKind.quote) else None
        if subtitle is not None:
            if lead and self._fill(slide, subtitle, [lead], report):
                top = max(top, subtitle.bbox.bottom)
            elif not lead:
                ops.remove_shape(slide, subtitle.shape_id)
        if force_dark:
            for sid in (p.title.shape_id, subtitle.shape_id if subtitle else None):
                el = ops.find_shape(slide, sid) if sid else None
                if el is not None:
                    ops.set_text_color(el, roles.on_emph)

        gap = t.slide_h // 25
        # низ зоны — над колонтитулом/номером, если они внизу
        bottom = min([t.slide_h - t.margins.bottom] + [c.bbox.y - gap // 2 for c in chrome
                                                       if c.bbox.y > t.slide_h * 0.5])
        area = synth.Area(box=BBox(x=t.margins.left, y=top + gap, w=t.slide_w - t.margins.left - t.margins.right,
                                   h=bottom - top - gap * 3 // 2), gap=gap)
        items = content.items
        kind = content.kind
        if kind == SlideKind.kpi and items:
            ok = synth.kpi_row(slide, area, items[:4], roles, on_dark)
        elif kind == SlideKind.chart and content.chart:
            ok = synth.chart_focus(slide, area, content, roles,
                                   lambda sl, box, ch: ops.add_chart(sl, box, ch, roles.series, self.body_font,
                                                                     roles.ink, self.body_size))
        elif kind in (SlideKind.process, SlideKind.timeline) and len(items) >= 2:
            ok = synth.steps(slide, area, items[:6], roles, timeline=kind == SlideKind.timeline)
        elif kind == SlideKind.table and content.table:
            rows = len(content.table.rows) + 1
            box = BBox(x=area.box.x, y=area.box.y, w=area.box.w,
                       h=min(area.box.h, int(self.body_size * 2.4 * fit.EMU_PER_PT) * rows))
            ops.add_table(slide, box, content.table, roles.accent, roles.on_emph, roles.ink, self.body_font,
                          max(self.body_size - 1, 9), roles.surface)
            ok = True
        elif kind in (SlideKind.section, SlideKind.closing, SlideKind.quote):
            ok = synth.statement(slide, area, content, roles, on_dark)
        elif kind == SlideKind.bullets and items:
            ok = synth.bullets(slide, area, items[:6], roles)
        elif items:
            ok = synth.cards_grid(slide, area, items[:6], roles)
        else:
            ok = synth.statement(slide, area, content, roles, on_dark)
        if not ok:
            report.warnings.append(f"текст не помещается (своя композиция {kind.value})")
        self._push_below_title()
        if content.notes:
            slide.notes_slide.notes_text_frame.text = content.notes

    def _chrome_slots(self, p: Pattern) -> list[Slot]:
        h = self.spec.tokens.slide_h
        return [s for s in p.slots if (s.sample in self.running or s.sample in self.pages)
                and (s.bbox.y > h * 0.86 or s.bbox.bottom < h * 0.1)]

    def _fill_chrome(self, slide, p: Pattern) -> set[int]:
        """Колонтитул — короткое название колоды (в регистре образца), номер — настоящий номер слайда."""
        done = set()
        for s in self._chrome_slots(p):
            if s.sample in self.pages:
                text_ = f"{self.index:0{len(s.sample.strip())}d}"
            else:
                text_ = self.footer.upper() if s.sample.isupper() else self.footer
            if text_ and self._fill(slide, s, [text_]):
                done.add(s.shape_id)
        return done

    def _drop_orphan_plates(self, slide) -> None:
        """Плашка («ГОД 20__», полоса квартала, блок дерева) без текста внутри — пустой прямоугольник.
        Если все текстовые слоты, что лежали на ней, убраны, а заполненных не осталось, — убираем и её."""
        p = self.pattern
        t = self.spec.tokens
        textual = [s for s in p.slots if s.kind in _TEXTUAL] + [
            s for g in p.groups for it in g.items for s in it if s.kind in _TEXTUAL]
        if p.title:
            textual.append(p.title)
        gone = [s.bbox for s in textual if ops.find_shape(slide, s.shape_id) is None]
        alive = [s.bbox for s in textual if ops.find_shape(slide, s.shape_id) is not None]
        if not gone:
            return

        def inside(inner: BBox, outer: BBox) -> bool:
            tol = t.slide_w // 150
            cx, cy = inner.x + inner.w / 2, inner.y + inner.h / 2
            return outer.x - tol <= cx <= outer.right + tol and outer.y - tol <= cy <= outer.bottom + tol

        for box in p.decor_boxes:
            b = box.bbox
            if box.service or box.bleed or b.w > 0.95 * t.slide_w or b.h > 0.5 * t.slide_h:
                continue
            if b.h < t.slide_h // 100:
                # тонкая линия-разделитель: относится к тексту сразу под/над ней
                pad = t.slide_h // 12
                b = BBox(x=b.x, y=b.y - pad, w=b.w, h=b.h + 2 * pad)
            if any(inside(g, b) for g in gone) and not any(inside(a, b) for a in alive):
                ops.remove_shape(slide, box.shape_id)

    def _clear(self, slide, slot: Slot) -> None:
        if slot.kind in _TEXTUAL or (slot.kind == SlotKind.image and slot.sample):
            ops.remove_shape(slide, slot.shape_id)

    def _fill_group(self, slide, g: RepeatGroup, items: list[Item], start: int = 1,
                    report: SlideReport | None = None) -> None:
        """items — не больше g.count; start — номер первого элемента (для нумерации «01», «02»…)."""
        for i, item_slots in enumerate(g.items):
            if i >= len(items):
                if g.fixed_count:                      # подложка из макета остаётся — чистим только текст
                    for s in item_slots:
                        if s.kind in _TEXTUAL:
                            ops.remove_shape(slide, s.shape_id)
                else:
                    for s in item_slots:
                        ops.remove_shape(slide, s.shape_id)
                continue
            texts = _assign_item(item_slots, items[i], start + i)
            limit = g.item_bboxes[i].bottom               # текст карточки не должен вылезать за карточку
            for s in item_slots:
                if s.kind in _TEXTUAL:
                    self._fill(slide, s, texts.get(s.shape_id, []), report, limit_bottom=limit)
                elif s.kind == SlotKind.image:
                    ops.remove_shape(slide, s.shape_id)   # фото-заглушки в карточках без своих фото убираем

    # ── области для графика/таблицы без готового слота ──

    def _content_area(self, p: Pattern) -> BBox:
        """Куда поставить график/таблицу, если своего слота нет: большая картинка, большой текстовый
        блок или область группы — и всегда ниже заголовка, в пределах полей."""
        t = self.spec.tokens
        images = [s.bbox for s in p.slots if s.kind == SlotKind.image]
        texts = [s.bbox for s in p.slots if s.kind in (SlotKind.text, SlotKind.caption)]
        groups = [b for b in (p.groups[0].item_bboxes if p.groups else [])]
        if images:
            box = max(images, key=lambda b: b.area)
        elif groups:
            x, y = min(b.x for b in groups), min(b.y for b in groups)
            box = BBox(x=x, y=y, w=max(b.right for b in groups) - x, h=max(b.bottom for b in groups) - y)
        elif texts and max(b.area for b in texts) > 0.15 * t.slide_w * t.slide_h:
            box = max(texts, key=lambda b: b.area)
        else:
            box = BBox(x=t.margins.left, y=0, w=t.slide_w - t.margins.left - t.margins.right,
                       h=t.slide_h - t.margins.bottom)
        top = max(box.y, self.title_bottom + t.slide_h // 40)
        left = max(box.x, t.margins.left)
        right = min(box.right, t.slide_w - t.margins.right)
        bottom = min(box.bottom, t.slide_h - t.margins.bottom)
        return BBox(x=left, y=top, w=max(right - left, t.slide_w // 4), h=max(bottom - top, t.slide_h // 4))

    def _below_title(self, box: BBox, p: Pattern) -> BBox:
        t = self.spec.tokens
        top = max(box.y, self.title_bottom + t.slide_h // 40)
        bottom = min(box.bottom, t.slide_h - t.margins.bottom)
        left, right = max(box.x, t.margins.left), min(box.right, t.slide_w - t.margins.right)
        return BBox(x=left, y=top, w=right - left, h=max(bottom - top, t.slide_h // 4))

    def _text_color(self, p: Pattern) -> str:
        pairs = [tp for tp in self.spec.tokens.text_pairs
                 if p.background and delta_e(tp.backdrop, p.background) < 8]
        if pairs:
            return pairs[0].text
        dark = p.background and relative_luminance(p.background) < 0.3
        return "FFFFFF" if dark else (self.spec.tokens.color("text") or "000000")

    # ── слайд целиком ──

    def compose(self, slide, ps: PatternSpec, content: SlideContent, report: SlideReport) -> None:
        p = ps.pattern
        self.pattern = p
        self.slide = slide
        self.title_bottom = p.title.bbox.bottom if p.title else self.spec.tokens.margins.top
        chrome = self._fill_chrome(slide, p)
        if p.title:
            self._fill(slide, p.title, [content.title], report)

        # крупные одиночные числа («герой» слайда) берут первые показатели, подпись — из блока под числом
        items = list(content.items)
        singles = [s for s in p.slots if s.kind in _TEXTUAL and s.shape_id not in chrome]
        taken: set[int] = set()
        heroes = sorted((s for s in singles if s.kind == SlotKind.number), key=lambda s: -s.bbox.area)
        while heroes and items and items[0].value:
            hero, item = heroes.pop(0), items.pop(0)
            self._fill(slide, hero, [_nbsp(item.value)], report)
            taken.add(hero.shape_id)
            below = [s for s in singles if s.shape_id not in taken and s.kind != SlotKind.number
                     and hero.bbox.bottom - hero.bbox.h // 2 <= s.bbox.y <= hero.bbox.bottom + 2 * hero.bbox.h
                     and s.bbox.x < hero.bbox.right and s.bbox.right > hero.bbox.x]
            caption = item.text or item.heading
            if below and caption:
                target = min(below, key=lambda s: s.bbox.y)
                self._fill(slide, target, [caption], report)
                taken.add(target.shape_id)

        # остальные пункты раскладываются по группам образца: основная берёт сколько вмещает, следующие — остаток
        placed_items = bool(p.groups) and bool(items)
        rest, start = items, 1
        placed = 0
        for gi, g in enumerate(p.groups):
            only_numbers = all(sl.kind == SlotKind.number for it in g.items for sl in it)
            if gi and only_numbers and placed:
                # отдельная группа номеров (кружки «1…4» цикла) — нумеруем, пунктов она не забирает
                self._fill_group(slide, g, [Item() for _ in range(min(placed, g.count))], 1, report)
                continue
            chunk, rest = rest[: g.count], rest[g.count:]
            self._fill_group(slide, g, chunk, start, report)
            start += len(chunk)
            placed += len(chunk)
        if placed_items and rest:
            report.warnings.append(f"не поместилось пунктов: {len(rest)} из {len(items)}")

        # подзаголовок — в верхний текстовый слот, список пунктов (если группы нет) — в самый большой
        subtitle_text = content.subtitle or (f"«{content.quote}»" if content.quote else None)
        by_top = sorted(singles, key=lambda s: (s.bbox.y, s.bbox.x))
        if items and not placed_items:
            target = max((s for s in singles if s.kind != SlotKind.number and s.shape_id not in taken),
                         key=lambda s: s.bbox.area, default=None)
            if target is not None:
                lines = [f"{i.heading}: {i.text}" if i.heading and i.text else (i.heading or i.text or i.value or "")
                         for i in items]
                self._fill(slide, target, lines, report)
                taken.add(target.shape_id)
                placed_items = True
        if subtitle_text:
            target = next((s for s in by_top if s.shape_id not in taken and s.kind != SlotKind.number), None)
            if target is not None:
                self._fill(slide, target, [subtitle_text], report)
                taken.add(target.shape_id)
        if content.author:
            target = next((s for s in by_top if s.shape_id not in taken and s.kind != SlotKind.number), None)
            if target is not None:
                self._fill(slide, target, [content.author], report)
                taken.add(target.shape_id)
        if items and not placed_items:
            report.warnings.append("пункты некуда поставить: в образце нет ни группы, ни текстового блока")
        # одиночные порядковые номера («1», «02» в кружках цикла) — часть схемы шагов: при заполненных шагах
        # оставляем их как в образце, лишние (шагов меньше) убираем
        ordinals = sorted((s for s in singles if s.kind == SlotKind.number and s.shape_id not in taken
                           and re.fullmatch(r"0?\d", s.sample.strip())), key=lambda s: int(s.sample))
        if placed_items and ordinals:
            n_steps = len(content.items)
            for s in ordinals:
                if int(s.sample) <= n_steps:
                    taken.add(s.shape_id)
        for s in singles:
            if s.shape_id not in taken:
                self._clear(slide, s)

        # картинка пользователя — в самый большой слот-картинку, остальные заглушки убираем
        image_slots = sorted((s for s in p.slots if s.kind == SlotKind.image), key=lambda s: -s.bbox.area)
        for i, s in enumerate(image_slots):
            el = ops.find_shape(slide, s.shape_id)
            if i == 0 and content.image and content.image in self.images and el is not None and not s.sample:
                ops.replace_picture(slide, el, self.images[content.image], s.bbox)
            else:
                ops.remove_shape(slide, s.shape_id)

        # график и таблица: в свой слот, иначе на место контентной области
        chart_slots = [s for s in p.slots if s.kind == SlotKind.chart]
        table_slots = [s for s in p.slots if s.kind == SlotKind.table]
        for s in chart_slots + table_slots:
            ops.remove_shape(slide, s.shape_id)        # копии чужих данных не оставляем
        text_color = self._text_color(p)
        example = None
        own = [pic for pic in p.pictures if not pic.from_layout and not pic.bleed]
        if ps.kind == SlideKind.chart and own and not (content.chart or content.table):
            # в образце-графике крупные картинки — это примеры чужих диаграмм; без своих данных — убираем
            for pic in own:
                ops.remove_shape(slide, pic.shape_id)
            own = []
        if (content.chart or content.table) and not (chart_slots or table_slots) \
                and ps.kind in (SlideKind.chart, SlideKind.image) and own:
            # примеры графиков картинками убираем все, настоящий график встаёт на их общее место
            example = own[0]
            x, y = min(pic.bbox.x for pic in own), min(pic.bbox.y for pic in own)
            r, b = max(pic.bbox.right for pic in own), max(pic.bbox.bottom for pic in own)
            example = example.model_copy(update={"bbox": BBox(x=x, y=y, w=r - x, h=b - y)})
            for pic in own:
                ops.remove_shape(slide, pic.shape_id)
        if content.chart:
            box = self._below_title(chart_slots[0].bbox, p) if chart_slots else (self._below_title(example.bbox, p) if example
                                                            else self._content_area(p))
            ops.add_chart(slide, box, content.chart, self.chart_colors, self.body_font, text_color, self.body_size)
        if content.table:
            box = self._below_title(table_slots[0].bbox, p) if table_slots else (
                self._below_title(example.bbox, p) if example else self._content_area(p))
            # высота строки — как в образце таблицы (или разумная по кеглю), а не растянутая на всю рамку
            rows = len(content.table.rows) + 1
            sample_rows = table_slots[0].table_rows if table_slots and table_slots[0].table_rows else 0
            row_h = box.h // sample_rows if sample_rows else int(self.body_size * 2.4 * fit.EMU_PER_PT)
            box = BBox(x=box.x, y=box.y, w=box.w, h=min(box.h, row_h * rows))
            ops.add_table(slide, box, content.table, self.chart_colors[0], "FFFFFF", text_color, self.body_font,
                          max(self.body_size - 2, 9), None)

        self._drop_orphan_plates(slide)
        self._push_below_title()
        if content.notes:
            slide.notes_slide.notes_text_frame.text = content.notes


# когда вариант берёт свою композицию вместо образца шаблона
POOR_FIT = 7.0          # оценка образца хуже этой — своя композиция честнее, чем натянутый образец
_SYNTH_VISUAL = {SlideKind.kpi, SlideKind.chart, SlideKind.timeline, SlideKind.process,
                 SlideKind.section, SlideKind.closing, SlideKind.quote}
_SYNTH_ANY = _SYNTH_VISUAL | {SlideKind.cards, SlideKind.bullets, SlideKind.table, SlideKind.agenda,
                              SlideKind.two_column, SlideKind.comparison}
_EMPHASIS = {SlideKind.section, SlideKind.closing, SlideKind.quote}
SYNTH_KINDS = _SYNTH_ANY                 # типы, которые вёрстка соберёт сама, даже если в шаблоне нет образца


def synth_mode(content: SlideContent, score: float, variant: int) -> bool:
    if content.kind not in _SYNTH_ANY:
        return False                     # титул, картинка, команда — только образцы шаблона
    if content.kind == SlideKind.chart and content.chart is None:
        return False
    if variant == 2:
        return content.kind in _SYNTH_VISUAL or (content.kind == SlideKind.cards and len(content.items) <= 4) \
            or score >= POOR_FIT
    if variant == 3:
        return content.kind == SlideKind.table or score >= POOR_FIT
    return score >= POOR_FIT and content.kind not in _EMPHASIS


def _short_title(title: str, limit: int = 32) -> str:
    """Для колонтитула: начало заголовка до двоеточия/тире, не длиннее limit."""
    head = re.split(r"[:—–.]", title, maxsplit=1)[0].strip()
    if len(head) <= limit:
        return head
    cut = head[:limit].rsplit(" ", 1)[0]
    return cut or head[:limit]


def compose_deck(template_pptx: str | Path, spec: TemplateSpec, contents: list[SlideContent],
                 images: dict[str, bytes] | None = None, variant: int = 1,
                 avoid: list[str] | None = None) -> tuple[bytes, list[SlideReport]]:
    """Колода по шаблону: .pptx в байтах + отчёт о раскладке (для аудита).
    variant — вариант вёрстки (1..3), avoid — образцы первого варианта по слайдам (для третьего)."""
    from prism.layout import synth

    prs = Presentation(str(template_pptx))
    originals = list(prs.slides)
    footer = _short_title(contents[0].title) if contents else ""
    composer = _Composer(spec, images or {}, footer)
    canvases = {False: synth.pick_canvas(spec), True: synth.pick_canvas(spec, dark=True)}
    roles = synth.roles(spec.tokens)
    used: Counter = Counter()
    reports = []
    for i, content in enumerate(contents, 1):
        if content.chart is not None and variant != 1:
            content = content.model_copy(update={"chart": chart_for_variant(content.chart, variant)})
        composer.index = i
        ps, score = choose_with_score(spec, content, used, variant, avoid[i - 1] if avoid else None)
        canvas = canvases[False]
        if synth_mode(content, score, variant) and canvas is not None:
            dark = variant == 2 and content.kind in _EMPHASIS
            base = canvases[True] if dark and canvases[True] else canvas
            slide = ops.duplicate_slide(prs, originals[base.pattern.slide_index - 1])
            report = SlideReport(index=i, pattern=f"synth:{content.kind.value}<-{base.pattern.id}", kind=content.kind.value)
            composer.compose_synth(slide, base, content, report, roles, dark and base is not canvases[True])
        else:
            used[ps.pattern.id] += 1
            slide = ops.duplicate_slide(prs, originals[ps.pattern.slide_index - 1])
            report = SlideReport(index=i, pattern=ps.pattern.id, kind=ps.kind.value)
            composer.compose(slide, ps, content, report)
        reports.append(report)
    ops.drop_slides(prs, len(originals))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), reports
