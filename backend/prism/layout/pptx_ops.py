"""Низкоуровневые операции над .pptx, которых нет в python-pptx: клонирование слайда шаблона,
замена текста с сохранением оформления, удаление фигур, картинки с обрезкой, нативные графики/таблицы.
"""

from __future__ import annotations

import io
from copy import deepcopy

from lxml import etree
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Emu, Pt

from prism.parsing.model import BBox

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
# связи, которые не копируем: макет у нового слайда свой, заметки пишем заново,
# графики пересоздаём нативно (копия общей chart-части сломала бы данные соседних слайдов)
_SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.CHART}


# ── слайды ─────────────────────────────────────────────────────────────────


def duplicate_slide(prs: Presentation, src: Slide) -> Slide:
    """Копия слайда-образца в конец колоды: фон, все фигуры, картинки, гиперссылки."""
    dst = prs.slides.add_slide(src.slide_layout)
    for shape in list(dst.shapes):                       # плейсхолдеры, которые add_slide создал сам
        shape._element.getparent().remove(shape._element)

    rid_map: dict[str, str] = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in _SKIP_RELS:
            continue
        if rel.is_external:
            rid_map[rid] = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = dst.part.relate_to(rel.target_part, rel.reltype)

    src_csld, dst_csld = src._element.cSld, dst._element.cSld
    bg = src_csld.find(qn("p:bg"))
    if bg is not None:
        dst_csld.insert(0, deepcopy(bg))
    tree = dst.shapes._spTree
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        tree.append(deepcopy(el))

    for el in dst._element.iter():
        for attr, value in el.attrib.items():
            if attr.startswith(f"{{{R_NS}}}") and value in rid_map:
                el.set(attr, rid_map[value])
    return dst


def drop_slides(prs: Presentation, keep_from: int) -> None:
    """Убирает первые keep_from слайдов (исходные образцы шаблона); их части не попадут в файл."""
    id_list = prs.slides._sldIdLst
    for sld_id in list(id_list)[:keep_from]:
        prs.part.drop_rel(sld_id.rId)
        id_list.remove(sld_id)


# ── фигуры ─────────────────────────────────────────────────────────────────


def find_shape(slide: Slide, shape_id: int) -> etree._Element | None:
    for c_nv_pr in slide.shapes._spTree.iter(qn("p:cNvPr")):
        if c_nv_pr.get("id") == str(shape_id):
            return c_nv_pr.getparent().getparent()
    return None


def remove_shape(slide: Slide, shape_id: int) -> None:
    el = find_shape(slide, shape_id)
    if el is None:
        return
    parent = el.getparent()
    parent.remove(el)
    # опустевшая группа — тоже мусор
    if parent.tag == qn("p:grpSp") and not any(
        c.tag in (qn("p:sp"), qn("p:pic"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:cxnSp")) for c in parent
    ):
        parent.getparent().remove(parent)


def shape_xfrm(el: etree._Element) -> etree._Element | None:
    for path in ("p:spPr/a:xfrm", "p:grpSpPr/a:xfrm", "p:xfrm"):
        found = el.find(path, {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": A_NS})
        if found is not None:
            return found
    return None


def set_box(el: etree._Element, box: BBox) -> None:
    """Задаёт фигуре положение и размер; плейсхолдеру без своего xfrm (наследует от макета) — создаёт его."""
    xfrm = shape_xfrm(el)
    if xfrm is None:
        sp_pr = el.find(qn("p:spPr"))
        if sp_pr is None:
            return
        xfrm = etree.Element(qn("a:xfrm"))
        sp_pr.insert(0, xfrm)
    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    if off is None:
        off = etree.SubElement(xfrm, qn("a:off"))
    if ext is None:
        ext = etree.SubElement(xfrm, qn("a:ext"))
    off.set("x", str(box.x))
    off.set("y", str(box.y))
    ext.set("cx", str(box.w))
    ext.set("cy", str(box.h))


def no_wrap(el: etree._Element) -> bool:
    body_pr = el.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    return body_pr is not None and body_pr.get("wrap") == "none"


def alignment(el: etree._Element) -> str:
    """Выравнивание первого абзаца: l | ctr | r | just…; "?" — наследуется от плейсхолдера макета."""
    tx_body = el.find(qn("p:txBody"))
    if tx_body is None:
        return "?"
    p_pr = tx_body.find(f"{qn('a:p')}/{qn('a:pPr')}")
    if p_pr is not None and p_pr.get("algn"):
        return p_pr.get("algn")
    lvl = tx_body.find(f"{qn('a:lstStyle')}/{qn('a:lvl1pPr')}")
    if lvl is not None and lvl.get("algn"):
        return lvl.get("algn")
    return "?" if el.find(f".//{qn('p:nvPr')}/{qn('p:ph')}") is not None else "l"


def offset_shape(el: etree._Element, dx: int, dy: int) -> None:
    xfrm = shape_xfrm(el)
    off = xfrm.find(qn("a:off")) if xfrm is not None else None
    if off is not None:
        off.set("x", str(int(off.get("x")) + dx))
        off.set("y", str(int(off.get("y")) + dy))


# ── текст ──────────────────────────────────────────────────────────────────


def _styled_paragraphs(tx_body: etree._Element) -> list[etree._Element]:
    paras = tx_body.findall(qn("a:p"))
    styled = [p for p in paras if "".join(t.text or "" for t in p.iter(qn("a:t"))).strip()]
    return styled or paras[:1]


def _paragraph_like(template: etree._Element, text: str) -> etree._Element:
    """Абзац с оформлением template (pPr + rPr первого рана) и новым текстом."""
    p = etree.SubElement(etree.Element(qn("a:txBody")), qn("a:p"))
    ppr = template.find(qn("a:pPr"))
    if ppr is not None:
        p.append(deepcopy(ppr))
    first_run = template.find(qn("a:r"))
    run = etree.SubElement(p, qn("a:r"))
    rpr = first_run.find(qn("a:rPr")) if first_run is not None else None
    if rpr is None:
        end = template.find(qn("a:endParaRPr"))
        rpr = deepcopy(end) if end is not None else etree.Element(qn("a:rPr"))
        rpr.tag = qn("a:rPr")
    else:
        rpr = deepcopy(rpr)
    rpr.set("lang", "ru-RU")
    rpr.attrib.pop("dirty", None)
    run.append(rpr)
    etree.SubElement(run, qn("a:t")).text = text
    end = template.find(qn("a:endParaRPr"))
    if end is not None:
        p.append(deepcopy(end))
    return p


def set_paragraphs(el: etree._Element, texts: list[str]) -> bool:
    """Заменяет текст фигуры: i-й текст получает оформление i-го абзаца образца (лишние — последнего).
    False — если текста нет (фигуру стоит удалить)."""
    texts = [t.strip() for t in texts if t and t.strip()]
    tx_body = el.find(qn("p:txBody"))
    if tx_body is None or not texts:
        return False
    templates = _styled_paragraphs(tx_body)
    if not templates:
        templates = [etree.Element(qn("a:p"))]
    new = [_paragraph_like(templates[min(i, len(templates) - 1)], t) for i, t in enumerate(texts)]
    for p in tx_body.findall(qn("a:p")):
        tx_body.remove(p)
    for p in new:
        tx_body.append(p)
    return True


def set_text_color(el: etree._Element, hex_color: str) -> None:
    """Перекрашивает весь текст фигуры (заголовок на акцентном тёмном фоне)."""
    tx_body = el.find(qn("p:txBody"))
    if tx_body is None:
        return
    for rpr in [*tx_body.iter(qn("a:rPr")), *tx_body.iter(qn("a:endParaRPr"))]:
        for old in rpr.findall(qn("a:solidFill")):
            rpr.remove(old)
        fill = etree.Element(qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hex_color)
        # solidFill должен идти раньше шрифтов (latin/ea/cs) по схеме DrawingML
        anchor = next((c for c in rpr if etree.QName(c).localname in ("latin", "ea", "cs", "sym", "hlinkClick")), None)
        if anchor is not None:
            anchor.addprevious(fill)
        else:
            rpr.append(fill)


def paragraph_count(el: etree._Element) -> int:
    tx_body = el.find(qn("p:txBody"))
    return len(_styled_paragraphs(tx_body)) if tx_body is not None else 0


# ── картинки ───────────────────────────────────────────────────────────────


def replace_picture(slide: Slide, el: etree._Element, image: bytes, box: BBox) -> None:
    """Меняет картинку в pic/фигуре с фото-заливкой; обрезает под пропорции рамки, не растягивая."""
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(image))
    blip_fill = el.find(".//" + qn("p:blipFill"))
    if blip_fill is None:
        blip_fill = el.find(".//" + qn("a:blipFill"))
    if blip_fill is None:
        return
    blip = blip_fill.find(qn("a:blip"))
    blip.set(f"{{{R_NS}}}embed", rid)
    with Image.open(io.BytesIO(image)) as im:
        img_ratio = im.width / im.height
    box_ratio = box.w / box.h if box.h else img_ratio
    src_rect = blip_fill.find(qn("a:srcRect"))
    if src_rect is None:
        src_rect = etree.Element(qn("a:srcRect"))
        blip.addnext(src_rect)
    src_rect.attrib.clear()
    if img_ratio > box_ratio:          # картинка шире — срезаем бока
        cut = round((1 - box_ratio / img_ratio) / 2 * 100000)
        src_rect.set("l", str(cut))
        src_rect.set("r", str(cut))
    elif img_ratio < box_ratio:        # выше — срезаем верх и низ
        cut = round((1 - img_ratio / box_ratio) / 2 * 100000)
        src_rect.set("t", str(cut))
        src_rect.set("b", str(cut))


# ── графики и таблицы ──────────────────────────────────────────────────────

_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def add_chart(slide: Slide, box: BBox, chart, colors: list[str], font: str | None, text_color: str,
              size_pt: float) -> None:
    data = CategoryChartData()
    data.categories = chart.categories
    for s in chart.series:
        data.add_series(s.name, [float(v) for v in s.values[: len(chart.categories)]])
    frame = slide.shapes.add_chart(_CHART_TYPES[chart.type], Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h), data)
    frame.name = "prism:chart"
    c = frame.chart
    c.font.size = Pt(size_pt)
    c.font.color.rgb = RGBColor.from_string(text_color)
    if font:
        c.font.name = font
    round_chart = chart.type in ("pie", "doughnut")
    c.has_legend = round_chart or len(chart.series) > 1
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM
        c.legend.include_in_layout = False
    plot = c.plots[0]
    if round_chart:
        points = plot.series[0].points
        for i in range(len(chart.categories)):
            points[i].format.fill.solid()
            points[i].format.fill.fore_color.rgb = RGBColor.from_string(colors[i % len(colors)])
        plot.has_data_labels = True
        plot.data_labels.number_format = '0"' + (chart.unit or "") + '"' if chart.unit == "%" else "General"
        plot.data_labels.number_format_is_linked = False
    else:
        for i, series in enumerate(plot.series):
            fmt = series.format
            if chart.type == "line":
                fmt.line.color.rgb = RGBColor.from_string(colors[i % len(colors)])
                fmt.line.width = Pt(2.25)
            else:
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = RGBColor.from_string(colors[i % len(colors)])
        value_axis = c.value_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(text_color)
        value_axis.major_gridlines.format.line.width = Pt(0.25)
        value_axis.format.line.fill.background()
        values = [float(v) for s in chart.series for v in s.values]
        if chart.type in ("column", "bar") and values and min(values) >= 0:
            value_axis.minimum_scale = 0      # столбцы от нуля: иначе 94 против 89 выглядит как «вдвое»
        if chart.unit:
            value_axis.has_title = True
            value_axis.axis_title.text_frame.text = chart.unit
        c.category_axis.format.line.color.rgb = RGBColor.from_string(text_color)


def add_table(slide: Slide, box: BBox, table, header_fill: str, header_text: str, body_text: str,
              font: str | None, size_pt: float, stripe: str | None) -> None:
    rows, cols = len(table.rows) + 1, len(table.columns)
    frame = slide.shapes.add_table(rows, cols, Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    frame.name = "prism:table"
    tbl = frame.table
    tbl.first_row = True
    for r in range(rows):
        values = table.columns if r == 0 else table.rows[r - 1]
        for ci in range(cols):
            cell = tbl.cell(r, ci)
            cell.text = str(values[ci]) if ci < len(values) else ""
            fill = header_fill if r == 0 else (stripe if stripe and r % 2 == 0 else None)
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            else:
                cell.fill.background()
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size_pt)
                    run.font.bold = r == 0
                    run.font.color.rgb = RGBColor.from_string(header_text if r == 0 else body_text)
                    if font:
                        run.font.name = font
