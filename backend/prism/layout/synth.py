"""Собственные композиции на дизайн-системе шаблона — там, где образца нет или варианту нужна другая подача.

Слайд строится на «холсте»: образце шаблона, из которого убран контент, но сохранена рамка (фон, логотип,
колонтитул, номер страницы, линия под заголовком). Поверх — нативные блоки: только цвета палитры, шрифты
и ступени кегля шаблона, поля и сетка — из его токенов. Так композиция своя, а правила — шаблона.
"""

from __future__ import annotations

from contextvars import ContextVar
from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

from prism.color import chroma, contrast_ratio, delta_e, relative_luminance
from prism.layout import fit
from prism.layout import pptx_ops as ops
from prism.parsing.model import BBox
from prism.parsing.patterns import Pattern, SlotKind
from prism.parsing.spec import PatternSpec, TemplateSpec
from prism.parsing.tokens import DesignTokens
from prism.planning.schemas import Item, SlideContent, SlideKind

EMU_PT = 12700


# ── роли цвета и типографики ───────────────────────────────────────────────


@dataclass
class Roles:
    bg: str
    ink: str             # основной текст на фоне
    muted: str           # второстепенный текст
    accent: str
    accent_text: str     # акцент для мелкого текста: читается и на фоне, и на подложке
    series: list[str]    # цвета серий/маркеров — из палитры, различимые между собой
    surface: str         # подложка карточек на основном фоне
    emph_bg: str         # акцентный (обычно тёмный) фон
    on_emph: str         # текст на акцентном фоне
    head_font: str | None
    body_font: str | None
    title: float
    heading: float
    body: float
    caption: float
    display: float


def _pick(colors: list[str], key) -> str | None:
    return min(colors, key=key) if colors else None


def roles(t: DesignTokens) -> Roles:
    bg = next((b.color for b in t.backgrounds if b.color), None) or t.color("background") or "FFFFFF"
    dark_bg = relative_luminance(bg) < 0.3
    on_bg = [p for p in t.text_pairs if delta_e(p.backdrop, bg) < 6 and p.share >= 0.02]
    readable = [p.text for p in on_bg if contrast_ratio(p.text, bg) >= 3]
    ink = _pick(readable, key=lambda c: -contrast_ratio(c, bg)) or ("FFFFFF" if dark_bg else "000000")
    accent = t.accent() or ink

    series: list[str] = []
    for c in [accent] + [p.hex for p in t.palette if p.usage > 0.003]:
        # цвета серий подписывают шаги и маркеры — должны читаться на фоне хотя бы как крупный текст (3:1)
        if chroma(c) >= 18 and 0.02 < relative_luminance(c) < 0.75 and contrast_ratio(c, bg) >= 3 \
                and all(delta_e(c, s) > 18 for s in series):
            series.append(c)
    series = series[:5] or [accent]

    surfaces = [p.hex for p in t.palette if ("surface" in p.roles or "background" in p.roles)
                and delta_e(p.hex, bg) > 3 and (relative_luminance(p.hex) < 0.3) == dark_bg]
    # подложка карточек — ближайший к фону различимый цвет: серый #BFBFBF на почти белом фоне давит на текст
    surface = min(surfaces, key=lambda c: delta_e(c, bg)) if surfaces else bg
    # второстепенный текст — только если читается и на фоне, и на подложке (WCAG 4.5:1 для мелкого текста)
    muted = next((c for c in readable if c != ink and chroma(c) < 25
                  and min(contrast_ratio(c, bg), contrast_ratio(c, surface)) >= 4.5), ink)
    # акцентный цвет мелким текстом (номера карточек) — если не читается на подложке, берём основной текст
    accent_text = accent if min(contrast_ratio(accent, bg), contrast_ratio(accent, surface)) >= 4.5 else ink

    dark_candidates = [b.color for b in t.backgrounds if b.color and b.dark and delta_e(b.color, bg) > 8]
    emph_bg = dark_candidates[0] if dark_candidates else accent
    if contrast_ratio(emph_bg, "FFFFFF") < 3 and contrast_ratio(emph_bg, "000000") < 3:
        emph_bg = ink
    on_emph_pairs = [p.text for p in t.text_pairs
                     if delta_e(p.backdrop, emph_bg) < 6 and contrast_ratio(p.text, emph_bg) >= 4.5]
    on_emph = on_emph_pairs[0] if on_emph_pairs else (
        "FFFFFF" if contrast_ratio("FFFFFF", emph_bg) >= contrast_ratio("000000", emph_bg) else "000000")

    def size(role: str, default: float) -> float:
        steps = [s.size_pt for s in t.type_scale if s.role == role and s.usage >= 0.005]
        return max(steps) if role in ("title", "display") and steps else (steps[0] if steps else default)

    body = size("body", 14)
    return Roles(bg=bg, ink=ink, muted=muted, accent=accent, accent_text=accent_text, series=series, surface=surface,
                 emph_bg=emph_bg,
                 on_emph=on_emph, head_font=t.fonts.heading, body_font=t.fonts.body,
                 title=size("title", 32), heading=max(size("heading", body * 1.25), body * 1.1), body=body,
                 caption=min(size("caption", body * 0.85), body), display=max(size("display", body * 3), body * 2.5))


# ── холст ──────────────────────────────────────────────────────────────────


def _is_top_title(p: Pattern, t: DesignTokens) -> bool:
    return bool(p.title) and p.title.bbox.y < t.slide_h * 0.3 and p.title.bbox.h < t.slide_h * 0.35


def pick_canvas(spec: TemplateSpec, dark: bool = False) -> PatternSpec | None:
    """Образец-рамка: заголовок сверху, основной (или тёмный — для акцентных слайдов) фон, мало крупного
    декора в зоне контента. Предпочтительно — с подзаголовком под заголовком."""
    t = spec.tokens
    base_bg = roles(t).bg
    best, best_score = None, float("inf")
    for ps in spec.patterns:
        p = ps.pattern
        # титул годится только как тёмная рамка акцентного слайда (обложка-«раздел»), и то не первый
        if not _is_top_title(p, t) or (ps.kind == SlideKind.title and not (dark and p.slide_index > 1)):
            continue
        is_dark = bool(p.background and relative_luminance(p.background) < 0.3)
        same_bg = bool(p.background and delta_e(p.background, base_bg) < 6)
        if dark and not (is_dark and not same_bg):
            continue
        if not dark and not same_bg:
            continue
        pics = sum(1 for pic in p.pictures if not pic.bleed and pic.bbox.y > p.title.bbox.bottom)
        # высокий заголовок-раздел (низ на середине слайда) оставляет контенту узкую полосу — график сплющится
        low = max(0.0, p.title.bbox.bottom / t.slide_h - 0.25) * 12
        score = pics * 3 + (0 if _subtitle(p) else 1.5) + (1 if ps.kind == SlideKind.closing else 0) + low
        score += 1.0 if ps.kind == SlideKind.section and not dark else 0.0
        # крупные плашки/рамки фото в зоне контента, что не уберутся с холста (стоят выше низа заголовка)
        slide_area = t.slide_w * t.slide_h
        big = [d for d in p.decor_boxes if not d.bleed and not d.service and d.bbox.area > 0.08 * slide_area
               and d.bbox.w < 0.9 * t.slide_w and d.bbox.y <= p.title.bbox.bottom and d.bbox.bottom > p.title.bbox.bottom]
        score += 3.0 * len(big) + (2.0 if ps.kind in (SlideKind.team, SlideKind.image) else 0.0)
        # картинки-объекты, которые холст не уберёт (из макета или «навылет»): 3D-иконка в углу и т. п.
        score += 3.0 * sum(1 for pic in p.pictures if (pic.from_layout or pic.bleed)
                           and pic.bbox.area < 0.35 * slide_area and pic.bbox.bottom > p.title.bbox.bottom)
        if score < best_score:
            best, best_score = ps, score
    return best


def _subtitle(p: Pattern):
    """Текстовый блок сразу под заголовком (подзаголовок-лид), если он есть."""
    if not p.title:
        return None
    tb = p.title.bbox
    cands = [s for s in p.slots if s.kind in (SlotKind.text, SlotKind.caption, SlotKind.heading)
             and 0 <= s.bbox.y - tb.bottom < tb.h * 1.2 and s.bbox.x < tb.x + tb.w * 0.3
             and s.bbox.h <= tb.h * 1.5]                  # не текстовое поле на полслайда
    return min(cands, key=lambda s: s.bbox.y, default=None)


def clear_canvas(slide, p: Pattern, t: DesignTokens, keep: set[int], head: set[int] | None = None) -> None:
    """Убирает с копии образца весь контент, оставляя рамку: слоты, группы, крупные картинки и декор в зоне
    контента уходят; служебные плейсхолдеры, фон во всю ширину и всё у нижнего края — остаются."""
    for s in p.slots:
        if s.shape_id not in keep:
            ops.remove_shape(slide, s.shape_id)
    for g in p.groups:
        for it in g.items:
            for s in it:
                ops.remove_shape(slide, s.shape_id)
    for pic in p.pictures:
        if not pic.from_layout and not pic.bleed:
            ops.remove_shape(slide, pic.shape_id)
    # верх зоны контента — по заголовку и подзаголовку (head), а не по колонтитулу внизу
    head = head if head is not None else keep
    top = max((ks.bbox.bottom for ks in p.slots + ([p.title] if p.title else []) if ks.shape_id in head),
              default=t.margins.top)
    kept = [s.bbox for s in p.slots if s.shape_id in keep and s.shape_id not in head]    # колонтитул, номер
    for box in p.decor_boxes:
        b = box.bbox
        if box.shape_id in keep or box.service or box.bleed or b.w > 0.9 * t.slide_w:
            continue
        holds_chrome = any(b.x <= k.x + k.w / 2 <= b.right and b.y <= k.y + k.h / 2 <= b.bottom for k in kept)
        # остатки контента образца: стрелки, линии, плашки, рамка «Вставить фото» у нижнего края;
        # полосу колонтитула (фигуру под номером страницы, линию у самого низа) не трогаем
        if b.y > top + t.slide_h // 60 and b.y < t.slide_h * 0.86 and not holds_chrome:
            ops.remove_shape(slide, box.shape_id)


# ── примитивы ──────────────────────────────────────────────────────────────


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def rect(slide, box: BBox, fill: str, rounded: bool = True, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    if rounded:
        shape.adjustments[0] = min(0.12, 0.25 * 360000 / max(min(box.w, box.h), 1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    shape.name = "prism:plate"                  # нарисовано вёрсткой — аудит проверяет поля и цвета
    return shape


def oval(slide, cx: int, cy: int, d: int, fill: str | None, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(cx - d // 2), Emu(cy - d // 2), Emu(d), Emu(d))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    shape.name = "prism:marker"
    return shape


def hline(slide, x1: int, x2: int, y: int, color: str, width_pt: float = 1.0):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y), Emu(x2), Emu(y))
    conn.line.color.rgb = _rgb(color)
    conn.line.width = Pt(width_pt)
    conn.name = "prism:line"
    return conn


# шкала кеглей шаблона для текущей вёрстки: свои композиции берут кегли только с неё (аудит size_off_scale)
TYPE_SCALE: ContextVar[tuple[float, ...]] = ContextVar("TYPE_SCALE", default=())


def use_scale(tokens: DesignTokens) -> None:
    """Задать шкалу кеглей для вёрстки в текущем потоке (compose_deck вызывает перед сборкой)."""
    TYPE_SCALE.set(tuple(sorted({round(s.size_pt, 1) for s in tokens.type_scale if s.usage >= 0.005})))


def _on_scale(size: float, scale: tuple[float, ...], down: int = 0) -> float:
    """Ступень шкалы не больше size (иначе — самая мелкая), и ещё down ступеней вниз."""
    if not scale:
        return size * (0.92 ** down)
    below = [s for s in scale if s <= size * 1.02] or [scale[0]]
    i = max(scale.index(below[-1]) - down, 0)
    return scale[i]


def text(slide, box: BBox, paras: list[tuple[str, str | None, float, str, bool]], align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, min_scale: float = 0.6) -> bool:
    """paras: (текст, шрифт, кегль, цвет, жирный). Кегли — ступени шкалы шаблона; не влезает — все абзацы
    спускаются на ступень ниже (не мельче min_scale от исходного). Возвращает, влез ли текст."""
    paras = [p for p in paras if p[0]]
    if not paras or box.w <= 0 or box.h <= 0:
        return True
    scale = TYPE_SCALE.get()
    width_pt, height_pt = box.w / fit.EMU_PER_PT, box.h / fit.EMU_PER_PT
    sizes, fits = [p[2] for p in paras], False
    for down in range(8):
        cand = [_on_scale(size, scale, down) for _, _, size, _, _ in paras]
        if down and any(c < size * min_scale for c, (_, _, size, _, _) in zip(cand, paras, strict=True)):
            break
        sizes = cand
        specs = [fit.ParaSpec(text=t, font=f, size_pt=sz, bold=b, spacing=1.15)
                 for (t, f, _, _, b), sz in zip(paras, sizes, strict=True)]
        _, h, too_wide = fit.measure(specs, width_pt)
        if h <= height_pt and not too_wide:
            fits = True
            break
    tb = slide.shapes.add_textbox(Emu(box.x), Emu(box.y), Emu(box.w), Emu(box.h))
    tb.name = "prism:text"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ((t, font, _, color, bold), size) in enumerate(zip(paras, sizes, strict=True)):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if i:
            para.space_before = Pt(size * 0.35)
        run = para.add_run()
        run.text = t
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
        if font:
            run.font.name = font
    return fits


# ── композиции ─────────────────────────────────────────────────────────────


@dataclass
class Area:
    box: BBox
    gap: int


def _columns(area: Area, n: int) -> list[BBox]:
    w = (area.box.w - area.gap * (n - 1)) // n
    return [BBox(x=area.box.x + i * (w + area.gap), y=area.box.y, w=w, h=area.box.h) for i in range(n)]


def _grid(area: Area, n: int) -> list[BBox]:
    cols = n if n <= 3 else (2 if n == 4 else 3)
    rows = (n + cols - 1) // cols
    w = (area.box.w - area.gap * (cols - 1)) // cols
    h = (area.box.h - area.gap * (rows - 1)) // rows
    return [BBox(x=area.box.x + (i % cols) * (w + area.gap), y=area.box.y + (i // cols) * (h + area.gap), w=w, h=h)
            for i in range(n)]


def kpi_row(slide, area: Area, items: list[Item], r: Roles, on_dark: bool = False) -> bool:
    """Крупные числа в ряд: число акцентом, подпись под ним, тонкая линия-акцент сверху."""
    ok = True
    ink = r.on_emph if on_dark else r.ink
    num_color = r.on_emph if on_dark else r.accent
    cols = _columns(area, len(items))
    for col, it in zip(cols, items, strict=True):
        hline(slide, col.x, col.x + min(col.w, col.w // 2), col.y, num_color, 2.5)
        num_h = int(col.h * 0.42)
        ok &= text(slide, BBox(x=col.x, y=col.y + area.gap // 2, w=col.w, h=num_h),
                   [(it.value or it.heading or "", r.head_font, r.display, num_color, True)], anchor=MSO_ANCHOR.BOTTOM,
                   min_scale=0.4)
        ok &= text(slide, BBox(x=col.x, y=col.y + area.gap // 2 + num_h + area.gap // 3, w=col.w,
                               h=col.h - num_h - area.gap),
                   [(it.heading if it.value and it.heading else "", r.head_font, r.heading, ink, True),
                    (it.text or "", r.body_font, r.body, ink if on_dark else r.muted, False)])
    return ok


def cards_grid(slide, area: Area, items: list[Item], r: Roles) -> bool:
    """Карточки на подложке из палитры: номер/значение акцентом, заголовок, текст."""
    ok = True
    for i, (cell, it) in enumerate(zip(_grid(area, len(items)), items, strict=True), 1):
        rect(slide, cell, r.surface)
        pad = max(area.gap // 2, int(min(cell.w, cell.h) * 0.08))
        inner = BBox(x=cell.x + pad, y=cell.y + pad, w=cell.w - 2 * pad, h=cell.h - 2 * pad)
        badge = it.value or f"{i:02d}"
        badge_h = int(r.heading * 1.6 * EMU_PT)
        ok &= text(slide, BBox(x=inner.x, y=inner.y, w=inner.w, h=badge_h),
                   [(badge, r.head_font, r.heading * (1.3 if it.value else 1), r.accent_text, True)])
        ok &= text(slide, BBox(x=inner.x, y=inner.y + badge_h + pad // 3, w=inner.w, h=inner.h - badge_h - pad // 3),
                   [(it.heading or "", r.head_font, r.heading, r.ink, True), (it.text or "", r.body_font, r.body,
                                                                              r.muted, False)])
    return ok


def steps(slide, area: Area, items: list[Item], r: Roles, timeline: bool) -> bool:
    """Шаги/таймлайн: линия-ось, маркеры цветами серий, подпись над осью (дата/номер), текст под ней."""
    ok = True
    n = len(items)
    cols = _columns(area, n)
    axis_y = area.box.y + int(area.box.h * 0.32)
    hline(slide, area.box.x, area.box.right, axis_y, r.muted, 1.0)
    d = int(min(r.heading * 1.9 * EMU_PT, cols[0].w * 0.4))
    for i, (col, it) in enumerate(zip(cols, items, strict=True), 1):
        color = r.series[(i - 1) % len(r.series)]
        cx = col.x + d // 2
        if timeline:
            oval(slide, cx, axis_y, d // 2, color)
            label = it.value or it.heading or str(i)
            ok &= text(slide, BBox(x=col.x, y=area.box.y, w=col.w, h=axis_y - area.box.y - d // 2),
                       [(label, r.head_font, r.heading * 1.2, color, True)], anchor=MSO_ANCHOR.BOTTOM)
            below = [(it.heading if it.value else "", r.head_font, r.body, r.ink, True),
                     (it.text or "", r.body_font, r.body, r.muted, False)]
        else:
            oval(slide, cx, axis_y, d, r.bg, line=color)
            text(slide, BBox(x=cx - d // 2, y=axis_y - d // 2, w=d, h=d), [(str(i), r.head_font, r.heading, color, True)],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            below = [(it.heading or "", r.head_font, r.heading, r.ink, True),
                     (it.text or "", r.body_font, r.body, r.muted, False)]
        ok &= text(slide, BBox(x=col.x, y=axis_y + d // 2 + area.gap // 2, w=col.w - area.gap // 3,
                               h=area.box.bottom - axis_y - d // 2 - area.gap // 2), below)
    return ok


def statement(slide, area: Area, content: SlideContent, r: Roles, on_dark: bool) -> bool:
    """Главная мысль крупно: для разделов, цитат и финала; на акцентном фоне — светлым."""
    color = r.on_emph if on_dark else r.ink
    main = content.quote and f"«{content.quote.strip('«»\"')}»" or content.subtitle or ""
    if main.strip("«»\" ").lower() == content.title.strip("«»\" ").lower():
        main = content.subtitle if content.quote and content.subtitle != content.title else ""   # не повторять заголовок
    lines = [(main, r.head_font, r.title * 0.9, color, False)]
    if content.author:
        lines.append((content.author, r.body_font, r.body, color, False))
    for it in content.items[:4]:
        lines.append((" — ".join(x for x in (it.heading, it.text) if x), r.body_font, r.body, color, False))
    box = BBox(x=area.box.x, y=area.box.y, w=int(area.box.w * 0.75), h=area.box.h)
    return text(slide, box, lines, anchor=MSO_ANCHOR.MIDDLE)


def chart_focus(slide, area: Area, content: SlideContent, r: Roles, add_chart) -> bool:
    """График на 62% ширины, справа — вывод крупно и пояснение."""
    left_w = int(area.box.w * 0.62)
    add_chart(slide, BBox(x=area.box.x, y=area.box.y, w=left_w, h=area.box.h), content.chart)
    side = BBox(x=area.box.x + left_w + area.gap, y=area.box.y, w=area.box.w - left_w - area.gap, h=area.box.h)
    rect(slide, side, r.surface)
    pad = area.gap // 2
    inner = BBox(x=side.x + pad, y=side.y + pad, w=side.w - 2 * pad, h=side.h - 2 * pad)
    lines = [(content.subtitle or "", r.head_font, r.heading * 1.15, r.ink, True)]
    lines += [(f"{i.value} — {i.text}" if i.value else (i.text or ""), r.body_font, r.body, r.muted, False)
              for i in content.items[:3]]
    return text(slide, inner, lines, anchor=MSO_ANCHOR.MIDDLE)


def bullets(slide, area: Area, items: list[Item], r: Roles) -> bool:
    """Список с маркерами-акцентами: крупный текст, воздух между пунктами."""
    ok = True
    n = len(items)
    row_h = area.box.h // max(n, 1)
    for i, it in enumerate(items):
        y = area.box.y + i * row_h
        color = r.series[i % len(r.series)]
        rect(slide, BBox(x=area.box.x, y=y + row_h // 6, w=int(r.body * 0.35 * EMU_PT), h=row_h * 2 // 3), color,
             rounded=False)
        ok &= text(slide, BBox(x=area.box.x + area.gap, y=y, w=area.box.w - area.gap, h=row_h),
                   [(it.heading or "", r.head_font, r.heading, r.ink, True),
                    (it.text or "", r.body_font, r.body * 1.1, r.muted, False)], anchor=MSO_ANCHOR.MIDDLE)
    return ok
