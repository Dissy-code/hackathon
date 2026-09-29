"""MCP-сервер «pptx»: собственные инструменты для работы с шаблонами и колодами.

Инструменты:
  template_summary(path)                         — дизайн-токены и образцы шаблона (палитра, шрифты, типы слайдов);
  inspect_deck(path)                             — слайды и фигуры: id, тип, рамка в долях слайда, текст;
  audit_deck(path)                               — детерминированный аудит (переполнение, контраст, наложения…);
  render_slide(path, slide, dpi)                 — PNG слайда (для VLM-аудита и превью);
  set_shape_text(path, slide, shape_id, paragraphs, out_path) — заменить текст фигуры с сохранением оформления.

    python -m prism.mcp.pptx_server --port 8002

Файлы — только внутри каталога данных сервиса (PRISM_DATA_DIR, по умолчанию backend/data): пути наружу
отклоняются, чтобы инструмент нельзя было направить на произвольный файл.
"""

from __future__ import annotations

import argparse
import asyncio
import os
import tempfile
from pathlib import Path

from mcp.server.fastmcp import FastMCP, Image
from pptx import Presentation

from prism.audit.deterministic import audit_deck as _audit
from prism.config import BACKEND_DIR
from prism.layout import fit
from prism.layout import pptx_ops as ops
from prism.parsing.spec import parse_template_sync
from prism.render import pdf_to_png, pptx_to_pdf

DATA_DIR = Path(os.environ.get("PRISM_DATA_DIR", BACKEND_DIR / "data")).resolve()


def _path(path: str, must_exist: bool = True) -> Path:
    p = Path(path)
    p = (p if p.is_absolute() else DATA_DIR / p).resolve()
    if not p.is_relative_to(DATA_DIR):
        raise ValueError(f"путь вне каталога данных {DATA_DIR}: {path}")
    if p.suffix.lower() not in (".pptx", ".potx"):
        raise ValueError(f"нужен .pptx: {path}")
    if must_exist and not p.is_file():
        raise FileNotFoundError(path)
    return p


def template_summary(path: str) -> dict:
    """Дизайн-токены шаблона и его образцы слайдов: палитра с ролями, шрифты, шкала кеглей, поля,
    список образцов (id, тип, число карточек)."""
    spec = parse_template_sync(_path(path))
    t = spec.tokens
    return {
        "slide": {"w": t.slide_w, "h": t.slide_h},
        "palette": [{"hex": c.hex, "roles": c.roles, "usage": round(c.usage, 3)} for c in t.palette[:12]],
        "fonts": t.fonts.model_dump(),
        "type_scale": [{"role": s.role, "size_pt": s.size_pt} for s in t.type_scale],
        "patterns": [{"id": ps.pattern.id, "kind": ps.kind.value,
                      "cards": [g.count for g in ps.pattern.groups]} for ps in spec.patterns],
    }


def inspect_deck(path: str) -> dict:
    """Слайды колоды: для каждой фигуры — id, имя, тип, рамка (x, y, w, h в долях слайда) и текст."""
    prs = Presentation(str(_path(path)))
    w, h = int(prs.slide_width), int(prs.slide_height)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        shapes = []
        for s in slide.shapes:
            if s.left is None:
                continue
            shapes.append({
                "id": s.shape_id, "name": s.name, "type": str(s.shape_type).split(".")[-1].split(" ")[0].lower(),
                "box": [round(int(s.left) / w, 3), round(int(s.top) / h, 3), round(int(s.width) / w, 3),
                        round(int(s.height) / h, 3)],
                "text": s.text_frame.text[:500] if s.has_text_frame else "",
            })
        out.append({"slide": i, "shapes": shapes})
    return {"slides": out}


def audit_deck(path: str) -> dict:
    """Детерминированный аудит колоды: переполнение текста, выход за край, наложения, контраст,
    мелкий кегль, оставшийся текст-образец. Каждая находка: slide, shape_id, kind, severity, message."""
    return {"issues": [i.model_dump() for i in _audit(_path(path))]}


async def render_slide(path: str, slide: int, dpi: int = 60) -> Image:
    """PNG одного слайда (нумерация с 1)."""
    src = _path(path)
    with tempfile.TemporaryDirectory() as tmp:
        pdf = await asyncio.to_thread(pptx_to_pdf, src, Path(tmp))
        pngs = await asyncio.to_thread(pdf_to_png, pdf, Path(tmp), max(30, min(dpi, 150)))
        if not 1 <= slide <= len(pngs):
            raise ValueError(f"в колоде {len(pngs)} слайдов")
        return Image(data=pngs[slide - 1].read_bytes(), format="png")


def set_shape_text(path: str, slide: int, shape_id: int, paragraphs: list[str], out_path: str | None = None) -> dict:
    """Заменить текст фигуры: абзацы получают оформление абзацев-образцов фигуры, кегль ужимается под рамку.
    out_path — куда сохранить (по умолчанию — поверх исходного файла)."""
    src = _path(path)
    dst = _path(out_path, must_exist=False) if out_path else src
    prs = Presentation(str(src))
    if not 1 <= slide <= len(prs.slides):
        raise ValueError(f"в колоде {len(prs.slides)} слайдов")
    s = prs.slides[slide - 1]
    el = ops.find_shape(s, shape_id)
    if el is None:
        raise ValueError(f"на слайде {slide} нет фигуры {shape_id}")
    if not ops.set_paragraphs(el, paragraphs):
        raise ValueError("пустой текст: чтобы убрать фигуру, это делает вёрстка, а не правка текста")
    shape = next(x for x in s.shapes if x.shape_id == shape_id)
    paras = fit.paragraph_specs(el, [])
    result = None
    if paras and shape.width and shape.height:
        result = fit.fit(paras, int(shape.width), int(shape.height))
        if result.scale != 1.0:
            fit.apply_scale(el, paras, result.scale)
    prs.save(str(dst))
    return {"saved": str(dst.relative_to(DATA_DIR)), "fits": result.fits if result else True,
            "scale": result.scale if result else 1.0}


def build(host: str = "127.0.0.1", port: int = 8002) -> FastMCP:
    server = FastMCP("prism-pptx", host=host, port=port,
                     instructions="Шаблоны и колоды .pptx: токены шаблона, осмотр слайдов, аудит, рендер, правка текста.")
    for tool in (template_summary, inspect_deck, audit_deck, render_slide, set_shape_text):
        server.add_tool(tool)
    return server


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--host", default=os.environ.get("MCP_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.environ.get("MCP_PORT", "8002")))
    args = ap.parse_args()
    build(args.host, args.port).run(transport="streamable-http")


if __name__ == "__main__":
    main()
