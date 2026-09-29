"""Экспорт колоды в HTML: один самодостаточный файл, слайды — живые элементы, а не картинки.

Каждая фигура становится элементом с абсолютной позицией в долях слайда; текст — настоящий текст с кеглем
в единицах ширины слайда (cqw), поэтому колода масштабируется под окно без потери чёткости. Картинки
встроены как data:URI, таблицы — <table>, диаграммы — SVG по данным из pptx. Стили берутся из того же
извлечения, что и разбор шаблонов (prism.parsing.extract): наследование от макета и мастера, цвета темы.

    python -m prism.export.html deck.pptx deck.html

Упрощения: у фигур нестандартной геометрии рисуется только прямоугольная/овальная заливка, у градиента —
первый цвет, эффекты (тени, свечение) не переносятся. Для сверки есть PDF — он из того же pptx.
"""

from __future__ import annotations

import base64
import hashlib
import html
import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from prism.parsing.extract import extract
from prism.parsing.model import BBox, RunStyle, Shape, ShapeKind

_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_GOOGLE = {"Play", "Montserrat", "Lato", "Poppins", "Carlito", "Unbounded", "Inter", "Roboto", "Open Sans",
           "PT Sans", "PT Serif", "Manrope", "Rubik", "Golos Text", "Onest"}
_PALETTE = ["4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]


class _Ctx:
    def __init__(self, w: int, h: int):
        self.w, self.h = w, h
        self.fonts: set[str] = set()
        self.images: dict[str, str] = {}           # sha1 -> css-класс; каждая картинка в файле один раз
        self.image_css: list[str] = []

    def image_class(self, blob: bytes, content_type: str) -> str:
        key = hashlib.sha1(blob).hexdigest()
        if key not in self.images:
            cls = f"i{len(self.images)}"
            self.images[key] = cls
            self.image_css.append(f".{cls}{{background-image:url({_data_uri(blob, content_type)})}}")
        return self.images[key]

    def pct(self, b: BBox) -> str:
        return (f"left:{b.x / self.w * 100:.3f}%;top:{b.y / self.h * 100:.3f}%;"
                f"width:{b.w / self.w * 100:.3f}%;height:{b.h / self.h * 100:.3f}%;")

    def size(self, pt: float) -> str:
        return f"{pt * 12700 / self.w * 100:.3f}cqw"       # кегль в долях ширины слайда


def _elements(shapes) -> dict[int, object]:
    """id фигуры -> python-pptx фигура (включая вложенные в группы)."""
    out = {}
    for sh in shapes:
        out[sh.shape_id] = sh
        if sh.shape_type is not None and "GROUP" in str(sh.shape_type):
            out.update(_elements(sh.shapes))
    return out


def _blob(el, part) -> tuple[bytes, str] | None:
    blip = el.find(f".//{qn('a:blip')}")
    rid = blip.get(_R_EMBED) if blip is not None else None
    if not rid:
        return None
    try:
        img = part.related_part(rid)
    except KeyError:
        return None
    return img.blob, img.content_type


def _data_uri(blob: bytes, content_type: str) -> str:
    return f"data:{content_type};base64,{base64.b64encode(blob).decode()}"


def _anchor(el, pptx_shape) -> str:
    """Вертикальное выравнивание текста: из фигуры, иначе из плейсхолдера макета/мастера."""
    node = pptx_shape
    while node is not None:
        body = node._element.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
        if body is not None and body.get("anchor"):
            return {"t": "flex-start", "ctr": "center", "b": "flex-end"}.get(body.get("anchor"), "flex-start")
        node = getattr(node, "_base_placeholder", None) if getattr(node, "is_placeholder", False) else None
    return "flex-start"


def _run_css(style: RunStyle, ctx: _Ctx, scale: float) -> str:
    css = []
    if style.font:
        ctx.fonts.add(style.font)
        css.append(f"font-family:'{style.font}',sans-serif")
    if style.size_pt:
        css.append(f"font-size:{ctx.size(style.size_pt * scale)}")
    if style.color:
        css.append(f"color:#{style.color}")
    if style.bold:
        css.append("font-weight:700")
    if style.italic:
        css.append("font-style:italic")
    return ";".join(css)


def _text(shape: Shape, pptx_shape, ctx: _Ctx) -> str:
    """Абзацы и раны из XML, стиль каждого рана — разрешённый при извлечении."""
    tx = pptx_shape._element.find(qn("p:txBody"))
    if tx is None:
        return ""
    scale = shape.autofit_scale or 1.0
    paras_xml = tx.findall(qn("a:p"))
    out = []
    for i, para in enumerate(shape.paragraphs):
        p_el = paras_xml[i] if i < len(paras_xml) else None
        runs = []
        if p_el is not None:
            styles = iter(para.run_styles)
            for r in p_el:
                tag = etree.QName(r).localname
                if tag == "br":
                    runs.append("<br>")
                elif tag in ("r", "fld"):
                    t = r.findtext(qn("a:t")) or ""
                    style = next(styles, para.style) if t.strip() else para.style
                    runs.append(f'<span style="{_run_css(style, ctx, scale)}">{html.escape(t)}</span>')
        if not runs:
            runs = [html.escape(para.text) or "&#8203;"]
        align = {"ctr": "center", "r": "right", "just": "justify"}.get(para.align or "l", "left")
        base = _run_css(para.style, ctx, scale)
        bullet = "padding-left:1.1em;text-indent:-1.1em;" if para.bullet else ""
        mark = "•&nbsp;" if para.bullet else ""
        out.append(f'<p style="text-align:{align};{bullet}{base}">{mark}{"".join(runs)}</p>')
    return "".join(out)


def _geometry_css(shape: Shape) -> str:
    if shape.geometry == "ellipse":
        return "border-radius:50%;"
    if shape.geometry in ("roundRect", "round2SameRect", "snipRoundRect"):
        return f"border-radius:{min(shape.bbox.w, shape.bbox.h) / max(shape.bbox.w, 1) * 16.7:.2f}% / " \
               f"{min(shape.bbox.w, shape.bbox.h) / max(shape.bbox.h, 1) * 16.7:.2f}%;"
    return ""


def _picture(shape: Shape, pptx_shape, part, ctx: _Ctx) -> str:
    got = _blob(pptx_shape._element, part)
    if got is None:
        return ""
    cls = ctx.image_class(*got)
    crop = pptx_shape._element.find(f".//{qn('a:srcRect')}")
    cut = [int(crop.get(k, 0)) / 100000 for k in ("l", "t", "r", "b")] if crop is not None else [0, 0, 0, 0]
    vis_w, vis_h = max(1 - cut[0] - cut[2], 0.01), max(1 - cut[1] - cut[3], 0.01)
    # обрезка srcRect через размер и позицию фона: видимая часть картинки заполняет рамку
    pos_x = cut[0] / (cut[0] + cut[2]) * 100 if cut[0] + cut[2] > 0 else 0
    pos_y = cut[1] / (cut[1] + cut[3]) * 100 if cut[1] + cut[3] > 0 else 0
    return (f'<div class="s {cls}" role="img" style="{ctx.pct(shape.bbox)}{_geometry_css(shape)}{_rot(shape)}'
            f'background-size:{100 / vis_w:.3f}% {100 / vis_h:.3f}%;background-position:{pos_x:.2f}% {pos_y:.2f}%;'
            f'background-repeat:no-repeat"></div>')


def _rot(shape: Shape) -> str:
    return f"transform:rotate({shape.rotation}deg);" if shape.rotation else ""


def _table(shape: Shape, pptx_shape, ctx: _Ctx) -> str:
    tbl = pptx_shape.table
    total_w = sum(c.width for c in tbl.columns) or 1
    cols = "".join(f'<col style="width:{c.width / total_w * 100:.2f}%">' for c in tbl.columns)
    rows = []
    for r in tbl.rows:
        cells = []
        for cell in r.cells:
            tc = cell._tc
            fill = tc.find(f"{qn('a:tcPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}")
            bg = f"background:#{fill.get('val')};" if fill is not None else ""
            run = cell._tc.find(f".//{qn('a:rPr')}")
            size = int(run.get("sz")) / 100 if run is not None and run.get("sz") else 12
            color = run.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}") if run is not None else None
            fg = f"color:#{color.get('val')};" if color is not None else ""
            bold = "font-weight:700;" if run is not None and run.get("b") in ("1", "true") else ""
            cells.append(f'<td style="{bg}{fg}{bold}font-size:{ctx.size(size)}">{html.escape(cell.text)}</td>')
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return (f'<div class="s" style="{ctx.pct(shape.bbox)}"><table class="t"><colgroup>{cols}</colgroup>'
            f'{"".join(rows)}</table></div>')


def _series_color(series, i: int) -> str:
    clr = series._element.find(f"{qn('c:spPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}")
    if clr is None:
        clr = series._element.find(f"{qn('c:spPr')}/{qn('a:ln')}/{qn('a:solidFill')}/{qn('a:srgbClr')}")
    return clr.get("val") if clr is not None else _PALETTE[i % len(_PALETTE)]


def _chart(shape: Shape, pptx_shape, ctx: _Ctx, text_color: str) -> str:
    """Диаграмма как SVG по данным pptx: столбцы, полосы, линия, круг/кольцо."""
    chart = pptx_shape.chart
    plot = chart.plots[0]
    cats = [str(c) for c in plot.categories]
    series = list(plot.series)
    kind = str(chart.chart_type)
    W, H = 1000, 1000 * shape.bbox.h / max(shape.bbox.w, 1)
    font = W * 0.028
    parts = []
    if "PIE" in kind or "DOUGHNUT" in kind:
        vals = [max(float(v or 0), 0) for v in series[0].values]
        total = sum(vals) or 1
        cx, cy, r = W * 0.35, H / 2, min(W * 0.3, H * 0.42)
        a0 = -math.pi / 2
        points = series[0].points
        for i, v in enumerate(vals):
            a1 = a0 + v / total * 2 * math.pi
            clr = points[i].format.fill.fore_color.rgb if points[i].format.fill.type == 1 else _PALETTE[i % 6]
            large = 1 if a1 - a0 > math.pi else 0
            x0, y0, x1, y1 = cx + r * math.cos(a0), cy + r * math.sin(a0), cx + r * math.cos(a1), cy + r * math.sin(a1)
            parts.append(f'<path d="M{cx:.1f},{cy:.1f} L{x0:.1f},{y0:.1f} A{r:.1f},{r:.1f} 0 {large} 1 {x1:.1f},{y1:.1f} Z" '
                         f'fill="#{clr}"/>')
            parts.append(f'<rect x="{W * 0.72:.1f}" y="{H * 0.2 + i * font * 1.8:.1f}" width="{font:.1f}" height="{font:.1f}" '
                         f'fill="#{clr}"/><text x="{W * 0.72 + font * 1.5:.1f}" y="{H * 0.2 + i * font * 1.8 + font * 0.85:.1f}">'
                         f'{html.escape(cats[i] if i < len(cats) else "")} — {v:g}</text>')
            a0 = a1
        if "DOUGHNUT" in kind:
            parts.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="{r * 0.55:.1f}" fill="var(--slide-bg,#fff)"/>')
    else:
        vals = [[float(v or 0) for v in s.values] for s in series]
        top = max((v for row in vals for v in row), default=1) or 1
        horizontal = "BAR" in kind and "COLUMN" not in kind
        left, right, top_pad, bottom = W * 0.1, W * 0.97, H * 0.06, H * 0.84
        n, k = max(len(cats), 1), max(len(series), 1)
        for g in range(5):                                          # сетка
            y = bottom - (bottom - top_pad) * g / 4
            if not horizontal:
                parts.append(f'<line x1="{left:.1f}" x2="{right:.1f}" y1="{y:.1f}" y2="{y:.1f}" class="g"/>'
                             f'<text x="{left - font * 0.4:.1f}" y="{y + font * 0.35:.1f}" text-anchor="end">'
                             f'{top * g / 4:.3g}</text>')
        for si, row in enumerate(vals):
            color = _series_color(series[si], si)
            if "LINE" in kind:
                pts = " ".join(f"{left + (i + 0.5) * (right - left) / n:.1f},{bottom - v / top * (bottom - top_pad):.1f}"
                               for i, v in enumerate(row))
                parts.append(f'<polyline points="{pts}" fill="none" stroke="#{color}" stroke-width="{font * 0.25:.1f}"/>')
                continue
            for i, v in enumerate(row):
                slot = (right - left) / n if not horizontal else (bottom - top_pad) / n
                bw = slot * 0.7 / k
                if horizontal:
                    y = top_pad + i * slot + slot * 0.15 + si * bw
                    parts.append(f'<rect x="{left:.1f}" y="{y:.1f}" width="{v / top * (right - left):.1f}" '
                                 f'height="{bw:.1f}" fill="#{color}"/>')
                else:
                    x = left + i * slot + slot * 0.15 + si * bw
                    h = v / top * (bottom - top_pad)
                    parts.append(f'<rect x="{x:.1f}" y="{bottom - h:.1f}" width="{bw:.1f}" height="{h:.1f}" fill="#{color}"/>')
        for i, c in enumerate(cats):
            if horizontal:
                y = top_pad + (i + 0.5) * (bottom - top_pad) / n
                parts.append(f'<text x="{left - font * 0.4:.1f}" y="{y:.1f}" text-anchor="end">{html.escape(c)}</text>')
            else:
                x = left + (i + 0.5) * (right - left) / n
                parts.append(f'<text x="{x:.1f}" y="{bottom + font * 1.4:.1f}" text-anchor="middle">{html.escape(c)}</text>')
        if len(series) > 1:
            for si, s in enumerate(series):
                parts.append(f'<rect x="{left + si * W * 0.2:.1f}" y="{H - font * 1.1:.1f}" width="{font * 0.8:.1f}" '
                             f'height="{font * 0.8:.1f}" fill="#{_series_color(s, si)}"/><text x="{left + si * W * 0.2 + font:.1f}" '
                             f'y="{H - font * 0.35:.1f}">{html.escape(s.name or "")}</text>')
    title = html.escape(chart.chart_title.text_frame.text) if chart.has_title and chart.chart_title.has_text_frame else ""
    svg = (f'<svg viewBox="0 0 {W} {H:.0f}" preserveAspectRatio="xMidYMid meet" role="img" aria-label="{title or "диаграмма"}" '
           f'style="width:100%;height:100%;font-size:{font:.1f}px;fill:#{text_color}">'
           f'<style>.g{{stroke:#{text_color};stroke-opacity:.18;stroke-width:1}}</style>{"".join(parts)}</svg>')
    return f'<div class="s" style="{ctx.pct(shape.bbox)}">{svg}</div>'


def _shape_html(shape: Shape, pptx_shape, part, ctx: _Ctx, text_color: str) -> str:
    if shape.kind == ShapeKind.group or pptx_shape is None:
        return ""
    if shape.kind == ShapeKind.picture:
        return _picture(shape, pptx_shape, part, ctx)
    if shape.kind == ShapeKind.table:
        return _table(shape, pptx_shape, ctx)
    if shape.kind == ShapeKind.chart:
        return _chart(shape, pptx_shape, ctx, text_color)
    if shape.kind == ShapeKind.line:
        b = shape.bbox
        color = shape.line or text_color
        return (f'<div class="s" style="{ctx.pct(b)}border-top:max(1px,.12cqw) solid #{color};height:0;'
                f'{_rot(shape)}"></div>') if b.w >= b.h else (
            f'<div class="s" style="{ctx.pct(b)}border-left:max(1px,.12cqw) solid #{color};width:0;{_rot(shape)}"></div>')
    if shape.kind not in (ShapeKind.text, ShapeKind.shape):
        return ""
    css = [ctx.pct(shape.bbox), _geometry_css(shape), _rot(shape)]
    classes = ""
    if shape.fill_kind in ("solid", "gradient") and shape.fill:
        css.append(f"background:#{shape.fill};")
    elif shape.fill_kind == "picture":
        got = _blob(pptx_shape._element, part)
        if got is not None:
            classes = f" {ctx.image_class(*got)}"
            css.append("background-position:center;background-size:cover;")
    if shape.line:
        css.append(f"border:max(1px,.08cqw) solid #{shape.line};")
    body = ""
    if shape.kind == ShapeKind.text and not shape.prompt_text and shape.text.strip():
        l_, t_, r_, b_ = shape.insets or (91440, 45720, 91440, 45720)
        pad = (f"padding:{t_ / ctx.w * 100:.3f}cqw {r_ / ctx.w * 100:.3f}cqw {b_ / ctx.w * 100:.3f}cqw "
               f"{l_ / ctx.w * 100:.3f}cqw;")
        css.append(f"display:flex;flex-direction:column;justify-content:{_anchor(None, pptx_shape)};{pad}")
        body = _text(shape, pptx_shape, ctx)
    if not body and len(css) == 3 and not shape.line and not classes:
        return ""                                                   # невидимая фигура
    return f'<div class="s{classes}" style="{"".join(css)}">{body}</div>'


def _layer(shapes: list[Shape], elements: dict, part, ctx: _Ctx, text_color: str, skip_placeholders: bool) -> str:
    out = []
    for sh in sorted(shapes, key=lambda s: s.z):
        if skip_placeholders and sh.placeholder is not None:
            continue                                                # плейсхолдеры макета — подсказки, не контент
        out.append(_shape_html(sh, elements.get(sh.id), part, ctx, text_color))
    return "".join(out)


def _background(bg, owner_part, owner_el, ctx: _Ctx) -> tuple[str, str]:
    """(css фона, класс с картинкой фона)."""
    if bg is None:
        return "#FFFFFF", ""
    if bg.kind == "picture" and owner_el is not None:
        bg_pr = owner_el.find(f"{qn('p:cSld')}/{qn('p:bg')}/{qn('p:bgPr')}")
        got = _blob(bg_pr, owner_part) if bg_pr is not None else None
        if got is not None:
            return f"#{bg.color or 'FFFFFF'};background-size:cover;background-position:center", ctx.image_class(*got)
    return f"#{bg.color or 'FFFFFF'}", ""


def deck_to_html(pptx: str | Path, title: str = "Презентация") -> str:
    raw = extract(pptx)
    prs = Presentation(str(pptx))
    ctx = _Ctx(raw.slide_w, raw.slide_h)
    slides_html = []
    for slide, s_raw in zip(prs.slides, raw.slides, strict=True):
        layout, master = slide.slide_layout, slide.slide_layout.slide_master
        l_raw = next((lay for lay in raw.layouts if lay.name == s_raw.layout and lay.master == s_raw.master), None)
        m_raw = raw.masters[s_raw.master] if s_raw.master < len(raw.masters) else None
        owner = {"slide": slide, "layout": layout, "master": master}.get(s_raw.background.source, slide)
        bg, bg_cls = _background(s_raw.background, owner.part, owner._element, ctx)
        bg_color = f"#{s_raw.background.color or 'FFFFFF'}"
        text_color = "FFFFFF" if s_raw.background.color and _dark(s_raw.background.color) else "333333"
        layers = []
        show_master = layout._element.get("showMasterSp", "1") not in ("0", "false")
        if m_raw is not None and show_master:
            layers.append(_layer(m_raw.shapes, _elements(master.shapes), master.part, ctx, text_color, True))
        if l_raw is not None:
            layers.append(_layer(l_raw.shapes, _elements(layout.shapes), layout.part, ctx, text_color, True))
        layers.append(_layer(s_raw.shapes, _elements(slide.shapes), slide.part, ctx, text_color, False))
        notes = f'<aside class="notes">{html.escape(s_raw.notes)}</aside>' if s_raw.notes.strip() else ""
        slides_html.append(f'<section class="slide {bg_cls}" id="s{s_raw.index}" style="background-color:{bg};--slide-bg:{bg_color}" '
                           f'aria-label="Слайд {s_raw.index}">{"".join(layers)}</section>{notes}')
    fonts = sorted(f for f in ctx.fonts if f in _GOOGLE)
    link = ("<link rel=\"stylesheet\" href=\"https://fonts.googleapis.com/css2?"
            + "&".join(f"family={f.replace(' ', '+')}:ital,wght@0,400;0,700;1,400" for f in fonts)
            + "&display=swap\">") if fonts else ""
    return _PAGE.format(title=html.escape(title), fonts=link, ratio=f"{raw.slide_w} / {raw.slide_h}",
                        slides="\n".join(slides_html), count=len(slides_html), images="\n".join(ctx.image_css))


def _dark(hex_color: str) -> bool:
    r, g, b = (int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b < 0.4


def export_html(pptx: str | Path, out: str | Path, title: str = "Презентация") -> Path:
    out = Path(out)
    out.write_text(deck_to_html(pptx, title), encoding="utf-8")
    return out


_PAGE = """<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
{fonts}
<style>
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; background: #111; font-family: Arial, sans-serif; }}
  main {{ display: flex; flex-direction: column; align-items: center; gap: 24px; padding: 24px; }}
  .slide {{ position: relative; width: min(100%, 1280px); aspect-ratio: {ratio}; overflow: hidden;
            container-type: inline-size; box-shadow: 0 8px 30px rgba(0,0,0,.5); }}
  .s {{ position: absolute; margin: 0; overflow: visible; }}
  .s p {{ margin: 0; line-height: 1.15; white-space: pre-wrap; overflow-wrap: break-word; }}
  .t {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
  .t td {{ padding: .5cqw .8cqw; border-bottom: 1px solid rgba(0,0,0,.08); vertical-align: top; }}
  .notes {{ display: none; }}
  body.present {{ overflow: hidden; }}
  body.present main {{ padding: 0; gap: 0; height: 100vh; justify-content: center; }}
  body.present .slide {{ display: none; width: min(100vw, calc(100vh * {ratio})); box-shadow: none; }}
  body.present .slide.on {{ display: block; }}
  {images}
  .hint {{ position: fixed; right: 12px; bottom: 10px; color: #aaa; font: 12px Arial, sans-serif; }}
  body.present .hint {{ display: none; }}
</style>
</head>
<body>
<main>
{slides}
</main>
<div class="hint">{count} слайдов · F — показ, ← → — листать, Esc — выход</div>
<script>
  const slides = [...document.querySelectorAll('.slide')];
  let at = 0;
  const show = (i) => {{ at = Math.max(0, Math.min(slides.length - 1, i));
    slides.forEach((s, k) => s.classList.toggle('on', k === at)); }};
  addEventListener('keydown', (e) => {{
    if (e.key === 'f' || e.key === 'F') {{ document.body.classList.add('present'); show(at);
      document.documentElement.requestFullscreen?.().catch(() => {{}}); }}
    else if (e.key === 'Escape') document.body.classList.remove('present');
    else if (document.body.classList.contains('present')) {{
      if (['ArrowRight', 'ArrowDown', ' ', 'PageDown'].includes(e.key)) show(at + 1);
      if (['ArrowLeft', 'ArrowUp', 'PageUp'].includes(e.key)) show(at - 1);
    }}
  }});
  slides.forEach((s, k) => s.addEventListener('dblclick', () => {{ at = k;
    document.body.classList.add('present'); show(k); }}));
</script>
</body>
</html>
"""


if __name__ == "__main__":
    import sys

    print(export_html(sys.argv[1], sys.argv[2]))
