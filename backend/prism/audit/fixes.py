"""Исправления находок аудита — те, что пользователь выбрал в панели аудита.

Большинство исправлений детерминированные: геометрия, кегль, цвет, шрифт, обрезка картинки, удаление
заглушки или слайда. Два — через модель: «сократить текст» (скилл text_shortener) и «переписать текст
слайда» (скилл slide_fixer, для смысловых замечаний). Модель меняет только текст существующих фигур:
вёрстка остаётся той, что собрал compose, и после правки текст снова подгоняется под рамку.

    report = await apply_fixes(pptx, [(issue, "shrink"), …], spec, llm=…, skills=…, context=…)
"""

from __future__ import annotations

import asyncio
import json
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from pydantic import BaseModel

from prism.audit.deterministic import (
    MIN_PT,
    Issue,
    backdrop_of,
    ink_box,
    insets,
    run_colors,
    run_fonts,
    run_sizes,
    shape_box,
    text_height,
)
from prism.color import contrast_ratio, delta_e
from prism.layout import fit
from prism.layout import pptx_ops as ops
from prism.llm.structured import StructuredOutputError, ainvoke_structured
from prism.parsing.model import BBox
from prism.parsing.spec import TemplateSpec

MIN_SHRINK = 0.55            # «уменьшить кегль» не опускает текст ниже 55% исходного


@dataclass
class FixReport:
    applied: list[str] = field(default_factory=list)          # id находок
    failed: dict[str, str] = field(default_factory=dict)      # id находки -> почему не вышло
    touched_slides: set[int] = field(default_factory=set)
    deleted_slides: list[int] = field(default_factory=list)


class _Block(BaseModel):
    id: str
    paragraphs: list[str]


class _Blocks(BaseModel):
    blocks: list[_Block]


class _Edit(BaseModel):
    shape_id: int
    paragraphs: list[str]


class _Edits(BaseModel):
    edits: list[_Edit]


# ── помощники ──────────────────────────────────────────────────────────────


def _shape(slide, shape_id: int | None):
    return next((s for s in slide.shapes if s.shape_id == shape_id), None) if shape_id else None


def _paragraphs(shape) -> list[str]:
    return [p.text for p in shape.text_frame.paragraphs if p.text.strip()]


def _scale_steps(spec: TemplateSpec | None, el) -> list[float]:
    """Множители кегля, при которых крупнейший абзац встаёт на ступень шкалы шаблона."""
    sizes = run_sizes(el)
    top = max(sizes) if sizes else 14.0
    steps = sorted({round(s.size_pt, 1) for s in spec.tokens.type_scale}, reverse=True) if spec else []
    low = max(MIN_SHRINK * top, MIN_PT * top / min(sizes)) if sizes else MIN_SHRINK * top   # не мельче 9 pt
    out = [1.0] + [s / top for s in steps if low <= s < top]
    return out if len(out) > 1 else [x for x in (1.0, 0.92, 0.85, 0.78, 0.7, 0.62, MIN_SHRINK) if x * top >= low]


def _refit(shape, spec: TemplateSpec | None) -> bool:
    """Подогнать кегль текста под рамку фигуры; True — влез."""
    el = shape._element
    paras = fit.paragraph_specs(el, [])
    box = shape_box(shape)
    if not paras or box is None:
        return True
    result = fit.fit(paras, box.w, box.h, insets(el), _scale_steps(spec, el))
    if result.scale != 1.0:
        fit.apply_scale(el, paras, result.scale)
    return result.fits


def _set_box(shape, box: BBox) -> None:
    shape.left, shape.top, shape.width, shape.height = box.x, box.y, box.w, box.h


def _nearest(color: str, pool: list[str]) -> str:
    return min(pool, key=lambda c: delta_e(color, c))


# ── детерминированные исправления ──────────────────────────────────────────


def _clamp(shape, w: int, h: int, margins=None) -> None:
    box = shape_box(shape)
    left, top = (margins.left, margins.top) if margins else (0, 0)
    right, bottom = (w - margins.right, h - margins.bottom) if margins else (w, h)
    bw, bh = min(box.w, right - left), min(box.h, bottom - top)
    x = min(max(box.x, left), right - bw)
    y = min(max(box.y, top), bottom - bh)
    _set_box(shape, BBox(x=x, y=y, w=bw, h=bh))


def _move_apart(slide, issue: Issue, shape, spec, h: int) -> None:
    """Наложение: нижний блок уходит под верхний; не хватает места до низа — ужимаем его текст."""
    other_id = int(issue.id.rsplit("-", 1)[-1]) if "-" in issue.id else None
    other = _shape(slide, other_id)
    if other is None:
        raise ValueError("второй блок наложения не найден")
    upper, lower = (other, shape) if shape_box(other).y <= shape_box(shape).y else (shape, other)
    ink_up = ink_box(upper._element, shape_box(upper))
    lb = shape_box(lower)
    gap = h // 60
    new_y = ink_up.bottom + gap
    bottom_limit = h - (spec.tokens.margins.bottom if spec else h // 20)
    if new_y + ink_box(lower._element, lb).h <= bottom_limit:
        shift = new_y - lb.y
        _set_box(lower, BBox(x=lb.x, y=new_y, w=lb.w, h=max(lb.h - max(shift, 0), bottom_limit - new_y)))
    else:
        _refit(upper, spec)
        _refit(lower, spec)


def _crop_to_box(shape) -> None:
    box = shape_box(shape)
    px_w, px_h = shape.image.size
    blip_fill = shape._element.find(f".//{qn('p:blipFill')}")
    blip = blip_fill.find(qn("a:blip"))
    src = blip_fill.find(qn("a:srcRect"))
    if src is None:
        src = etree.Element(qn("a:srcRect"))
        blip.addnext(src)
    src.attrib.clear()
    img_ratio, box_ratio = px_w / px_h, box.w / box.h
    if img_ratio > box_ratio:
        cut = round((1 - box_ratio / img_ratio) / 2 * 100000)
        src.set("l", str(cut))
        src.set("r", str(cut))
    elif img_ratio < box_ratio:
        cut = round((1 - img_ratio / box_ratio) / 2 * 100000)
        src.set("t", str(cut))
        src.set("b", str(cut))


def _recolor(slide, shape, spec, bg_hint: str | None) -> None:
    backdrop = backdrop_of(slide, shape) or bg_hint or "FFFFFF"
    pool = ["FFFFFF", "000000"]
    if spec:
        pool = [p.text for p in spec.tokens.text_pairs] + [c.hex for c in spec.tokens.palette] + pool
    sizes = run_sizes(shape._element)
    need = 3.0 if sizes and min(sizes) >= 18 else 4.5
    ok = [c for c in pool if contrast_ratio(c, backdrop) >= need]
    current = run_colors(shape._element)
    # ближе всего к нынешнему цвету из достаточно контрастных — чтобы не ломать задумку шаблона
    color = min(ok, key=lambda c: delta_e(c, current[0])) if ok and current else (
        max(pool, key=lambda c: contrast_ratio(c, backdrop)))
    ops.set_text_color(shape._element, color)


def _fonts(shape, spec) -> None:
    if spec is None:
        raise ValueError("нет спецификации шаблона")
    t = spec.tokens
    allowed = {f for f, share in t.fonts.usage.items() if share >= 0.01} | {t.fonts.heading, t.fonts.body}
    heading_size = next((s.size_pt for s in t.type_scale if s.role == "heading"), 20)
    for rpr in shape._element.iter(qn("a:rPr")):
        latin = rpr.find(qn("a:latin"))
        if latin is None or latin.get("typeface") in allowed:
            continue
        big = rpr.get("sz") and int(rpr.get("sz")) / 100 >= heading_size
        latin.set("typeface", (t.fonts.heading if big else t.fonts.body) or t.fonts.body or "Arial")


def _snap(shape, spec) -> None:
    if spec is None:
        raise ValueError("нет спецификации шаблона")
    scale = sorted({round(s.size_pt, 1) for s in spec.tokens.type_scale})
    for rpr in [*shape._element.iter(qn("a:rPr")), *shape._element.iter(qn("a:endParaRPr"))]:
        if rpr.get("sz"):
            size = int(rpr.get("sz")) / 100
            below = [s for s in scale if s <= size + 0.05] or scale[:1]
            rpr.set("sz", str(round(below[-1] * 100)))


def _palette(shape, spec) -> None:
    if spec is None:
        raise ValueError("нет спецификации шаблона")
    pool = [c.hex for c in spec.tokens.palette] + [p.text for p in spec.tokens.text_pairs] + ["FFFFFF", "000000"]
    for clr in shape._element.iter(qn("a:srgbClr")):
        val = clr.get("val").upper()
        if all(delta_e(val, c) >= 4 for c in pool):
            clr.set("val", _nearest(val, pool))


def _enlarge(shape) -> None:
    for rpr in [*shape._element.iter(qn("a:rPr")), *shape._element.iter(qn("a:endParaRPr"))]:
        if rpr.get("sz") and int(rpr.get("sz")) < 900:
            rpr.set("sz", "900")


def _legend(shape) -> None:
    chart = shape.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.has_data_labels = True


def _delete_slide(prs, index: int) -> None:
    id_list = prs.slides._sldIdLst
    sld_id = list(id_list)[index - 1]
    prs.part.drop_rel(sld_id.rId)
    id_list.remove(sld_id)


# ── исправления через модель ───────────────────────────────────────────────


def _target_chars(shape) -> tuple[int, int | None]:
    """Сколько символов влезет в рамку при текущем кегле (с запасом), и лимит слов на пункт."""
    text = "\n".join(_paragraphs(shape))
    _, height, limit, _ = text_height(shape._element, shape_box(shape))
    ratio = min(1.0, limit / height) if height else 1.0
    return max(20, int(len(text) * ratio * 0.9)), None


async def _shorten(jobs: list[tuple[Issue, object]], llm, skills, language: str) -> dict[str, list[str]]:
    skill = skills.get("text_shortener")
    model = llm.for_role(skill.role)
    blocks = []
    for issue, shape in jobs:
        max_chars, _ = _target_chars(shape)
        block = {"id": issue.id, "paragraphs": _paragraphs(shape), "max_chars": max_chars}
        if issue.check == "bullet_long":
            block["max_words"] = 15
        if issue.check == "bullets_count":
            block["max_paragraphs"] = 6
        blocks.append(block)
    out = await ainvoke_structured(model, skill.render(language=language, blocks=json.dumps(
        blocks, ensure_ascii=False, indent=1)), _Blocks, retries=1)
    return {b.id: b.paragraphs for b in out.blocks}


async def _rewrite(slide, index: int, problems: list[str], llm, skills, context: dict) -> list[_Edit]:
    skill = skills.get("slide_fixer")
    model = llm.for_role(skill.role)
    shapes = [{"id": s.shape_id, "paragraphs": _paragraphs(s)} for s in slide.shapes
              if s.has_text_frame and s.text_frame.text.strip()]
    msgs = skill.render(language=context.get("language", "ru"), brief=context.get("brief", ""),
                        materials=context.get("materials", []), slide=index,
                        shapes=json.dumps(shapes, ensure_ascii=False, indent=1), problems=problems)
    return (await ainvoke_structured(model, msgs, _Edits, retries=1)).edits


# ── применение ─────────────────────────────────────────────────────────────


async def apply_fixes(pptx: str | Path, choices: list[tuple[Issue, str]], spec: TemplateSpec | None = None,
                      *, llm=None, skills=None, context: dict | None = None, out: str | Path | None = None,
                      bg_hints: list[str | None] | None = None) -> FixReport:
    """choices — (находка, id исправления). Результат сохраняется в out (по умолчанию — поверх pptx).
    bg_hints — фон образца под каждым слайдом (если фон слайда задан темой, а не цветом)."""
    context = context or {}
    prs = Presentation(str(pptx))
    w, h = int(prs.slide_width), int(prs.slide_height)
    slides = list(prs.slides)
    report = FixReport()

    def fail(issue: Issue, why: str) -> None:
        report.failed[issue.id] = why

    # 1. модель: сокращения одним запросом, переписывания — по слайду (параллельно)
    shorten = [(i, _shape(slides[i.slide - 1], i.shape_id)) for i, f in choices if f == "shorten"]
    shorten = [(i, s) for i, s in shorten if s is not None and s.has_text_frame]
    rewrite: dict[int, list[Issue]] = {}
    for issue, f in choices:
        if f == "rewrite":
            rewrite.setdefault(issue.slide, []).append(issue)
    if (shorten or rewrite) and (llm is None or skills is None):
        for i, _ in shorten:
            fail(i, "модель недоступна")
        for issues in rewrite.values():
            for i in issues:
                fail(i, "модель недоступна")
        shorten, rewrite = [], {}
    tasks = []
    if shorten:
        tasks.append(_shorten(shorten, llm, skills, context.get("language", "ru")))
    rewrite_keys = list(rewrite)
    tasks += [_rewrite(slides[n - 1], n, [i.message for i in rewrite[n]], llm, skills, context) for n in rewrite_keys]
    results = await asyncio.gather(*tasks, return_exceptions=True) if tasks else []
    if shorten:
        short = results.pop(0)
        for issue, shape in shorten:
            if isinstance(short, BaseException) or issue.id not in short:
                fail(issue, f"модель не сократила текст: {short}" if isinstance(short, BaseException)
                     else "модель не вернула этот блок")
                continue
            ops.set_paragraphs(shape._element, short[issue.id])
            _refit(shape, spec)
            report.applied.append(issue.id)
            report.touched_slides.add(issue.slide)
    for n, edits in zip(rewrite_keys, results, strict=True):
        if isinstance(edits, BaseException):
            for i in rewrite[n]:
                fail(i, f"модель не переписала слайд: {edits}")
            continue
        for e in edits:
            shape = _shape(slides[n - 1], e.shape_id)
            if shape is not None and shape.has_text_frame and ops.set_paragraphs(shape._element, e.paragraphs):
                _refit(shape, spec)
        report.applied += [i.id for i in rewrite[n]]
        report.touched_slides.add(n)

    # 2. детерминированные исправления
    deletions = set()
    for issue, fix_id in choices:
        if fix_id in ("shorten", "rewrite"):
            continue
        if fix_id == "delete_slide":
            deletions.add(issue.slide)
            report.applied.append(issue.id)
            continue
        slide = slides[issue.slide - 1] if 1 <= issue.slide <= len(slides) else None
        shape = _shape(slide, issue.shape_id) if slide is not None else None
        try:
            if fix_id == "font" and shape is None and slide is not None:
                for s in slide.shapes:                      # font_count — вся гарнитура слайда
                    if s.has_text_frame and run_fonts(s._element):
                        _fonts(s, spec)
            elif shape is None:
                raise ValueError("фигура не найдена — возможно, уже исправлена")
            elif fix_id == "shrink":
                if not _refit(shape, spec):
                    raise ValueError("кегль уменьшен до нижней ступени шкалы, но текст всё ещё не влезает — "
                                     "выберите «Сократить текст»")
            elif fix_id == "clamp":
                _clamp(shape, w, h)
            elif fix_id == "clamp_margins":
                _clamp(shape, w, h, spec.tokens.margins if spec else None)
            elif fix_id == "move":
                _move_apart(slide, issue, shape, spec, h)
            elif fix_id == "aspect":
                _crop_to_box(shape)
            elif fix_id == "recolor":
                hint = bg_hints[issue.slide - 1] if bg_hints and issue.slide <= len(bg_hints) else None
                _recolor(slide, shape, spec, hint)
            elif fix_id == "font":
                _fonts(shape, spec)
            elif fix_id == "snap":
                _snap(shape, spec)
            elif fix_id == "palette":
                _palette(shape, spec)
            elif fix_id == "enlarge":
                _enlarge(shape)
            elif fix_id == "delete":
                ops.remove_shape(slide, shape.shape_id)
            elif fix_id == "legend":
                _legend(shape)
            else:
                raise ValueError(f"неизвестное исправление {fix_id}")
        except (ValueError, AttributeError, KeyError, StructuredOutputError) as e:
            fail(issue, str(e))
            continue
        report.applied.append(issue.id)
        report.touched_slides.add(issue.slide)
    for n in sorted(deletions, reverse=True):          # с конца — номера оставшихся не сдвигаются
        _delete_slide(prs, n)
    report.deleted_slides = sorted(deletions)
    prs.save(str(out or pptx))
    return report

