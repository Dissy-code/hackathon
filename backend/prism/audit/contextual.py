"""Контекстуальный (смысловой) аудит — вопросы «да/нет» из Приложения 1 ТЗ, отвечает модель.

    по слайду  — скилл slide_auditor (роль audit_vlm): картинка слайда + его текст + бриф и материалы;
                 вопросы 1–8 и 10: вывод в заголовке, соответствие заголовку, одна мысль, цифры из
                 источников, есть содержание, картинки по теме, нет мусора, нет опечаток, данные по делу;
    по колоде  — скилл deck_auditor: связаны ли соседние слайды (вопрос 11).
    Вопрос 9 (вся колода на одном языке) проверяется детерминированно — см. deterministic.language.

Недетерминированная проверка: на повторном запуске ответ может отличаться, поэтому в интерфейсе она
помечена «смысл · модель» и запускается отдельно от детерминированной.
"""

from __future__ import annotations

import asyncio
import base64
from pathlib import Path

from langchain_core.messages import HumanMessage
from pptx import Presentation
from pydantic import BaseModel, Field

from prism.audit.deterministic import Fix, Issue
from prism.llm.structured import StructuredOutputError, ainvoke_structured

PARALLEL = 6

# вопрос -> (текст для интерфейса, исправления)
QUESTIONS = {
    "conclusion_title": ("заголовок называет тему, а не вывод", [("rewrite", "Переписать заголовок (модель)")]),
    "content_matches_title": ("содержимое не соответствует заголовку", [("rewrite", "Переписать текст (модель)")]),
    "one_sentence": ("на слайде больше одной мысли", [("rewrite", "Сфокусировать текст (модель)")]),
    "facts_sourced": ("цифры или факты не из брифа и материалов", [("rewrite", "Убрать непроверенное (модель)")]),
    "has_content": ("на слайде нет содержания, кроме заголовка", [("delete_slide", "Удалить слайд"),
                                                                  ("rewrite", "Дописать текст (модель)")]),
    "visuals_relevant": ("картинки или иконки не по теме", []),
    "no_garbage": ("служебный мусор: реплики, куски промпта", [("rewrite", "Убрать мусор (модель)")]),
    "no_typos": ("опечатки в тексте", [("rewrite", "Исправить опечатки (модель)")]),
    "data_serves_point": ("строки таблицы или легенды не работают на мысль", [("rewrite", "Переписать (модель)")]),
}
SEVERITY = {"facts_sourced": "error", "no_garbage": "error", "has_content": "error", "content_matches_title": "error"}


class Answer(BaseModel):
    ok: bool
    reason: str = ""


class SlideVerdict(BaseModel):
    conclusion_title: Answer
    content_matches_title: Answer
    one_sentence: Answer
    facts_sourced: Answer
    has_content: Answer
    visuals_relevant: Answer
    no_garbage: Answer
    no_typos: Answer
    data_serves_point: Answer


class Break(BaseModel):
    after: int
    reason: str


class DeckVerdict(BaseModel):
    breaks: list[Break] = Field(default_factory=list)


def slide_texts(pptx: str | Path) -> list[str]:
    prs = Presentation(str(pptx))
    return ["\n".join(s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())
            for slide in prs.slides]


def _issue(check: str, slide: int, message: str) -> Issue:
    title, fixes = QUESTIONS.get(check, (check, []))
    return Issue(id=f"ctx:{check}:{slide}", slide=slide, shape_id=None, check=check, group="content",
                 mode="contextual", severity=SEVERITY.get(check, "warning"),
                 message=_sentence(message) if message else _sentence(title),
                 fixes=[Fix(id=i, label=lbl) for i, lbl in fixes])


def _sentence(text: str) -> str:
    text = text.strip()
    return text[:1].upper() + text[1:] if text else text


async def audit_slides(pngs: list[Path], texts: list[str], kinds: list[str], context: dict, llm, skills,
                       only: set[int] | None = None) -> tuple[list[Issue], list[str]]:
    """Находки по слайдам; only — номера слайдов (с 1), которые проверить (после правки — только их)."""
    skill = skills.get("slide_auditor")
    model = llm.for_role(skill.role)
    gate = asyncio.Semaphore(PARALLEL)
    errors: list[str] = []

    async def one(n: int) -> list[Issue]:
        png = pngs[n - 1]
        msgs = skill.render(language=context.get("language", "ru"), brief=context.get("brief", ""),
                            materials=context.get("materials", []), slide=n, total=len(pngs),
                            template=context.get("template_name", "корпоративный шаблон"),
                            kind=kinds[n - 1] if n <= len(kinds) else "slide", text=texts[n - 1] or "(нет текста)")
        image = base64.b64encode(png.read_bytes()).decode()
        msgs[-1] = HumanMessage(content=[
            {"type": "text", "text": msgs[-1].content},
            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image}"}},
        ])
        async with gate:
            try:
                verdict = await ainvoke_structured(model, msgs, SlideVerdict, retries=1)
            except StructuredOutputError as e:
                errors.append(f"слайд {n}: модель не ответила ({e})")
                return []
        return [_issue(q, n, a.reason.strip()) for q, a in verdict if not a.ok]

    numbers = [n for n in range(1, len(pngs) + 1) if only is None or n in only]
    found = await asyncio.gather(*(one(n) for n in numbers))
    return [i for batch in found for i in batch], errors


async def audit_storyline(texts: list[str], kinds: list[str], context: dict, llm, skills) -> list[Issue]:
    skill = skills.get("deck_auditor")
    model = llm.for_role(skill.role)
    slides = [{"n": i, "kind": kinds[i - 1] if i <= len(kinds) else "slide", "text": t.replace("\n", " / ")[:400]}
              for i, t in enumerate(texts, 1)]
    try:
        verdict = await ainvoke_structured(model, skill.render(language=context.get("language", "ru"), slides=slides),
                                           DeckVerdict, retries=1)
    except StructuredOutputError:
        return []
    out = []
    for b in verdict.breaks:
        if 1 <= b.after < len(texts):
            out.append(Issue(id=f"ctx:storyline:{b.after + 1}", slide=b.after + 1, shape_id=None, check="storyline",
                             group="content", mode="contextual", severity="warning",
                             message=f"не связан с предыдущим слайдом: {b.reason}",
                             fixes=[Fix(id="rewrite", label="Добавить связку (модель)")]))
    return out
