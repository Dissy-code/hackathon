"""Детерминированный аудит готовой колоды: то, что проверяется по файлу — координатам, размерам, кодам
цветов, шрифтам, связям с макетами. Один и тот же файл всегда даёт один и тот же список находок.

Проверки (ориентир — Приложение 1 ТЗ; код проверки -> группа):

    вёрстка     off_slide, overlap, overflow (в т.ч. текст срезан краем), margins, stretched_image
    шаблон      font_foreign, font_count, size_off_scale, small_font, color_off_palette, contrast, layout_foreign
    плотность   bullets_count, bullet_long, table_size, chart_series, fill_ratio
    целостность file_open, sample_text, empty_slide, raster_slide, chart_labels, duplicate_slide
    смысл       language (детерминированная часть: язык слайда не совпадает с языком колоды)

Проверки шаблона работают, если передана спецификация шаблона (TemplateSpec): без неё неизвестно, какие
шрифты, кегли и цвета «свои». Каждая находка — Issue: номер слайда, id фигуры, рамка для подсветки в
интерфейсе и варианты исправления (prism/audit/fixes.py), из которых пользователь выбирает.
"""

from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pydantic import BaseModel, Field

from prism.color import contrast_ratio, delta_e
from prism.layout import fit
from prism.parsing.model import BBox
from prism.parsing.spec import TemplateSpec

_SAMPLE = re.compile(r"^\s*(\[[^\]]+\]|заголовок( слайда)?|подзаголовок|текст( карточки)?|lorem ipsum.*|xxx+|todo|"
                     r"title|subtitle|click to (add|edit).*|вставьте текст.*|вставить (фото|текст).*|"
                     r"имя фамилия|должность)\s*$", re.IGNORECASE)
MIN_PT = 9.0
MAX_BULLETS = 6
MAX_BULLET_WORDS = 15
MAX_TABLE_ROWS, MAX_TABLE_COLS = 7, 5
MAX_SERIES = 5
FILL_MIN, FILL_MAX = 0.25, 0.75
_SPARSE_KINDS = {"title", "section", "closing", "quote"}      # у них пустота — замысел, а не дефект

GROUPS = {
    "off_slide": "layout", "overlap": "layout", "overflow": "layout", "margins": "layout",
    "stretched_image": "layout",
    "font_foreign": "template", "font_count": "template", "size_off_scale": "template", "small_font": "template",
    "color_off_palette": "template", "contrast": "template", "layout_foreign": "template",
    "bullets_count": "density", "bullet_long": "density", "table_size": "density", "chart_series": "density",
    "fill_ratio": "density",
    "file_open": "integrity", "sample_text": "integrity", "empty_slide": "integrity", "raster_slide": "integrity",
    "chart_labels": "integrity", "duplicate_slide": "integrity",
    "language": "content",
}

# варианты исправления по проверке: (id, подпись). Первый — предлагаемый по умолчанию.
FIXES: dict[str, list[tuple[str, str]]] = {
    "overflow": [("shrink", "Уменьшить кегль под рамку"), ("shorten", "Сократить текст (модель)")],
    "off_slide": [("clamp", "Вернуть в границы слайда")],
    "margins": [("clamp_margins", "Вернуть внутрь полей")],
    "overlap": [("move", "Сдвинуть нижний блок"), ("shrink", "Ужать текст в свою рамку")],
    "stretched_image": [("aspect", "Обрезать без искажения")],
    "contrast": [("recolor", "Цвет текста с достаточным контрастом")],
    "font_foreign": [("font", "Шрифт шаблона")],
    "font_count": [("font", "Оставить шрифты шаблона")],
    "size_off_scale": [("snap", "Кегль по шкале шаблона")],
    "color_off_palette": [("palette", "Ближайший цвет палитры")],
    "small_font": [("enlarge", "Увеличить до 9 pt")],
    "bullets_count": [("shorten", "Сократить список (модель)")],
    "bullet_long": [("shorten", "Сократить пункты (модель)")],
    "sample_text": [("delete", "Удалить заглушку")],
    "empty_slide": [("delete_slide", "Удалить слайд")],
    "raster_slide": [("delete_slide", "Удалить слайд")],
    "duplicate_slide": [("delete_slide", "Удалить повтор")],
    "chart_labels": [("legend", "Добавить легенду и подписи")],
    "language": [("rewrite", "Переписать на языке колоды (модель)")],
}


class Fix(BaseModel):
    id: str
    label: str


class Issue(BaseModel):
    id: str
    slide: int                        # с 1
    shape_id: int | None
    check: str                        # код проверки (см. GROUPS)
    group: str                        # layout | template | density | integrity | content
    mode: str = "deterministic"       # deterministic | contextual
    severity: str                     # error | warning
    message: str
    box: list[float] | None = None    # x, y, w, h в долях слайда — что подсветить
    fixes: list[Fix] = Field(default_factory=list)


def make_issue(check: str, slide: int, shape_id: int | None, severity: str, message: str, box: BBox | None,
               size: tuple[int, int], extra_id: str = "") -> Issue:
    w, h = size
    frac = [round(box.x / w, 4), round(box.y / h, 4), round(box.w / w, 4), round(box.h / h, 4)] if box else None
    return Issue(id=f"{check}:{slide}:{shape_id or 0}{extra_id}", slide=slide, shape_id=shape_id, check=check,
                 group=GROUPS[check], severity=severity, message=message, box=frac,
                 fixes=[Fix(id=i, label=lbl) for i, lbl in FIXES.get(check, [])])


# ── чтение фигур ───────────────────────────────────────────────────────────


def shape_box(shape) -> BBox | None:
    if shape.left is None or shape.width is None:
        return None
    return BBox(x=int(shape.left), y=int(shape.top), w=int(shape.width), h=int(shape.height))


def _inter(a: BBox, b: BBox) -> int:
    w = min(a.right, b.right) - max(a.x, b.x)
    h = min(a.bottom, b.bottom) - max(a.y, b.y)
    return max(w, 0) * max(h, 0)


def _solid(parent: etree._Element | None) -> str | None:
    if parent is None:
        return None
    clr = parent.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
    return clr.get("val").upper() if clr is not None else None


def slide_background(slide) -> str | None:
    for part in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        color = _solid(part._element.find(f"{qn('p:cSld')}/{qn('p:bg')}/{qn('p:bgPr')}"))
        if color:
            return color
    return None


def backdrop_of(slide, shape) -> str | None:
    """Цвет под текстом: ближайшая залитая фигура ниже по z, накрывающая текст, иначе фон слайда."""
    box = shape_box(shape)
    best = None
    for other in slide.shapes:
        if other.shape_id == shape.shape_id:
            break
        ob = shape_box(other)
        color = _solid(other._element.find(qn("p:spPr")))
        if color and ob is not None and box is not None and _inter(ob, box) >= 0.6 * box.area \
                and (best is None or ob.area < best[0].area):
            best = (ob, color)
    return best[1] if best else slide_background(slide)


def run_colors(el: etree._Element) -> list[str]:
    return [c.get("val").upper() for c in el.iter(qn("a:srgbClr"))
            if c.getparent().tag == qn("a:solidFill") and c.getparent().getparent().tag == qn("a:rPr")]


def run_fonts(el: etree._Element) -> set[str]:
    out = set()
    for rpr in el.iter(qn("a:rPr")):
        latin = rpr.find(qn("a:latin"))
        if latin is not None and latin.get("typeface") and not latin.get("typeface").startswith("+"):
            out.add(latin.get("typeface"))
    return out


def run_sizes(el: etree._Element) -> list[float]:
    return [int(r.get("sz")) / 100 for r in el.iter(qn("a:rPr")) if r.get("sz")]


def _auto_grows(el: etree._Element) -> bool:
    body = el.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    return body is not None and body.find(qn("a:spAutoFit")) is not None


def _explicit_sizes(el: etree._Element) -> bool:
    runs = list(el.iter(qn("a:r")))
    return bool(runs) and all(r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).get("sz") for r in runs)


def insets(el: etree._Element) -> tuple[int, int, int, int]:
    body = el.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    d = fit.DEFAULT_INSETS
    if body is None:
        return d
    return tuple(int(body.get(k, dv)) for k, dv in zip(("lIns", "tIns", "rIns", "bIns"), d, strict=True))


def _is_title(shape) -> bool:
    return bool(shape.is_placeholder and "TITLE" in str(shape.placeholder_format.type))


def _paragraphs(shape) -> list[str]:
    return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]


def text_height(el: etree._Element, box: BBox) -> tuple[int, float, float, bool]:
    """(строк, высота текста pt, высота рамки pt, есть слово шире строки) — по метрикам шрифтов."""
    paras = fit.paragraph_specs(el, [])
    l_, t_, r_, b_ = insets(el)
    lines, height, too_wide = fit.measure(paras, max((box.w - l_ - r_) / fit.EMU_PER_PT, 1))
    limit = (box.h - t_ - b_) / fit.EMU_PER_PT
    if lines == 1 and not too_wide and paras:
        limit = max(limit, max(p.size_pt for p in paras) * fit.LINE_SPACING * 1.05)
    return lines, height, limit, too_wide


def ink_box(el: etree._Element, box: BBox) -> BBox:
    """Где на самом деле лежит текст: рамки шаблона часто выше и шире текста (под длинный образец),
    и наложение рамок ещё не значит, что наложились буквы."""
    paras = fit.paragraph_specs(el, [])
    if not paras:
        return box
    l_, t_, r_, b_ = insets(el)
    width_pt = max((box.w - l_ - r_) / fit.EMU_PER_PT, 1)
    _, height, _ = fit.measure(paras, width_pt)
    h = int(height * fit.EMU_PER_PT) + t_ + b_
    w = min(int(fit.widest_line_pt(paras, width_pt, 1.0) * fit.EMU_PER_PT) + l_ + r_, box.w)
    body = el.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    anchor = body.get("anchor", "t") if body is not None else "t"
    y = box.y if anchor == "t" else (box.bottom - h if anchor == "b" else box.y + (box.h - h) // 2)
    p = el.find(f"{qn('p:txBody')}/{qn('a:p')}/{qn('a:pPr')}")
    algn = p.get("algn", "l") if p is not None else "l"
    x = box.x if algn in ("l", "just") else (box.right - w if algn == "r" else box.x + (box.w - w) // 2)
    return BBox(x=x, y=y, w=w, h=max(h, 1))


# ── правила шаблона ────────────────────────────────────────────────────────


class Rules:
    """Что считается «своим» для шаблона: шрифты, кегли шкалы, цвета палитры, поля, макеты."""

    def __init__(self, spec: TemplateSpec | None, layouts: set[str] | None):
        self.spec = spec
        self.layouts = layouts
        self.fonts: set[str] | None = None
        self.sizes: list[float] | None = None
        self.colors: set[str] | None = None
        self.margins = None
        if spec is None:
            return
        t = spec.tokens
        self.fonts = {f for f, share in t.fonts.usage.items() if share >= 0.01} | {t.fonts.heading, t.fonts.body}
        self.fonts.discard(None)
        self.sizes = sorted({round(x, 1) for s in t.type_scale for x in [s.size_pt, *s.sizes]})
        self.colors = ({c.hex.upper() for c in t.palette} | {m.upper() for c in t.palette for m in c.members}
                       | {p.text.upper() for p in t.text_pairs} | {p.backdrop.upper() for p in t.text_pairs}
                       | {b.color.upper() for b in t.backgrounds if b.color} | {"FFFFFF", "000000"})
        self.margins = t.margins

    def size_ok(self, size: float) -> bool:
        return self.sizes is None or any(abs(size - s) <= 0.6 for s in self.sizes)

    def color_ok(self, color: str) -> bool:
        return self.colors is None or any(delta_e(color, c) < 4 for c in self.colors)


# ── проверки слайда ────────────────────────────────────────────────────────


def _is_picture(shape) -> bool:
    return shape.shape_type is not None and "PICTURE" in str(shape.shape_type)


def audit_slide(slide, index: int, size: tuple[int, int], rules: Rules, kind: str | None,
                bg_hint: str | None = None) -> list[Issue]:
    """bg_hint — цвет фона образца шаблона, если фон слайда задан темой/картинкой, а не цветом в XML."""
    slide_w, slide_h = size
    issues: list[Issue] = []
    bg = slide_background(slide) or bg_hint
    texts = []                   # (shape, bbox, el, z)
    fills = []                   # (bbox, цвет, z) — залитые фигуры как фон под текстом
    content_boxes = []           # то, что «занимает» слайд, — для плотности заполнения
    pictures, charts, tables = 0, 0, 0
    tol_x, tol_y = slide_w // 100, slide_h // 100

    def add(check, shape, severity, message, box=None, extra=""):
        issues.append(make_issue(check, index, shape.shape_id if shape is not None else None, severity, message,
                                 box, size, extra))

    if rules.layouts is not None and slide.slide_layout.name not in rules.layouts:
        add("layout_foreign", None, "warning", f"слайд на макете «{slide.slide_layout.name}», которого нет в шаблоне")

    for z, shape in enumerate(slide.shapes):
        box = shape_box(shape)
        if box is None:
            continue
        el = shape._element
        ours = shape.name.startswith("prism:")                 # нарисовано вёрсткой, а не взято из шаблона
        big = box.area > 0.6 * slide_w * slide_h
        fill = _solid(el.find(qn("p:spPr")))
        if fill:
            fills.append((box, fill, z))
            if ours and not rules.color_ok(fill):
                add("color_off_palette", shape, "warning", f"цвет подложки #{fill} не из палитры шаблона", box)
        if _is_picture(shape):
            pictures += 1
            if not big:
                content_boxes.append(box)
            _check_aspect(shape, box, add)
        if getattr(shape, "has_chart", False) and shape.has_chart:
            charts += 1
            content_boxes.append(box)
            _check_chart(shape, box, add)
        if getattr(shape, "has_table", False) and shape.has_table:
            tables += 1
            content_boxes.append(box)
            rows, cols = len(shape.table.rows), len(shape.table.columns)
            if rows - 1 > MAX_TABLE_ROWS or cols > MAX_TABLE_COLS:
                add("table_size", shape, "warning",
                    f"таблица {rows - 1}×{cols}: больше {MAX_TABLE_ROWS} строк или {MAX_TABLE_COLS} колонок", box)
        if ours and rules.margins is not None:
            _check_margins(shape, box, size, rules, add)
        if fill and not big and not shape.has_text_frame:
            content_boxes.append(box)
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        texts.append((shape, box, el, z))
        content_boxes.append(box)
        text = shape.text_frame.text.strip()
        if _SAMPLE.match(text):
            add("sample_text", shape, "error", f"остался текст-образец: «{text[:40]}»", box)
        if box.x < -tol_x or box.y < -tol_y or box.right > slide_w + tol_x or box.bottom > slide_h + tol_y:
            add("off_slide", shape, "error", f"текстовый блок выходит за край слайда: «{text[:40]}»", box)
        sizes = run_sizes(el)
        if sizes and min(sizes) < MIN_PT - 0.01:
            add("small_font", shape, "warning", f"кегль {min(sizes):g} pt — мельче {MIN_PT:g} pt", box)
        off = sorted({s for s in sizes if not rules.size_ok(s)})
        if off:
            add("size_off_scale", shape, "warning",
                f"кегль {', '.join(f'{s:g}' for s in off[:3])} pt не из шкалы шаблона", box)
        foreign = sorted(f for f in run_fonts(el) if rules.fonts is not None and f not in rules.fonts)
        if foreign:
            add("font_foreign", shape, "warning", f"шрифт {', '.join(foreign)} не из шаблона", box)
        colors = sorted(c for c in set(run_colors(el)) if not rules.color_ok(c))
        if colors:
            add("color_off_palette", shape, "warning", f"цвет текста #{colors[0]} не из палитры шаблона", box)
        paras = _paragraphs(shape)
        if not _is_title(shape):
            if len(paras) > MAX_BULLETS:
                add("bullets_count", shape, "warning", f"{len(paras)} пунктов в одном блоке — больше {MAX_BULLETS}", box)
            long = [p for p in paras if len(p.split()) > MAX_BULLET_WORDS]
            if long:
                add("bullet_long", shape, "warning",
                    f"пункт из {len(long[0].split())} слов — длиннее {MAX_BULLET_WORDS}: «{long[0][:40]}»", box)
        if _explicit_sizes(el) and not _auto_grows(el):
            _check_overflow(shape, box, el, size, add)

    _refine_overflow(issues, texts, size)
    _check_contrast(texts, fills, bg, add)
    _check_overlap(texts, add)

    fonts = set().union(*(run_fonts(el) for _, _, el, _ in texts)) if texts else set()
    if len(fonts) > 2:
        add("font_count", None, "warning", f"на слайде {len(fonts)} гарнитуры: {', '.join(sorted(fonts))}")

    words = sum(len(s.text_frame.text.split()) for s, *_ in texts)
    sparse_ok = kind in _SPARSE_KINDS and texts       # разделитель из одного крупного тезиса — норма
    if not charts and not tables and len(texts) <= 1 and not sparse_ok and not (pictures and kind == "image"):
        add("empty_slide", None, "error", "на слайде только заголовок" if texts else "пустой слайд")
    if pictures == 1 and not texts and not charts and not tables:
        pic = next(s for s in slide.shapes if _is_picture(s))
        if (pb := shape_box(pic)) is not None and pb.area > 0.85 * slide_w * slide_h:
            add("raster_slide", None, "error", "слайд — одна картинка, а не редактируемые объекты")
    if kind not in _SPARSE_KINDS and (texts or charts or tables):
        ratio = _fill_ratio(content_boxes, size)
        if ratio < FILL_MIN and words > 0:
            add("fill_ratio", None, "warning", f"слайд заполнен на {ratio:.0%} — меньше четверти")
        elif ratio > FILL_MAX:
            add("fill_ratio", None, "warning", f"слайд заполнен на {ratio:.0%} — больше трёх четвертей")
    return issues


def _check_aspect(shape, box: BBox, add) -> None:
    """Картинка растянута: пропорции рамки не совпадают с пропорциями видимой части изображения."""
    try:
        px_w, px_h = shape.image.size
    except (AttributeError, ValueError, KeyError):
        return
    crop = shape._element.find(f".//{qn('a:srcRect')}")
    cut = [int(crop.get(k, 0)) / 100000 for k in ("l", "t", "r", "b")] if crop is not None else [0, 0, 0, 0]
    vis_w, vis_h = px_w * (1 - cut[0] - cut[2]), px_h * (1 - cut[1] - cut[3])
    if vis_w <= 0 or vis_h <= 0 or box.h <= 0:
        return
    ratio = (box.w / box.h) / (vis_w / vis_h)
    if abs(ratio - 1) > 0.04:
        add("stretched_image", shape, "error", f"картинка растянута на {abs(ratio - 1):.0%}", box)


def _check_chart(shape, box: BBox, add) -> None:
    chart = shape.chart
    series = list(chart.plots[0].series) if chart.plots else []
    if len(series) > MAX_SERIES:
        add("chart_series", shape, "warning", f"на диаграмме {len(series)} серий — больше {MAX_SERIES}", box)
    missing = []
    if len(series) > 1 and not chart.has_legend:
        missing.append("легенды")
    round_chart = "PIE" in str(chart.chart_type) or "DOUGHNUT" in str(chart.chart_type)
    if round_chart and not chart.plots[0].has_data_labels and not chart.has_legend:
        missing.append("подписей долей")
    if not round_chart:
        try:
            has_unit = chart.value_axis.has_title
        except ValueError:
            has_unit = True
        if not has_unit and not chart.plots[0].has_data_labels:
            missing.append("единиц на оси значений")
    if missing:
        add("chart_labels", shape, "warning", "у диаграммы нет " + " и ".join(missing), box)


def _check_margins(shape, box: BBox, size: tuple[int, int], rules: Rules, add) -> None:
    w, h = size
    m = rules.margins
    tol = w // 200
    if box.x < m.left - tol or box.right > w - m.right + tol or box.y < m.top - tol or box.bottom > h - m.bottom + tol:
        add("margins", shape, "warning", "блок заходит в поля шаблона", box)


def _check_overflow(shape, box: BBox, el, size: tuple[int, int], add) -> None:
    _, height, limit, too_wide = text_height(el, box)
    if height <= limit * 1.1 + 2 and not too_wide:
        return
    _, t_, _, b_ = insets(el)
    grown = BBox(x=box.x, y=box.y, w=box.w, h=int(height * fit.EMU_PER_PT) + t_ + b_)
    cut = grown.bottom > size[1]
    why = "слово шире рамки" if too_wide else f"{height:.0f} из {limit:.0f} pt"
    text = shape.text_frame.text.strip()
    msg = f"текст {'срезан краем слайда' if cut else 'не помещается в рамку'} ({why}): «{text[:40]}»"
    add("overflow", shape, "error" if cut or too_wide else "warning", msg, grown)


def _refine_overflow(issues: list[Issue], texts, size: tuple[int, int]) -> None:
    """Переполнение, что налезает на соседний текст, — ошибка, а не предупреждение."""
    boxes = {s.shape_id: b for s, b, _, _ in texts}
    for issue in issues:
        if issue.check != "overflow" or issue.severity == "error" or issue.shape_id not in boxes or not issue.box:
            continue
        own = boxes[issue.shape_id]
        grown = BBox(x=own.x, y=own.y, w=own.w, h=int(issue.box[3] * size[1]))
        if any(sid != issue.shape_id and _inter(grown, b) > 0 and _inter(own, b) == 0 for sid, b in boxes.items()):
            issue.severity = "error"
            issue.message = issue.message.replace("не помещается в рамку", "не помещается и налезает на соседний блок")


def _check_contrast(texts, fills, bg: str | None, add) -> None:
    for shape, box, el, z in texts:
        under = [(b, c) for b, c, fz in fills if fz < z and _inter(b, box) >= 0.6 * box.area]
        backdrop = min(under, key=lambda bc: bc[0].area)[1] if under else bg
        colors = set(run_colors(el))
        if not backdrop or not colors:
            continue
        sizes = run_sizes(el)
        bold = any(r.get("b") in ("1", "true") for r in el.iter(qn("a:rPr")))
        large = bool(sizes) and (min(sizes) >= 18 or (bold and min(sizes) >= 14))
        need = 3.0 if large else 4.5          # WCAG: крупный текст — 3:1, обычный — 4.5:1
        worst = min(colors, key=lambda c: contrast_ratio(c, backdrop))
        ratio = contrast_ratio(worst, backdrop)
        if ratio < need:
            add("contrast", shape, "error" if ratio < 3 else "warning",
                f"контраст #{worst} на #{backdrop}: {ratio:.1f}:1, нужно ≥ {need:g}:1", box)


def _check_overlap(texts, add) -> None:
    inks = [ink_box(el, box) for _, box, el, _ in texts]
    for i, (sa, ba, _, _) in enumerate(texts):
        for j, (sb, bb, _, _) in enumerate(texts[i + 1:], i + 1):
            ia, ib = inks[i], inks[j]
            if _inter(ia, ib) > 0.15 * min(ia.area, ib.area):
                lower = sb if bb.y >= ba.y else sa
                add("overlap", lower, "warning", f"текстовые блоки перекрываются: «{sa.text_frame.text[:25]}» "
                    f"и «{sb.text_frame.text[:25]}»", _union(ba, bb), extra=f"-{sa.shape_id}")


def _union(a: BBox, b: BBox) -> BBox:
    x, y = min(a.x, b.x), min(a.y, b.y)
    return BBox(x=x, y=y, w=max(a.right, b.right) - x, h=max(a.bottom, b.bottom) - y)


def _fill_ratio(boxes: list[BBox], size: tuple[int, int], n: int = 32) -> float:
    """Доля слайда под содержимым: сетка n×n, клетка занята, если её центр внутри какого-то блока."""
    w, h = size
    covered = 0
    for gy in range(n):
        cy = (gy + 0.5) * h / n
        for gx in range(n):
            cx = (gx + 0.5) * w / n
            if any(b.x <= cx <= b.right and b.y <= cy <= b.bottom for b in boxes):
                covered += 1
    return covered / (n * n)


# ── проверки колоды ────────────────────────────────────────────────────────


def _words(text: str) -> set[str]:
    return set(re.findall(r"\w{4,}", text.lower()))


def _deck_checks(prs, size) -> list[Issue]:
    issues = []
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
    words = [_words(t) for t in texts]
    for i in range(len(words)):
        for j in range(i + 1, len(words)):
            a, b = words[i], words[j]
            if len(a) >= 5 and len(b) >= 5 and len(a & b) / len(a | b) > 0.8:
                issues.append(make_issue("duplicate_slide", j + 1, None, "warning",
                                         f"слайд почти повторяет слайд {i + 1}", None, size))
    # язык: слайд, где больше трети слов на другом алфавите, выбивается из колоды
    cyr = sum(len(re.findall(r"[а-яё]", t, re.IGNORECASE)) for t in texts)
    lat = sum(len(re.findall(r"[a-z]", t, re.IGNORECASE)) for t in texts)
    deck_ru = cyr >= lat
    for i, t in enumerate(texts, 1):
        # имена собственные (Nestlé, KitKat) — с заглавной буквы, их не считаем
        foreign = re.findall(r"\b[a-z]{3,}\b", t) if deck_ru else re.findall(r"\b[а-яё]{3,}\b", t)
        total = len(re.findall(r"\b\w{3,}\b", t))
        if total >= 6 and len(foreign) / total > 0.33:
            issues.append(make_issue("language", i, None, "warning",
                                     f"слайд не на языке колоды ({'русский' if deck_ru else 'английский'})",
                                     None, size))
    return issues


def pattern_bg(spec: TemplateSpec | None, patterns: list[str] | None, index: int) -> str | None:
    """«slide:7» или «synth:kpi<-slide:3» -> фон этого образца по спецификации шаблона."""
    if spec is None or not patterns or index > len(patterns):
        return None
    pid = patterns[index - 1].split("<-")[-1]
    return next((ps.pattern.background for ps in spec.patterns if ps.pattern.id == pid), None)


def template_layouts(template_pptx: str | Path | None) -> set[str] | None:
    if template_pptx is None or not Path(template_pptx).exists():
        return None
    prs = Presentation(str(template_pptx))
    return {layout.name for master in prs.slide_masters for layout in master.slide_layouts}


def audit_deck(pptx: str | Path, spec: TemplateSpec | None = None, template_pptx: str | Path | None = None,
               kinds: list[str] | None = None, patterns: list[str] | None = None) -> list[Issue]:
    """kinds — тип каждого слайда из отчёта вёрстки (для плотности: разделителю пустота положена);
    patterns — образец шаблона под каждым слайдом (оттуда цвет фона для проверки контраста)."""
    try:
        prs = Presentation(str(pptx))
    except Exception as e:  # noqa: BLE001 — любой сбой открытия и есть находка
        return [Issue(id="file_open:0:0", slide=0, shape_id=None, check="file_open", group="integrity",
                      severity="error", message=f"файл не открывается: {e}")]
    size = (int(prs.slide_width), int(prs.slide_height))
    rules = Rules(spec, template_layouts(template_pptx))
    issues = []
    for i, slide in enumerate(prs.slides, 1):
        kind = kinds[i - 1] if kinds and i <= len(kinds) else None
        issues += audit_slide(slide, i, size, rules, kind, pattern_bg(spec, patterns, i))
    return issues + _deck_checks(prs, size)


if __name__ == "__main__":
    import sys

    for issue in audit_deck(sys.argv[1]):
        print(f"слайд {issue.slide:>2}  {issue.severity:<7} {issue.check:<17} #{issue.shape_id}  {issue.message}")
