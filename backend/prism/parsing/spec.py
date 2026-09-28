"""Полный разбор шаблона: извлечение -> классификация -> токены -> паттерны -> ассеты = TemplateSpec.

    python -m prism.parsing.spec data/templates/x.pptx            # только эвристики, без модели
    python -m prism.parsing.spec data/templates/x.pptx --llm      # + скилл template_classifier

Результат кешируется по sha256 файла в data/specs/<sha>/ (spec.json + assets/), повторный разбор мгновенный.
"""

from __future__ import annotations

import argparse
import asyncio
from pathlib import Path

from pptx import Presentation
from pydantic import BaseModel, Field

from prism.config import BACKEND_DIR
from prism.parsing.classify import (
    SlideClass,
    TemplateClassification,
    Usage,
    classify_with_llm,
    heuristic_classification,
)
from prism.parsing.extract import extract
from prism.parsing.model import BBox, ShapeKind, TemplateRaw
from prism.parsing.patterns import Pattern, build_pattern
from prism.parsing.tokens import DesignTokens, build_tokens
from prism.planning.schemas import SlideKind

SPECS_DIR = BACKEND_DIR / "data" / "specs"
# повышать при изменении парсера: кеш со старой версией разбора пересобирается автоматически
SPEC_VERSION = 11


class Asset(BaseModel):
    sha1: str
    path: str                        # относительно каталога спецификации
    kind: str                        # icon | logo | illustration
    category: str | None = None      # подпись раздела библиотеки («Бизнес», «Графики»…)
    px_w: int | None = None
    px_h: int | None = None
    avg_color: str | None = None
    slide_index: int


class PatternSpec(BaseModel):
    pattern: Pattern
    kind: SlideKind
    purpose: str = ""
    classified_by: str


class TemplateSpec(BaseModel):
    version: int = 0
    source: str
    sha256: str
    tokens: DesignTokens
    patterns: list[PatternSpec]
    service_slides: list[int]
    asset_library_slides: list[int]
    assets: list[Asset] = Field(default_factory=list)
    classifier: dict | None = None   # версия скилла template_classifier (name@vN + sha256) для manifest

    def by_kind(self, kind: SlideKind) -> list[PatternSpec]:
        return [p for p in self.patterns if p.kind == kind]


def _category(raw: TemplateRaw, slide_index: int, bbox: BBox) -> str | None:
    """Ближайшая короткая подпись выше/левее иконки — название раздела библиотеки."""
    slide = raw.slides[slide_index - 1]
    labels = [s for s in slide.shapes if s.kind == ShapeKind.text and 0 < len(s.text.strip()) <= 30
              and not s.text.strip().isdigit() and s.bbox.y <= bbox.y + bbox.h]
    if not labels:
        return None
    cx, cy = bbox.x + bbox.w / 2, bbox.y + bbox.h / 2

    def dist(s):
        # подпись раздела обычно над колонкой иконок: вертикальное расстояние важнее
        return abs(s.bbox.x + s.bbox.w / 2 - cx) * 0.5 + abs(cy - (s.bbox.y + s.bbox.h))

    return min(labels, key=dist).text.strip()


def _export_assets(pptx_path: Path, raw: TemplateRaw, slides: list[int], out_dir: Path) -> list[Asset]:
    prs = Presentation(str(pptx_path))
    assets: dict[str, Asset] = {}
    (out_dir / "assets").mkdir(parents=True, exist_ok=True)
    for idx in slides:
        raw_by_id = {s.id: s for s in raw.slides[idx - 1].shapes}
        stack = list(prs.slides[idx - 1].shapes)
        while stack:
            sh = stack.pop()
            if sh.shape_type is not None and hasattr(sh, "shapes"):
                stack.extend(sh.shapes)
                continue
            raw_shape = raw_by_id.get(sh.shape_id)
            if raw_shape is None or raw_shape.image is None or raw_shape.image.sha1 in assets:
                continue
            ext = raw_shape.image.content_type.split("/")[-1].replace("jpeg", "jpg").replace("x-", "")
            rel = f"assets/{raw_shape.image.sha1}.{ext}"
            try:
                (out_dir / rel).write_bytes(sh.image.blob)
            except (AttributeError, ValueError):
                continue
            side = max(raw_shape.bbox.w, raw_shape.bbox.h)
            kind = "icon" if side <= raw.slide_w * 0.06 else "logo" if raw_shape.bbox.h < raw_shape.bbox.w / 2 \
                else "illustration"
            assets[raw_shape.image.sha1] = Asset(
                sha1=raw_shape.image.sha1, path=rel, kind=kind, category=_category(raw, idx, raw_shape.bbox),
                px_w=raw_shape.image.px_w, px_h=raw_shape.image.px_h, avg_color=raw_shape.image.avg_color,
                slide_index=idx,
            )
    return list(assets.values())


def build_spec(raw: TemplateRaw, classification: TemplateClassification,
               patterns: dict[int, Pattern]) -> TemplateSpec:
    by_index = {c.index: c for c in classification.slides}
    pattern_idx = [i for i, c in by_index.items() if c.usage == Usage.pattern]
    tokens = build_tokens(raw, pattern_idx or None)
    return TemplateSpec(
        version=SPEC_VERSION, source=raw.source, sha256=raw.sha256, tokens=tokens,
        patterns=[
            PatternSpec(pattern=patterns[i], kind=by_index[i].kind or SlideKind.bullets,
                        purpose=by_index[i].purpose, classified_by=by_index[i].source)
            for i in pattern_idx
        ],
        service_slides=[i for i, c in by_index.items() if c.usage == Usage.service],
        asset_library_slides=[i for i, c in by_index.items() if c.usage == Usage.asset_library],
        classifier=classification.skill,
    )


def _save(path: Path, raw: TemplateRaw, spec: TemplateSpec) -> TemplateSpec:
    out_dir = SPECS_DIR / raw.sha256[:16]
    out_dir.mkdir(parents=True, exist_ok=True)
    spec.assets = _export_assets(path, raw, spec.asset_library_slides, out_dir)
    (out_dir / "spec.json").write_text(spec.model_dump_json(indent=1, exclude_defaults=True), encoding="utf-8")
    (out_dir / "raw.json").write_text(raw.model_dump_json(exclude_defaults=True), encoding="utf-8")
    return spec


def _cached(raw: TemplateRaw) -> TemplateSpec | None:
    cached = SPECS_DIR / raw.sha256[:16] / "spec.json"
    if not cached.exists():
        return None
    spec = TemplateSpec.model_validate_json(cached.read_text(encoding="utf-8"))
    return spec if spec.version == SPEC_VERSION else None


def _patterns(raw: TemplateRaw) -> dict[int, Pattern]:
    tokens_all = build_tokens(raw)   # предварительные токены: нужны паттернам для ролей кеглей
    return {s.index: build_pattern(raw, s, tokens_all) for s in raw.slides}


def parse_template_sync(path: str | Path, *, use_cache: bool = True) -> TemplateSpec:
    """Разбор на эвристиках, без модели: ~1-3 с, годится прямо в обработчике загрузки."""
    path = Path(path)
    raw = extract(path)
    if use_cache and (spec := _cached(raw)) is not None:
        return spec
    patterns = _patterns(raw)
    return _save(path, raw, build_spec(raw, heuristic_classification(raw.slides, patterns), patterns))


async def parse_template(path: str | Path, *, llm=None, skills=None, use_cache: bool = True) -> TemplateSpec:
    """Полный разбор; с llm/skills классификацию уточняет скилл template_classifier."""
    if llm is None or skills is None:
        return parse_template_sync(path, use_cache=use_cache)
    path = Path(path)
    raw = extract(path)
    if use_cache and (spec := _cached(raw)) is not None and all(p.classified_by == "llm" for p in spec.patterns):
        return spec
    patterns = _patterns(raw)
    classification = await classify_with_llm(raw.slides, patterns, llm, skills, slide_area=raw.slide_w * raw.slide_h)
    return _save(path, raw, build_spec(raw, classification, patterns))


def _summary(spec: TemplateSpec, classes: list[SlideClass] | None = None) -> str:
    t = spec.tokens
    lines = [
        (f"{spec.source}: {len(spec.patterns)} образцов, служебные {spec.service_slides}, "
         f"библиотеки {spec.asset_library_slides}, ассетов {len(spec.assets)}"),
        f"  шрифты: {t.fonts.heading} / {t.fonts.body}; шкала: "
        + ", ".join(f"{s.size_pt:g}({s.role})" for s in t.type_scale if s.usage >= 0.01),
        f"  акцент {t.color('accent')}, фон {t.color('background')}, текст {t.color('text')}",
    ]
    for p in spec.patterns:
        groups = "; ".join(f"{g.count}x{g.arrangement}{'!' if g.fixed_count else ''}" for g in p.pattern.groups)
        lines.append(f"  #{p.pattern.slide_index:<3} {p.kind.value:<11} [{p.classified_by[0]}] {groups:<18} {p.purpose}")
    return "\n".join(lines)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--llm", action="store_true", help="уточнить классификацию скиллом template_classifier")
    ap.add_argument("--no-cache", action="store_true")
    args = ap.parse_args()
    llm = skills = None
    if args.llm:
        from prism.config import load_config
        from prism.llm.client import LLMFactory
        from prism.skills.registry import SkillRegistry

        cfg = load_config()
        llm, skills = LLMFactory(cfg), SkillRegistry(pins=cfg.skills)
    spec = asyncio.run(parse_template(args.pptx, llm=llm, skills=skills, use_cache=not args.no_cache))
    print(_summary(spec))


if __name__ == "__main__":
    main()
