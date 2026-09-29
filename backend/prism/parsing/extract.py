"""Извлечение шаблона .pptx в TemplateRaw: геометрия, стили текста, цвета, картинки — с наследованием.

    python -m prism.parsing.extract data/templates/x.pptx [-o out.json]
"""

from __future__ import annotations

import argparse
import hashlib
import io
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pptx.slide import SlideLayout, SlideMaster

from prism.parsing.colors import A, ColorContext, color_element, resolve, resolve_fill
from prism.parsing.model import (
    Background,
    BBox,
    Guide,
    ImageRef,
    LayoutRaw,
    MasterRaw,
    Paragraph,
    PlaceholderRef,
    RunStyle,
    Shape,
    ShapeKind,
    SlideRaw,
    TemplateRaw,
    ThemeRaw,
)

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NS = {"a": A, "p": P, "c": C}

_TITLE_PH = {"title", "ctrTitle"}
_OTHER_PH = {"dt", "ftr", "sldNum"}


# ── геометрия ──────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Xform:
    """Аффинное преобразование дочерних координат группы в координаты слайда (без поворота)."""

    sx: float = 1.0
    sy: float = 1.0
    tx: float = 0.0
    ty: float = 0.0

    def apply(self, x: float, y: float, w: float, h: float) -> BBox:
        return BBox(x=round(x * self.sx + self.tx), y=round(y * self.sy + self.ty),
                    w=round(w * self.sx), h=round(h * self.sy))

    def child(self, grp_xfrm: etree._Element | None) -> Xform:
        if grp_xfrm is None:
            return self
        off, ext = grp_xfrm.find("a:off", NS), grp_xfrm.find("a:ext", NS)
        ch_off, ch_ext = grp_xfrm.find("a:chOff", NS), grp_xfrm.find("a:chExt", NS)
        if off is None or ext is None or ch_off is None or ch_ext is None:
            return self
        cw, ch = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
        gx = int(ext.get("cx")) / cw if cw else 1.0
        gy = int(ext.get("cy")) / ch if ch else 1.0
        # child -> group space: off + (p - chOff) * g ; затем -> родитель через self
        tx = int(off.get("x")) - int(ch_off.get("x")) * gx
        ty = int(off.get("y")) - int(ch_off.get("y")) * gy
        return Xform(self.sx * gx, self.sy * gy, self.tx + tx * self.sx, self.ty + ty * self.sy)


IDENTITY = Xform()


def _xfrm(el: etree._Element) -> etree._Element | None:
    for path in ("p:spPr/a:xfrm", "p:grpSpPr/a:xfrm", "p:xfrm"):
        found = el.find(path, NS)
        if found is not None:
            return found
    return None


# ── стили текста ───────────────────────────────────────────────────────────


class TextStyles:
    """Цепочка источников стиля: lstStyle фигуры -> плейсхолдеры макета/мастера -> txStyles -> defaults."""

    def __init__(self, master: SlideMaster, default_text_style: etree._Element | None, theme: ThemeRaw):
        tx = master._element.find("p:txStyles", NS)
        self.title = tx.find("p:titleStyle", NS) if tx is not None else None
        self.body = tx.find("p:bodyStyle", NS) if tx is not None else None
        self.other = tx.find("p:otherStyle", NS) if tx is not None else None
        self.default = default_text_style
        self.theme = theme

    def chain(self, shape: BaseShape) -> list[etree._Element]:
        chain: list[etree._Element] = []
        node = shape
        while node is not None:
            lst = node._element.find("p:txBody/a:lstStyle", NS)
            if lst is not None:
                chain.append(lst)
            node = getattr(node, "_base_placeholder", None) if node.is_placeholder else None
        if shape.is_placeholder:
            ph_type = shape._element.ph_type
            ph = str(ph_type).split(" ")[0].lower() if ph_type is not None else "body"
            ph = {"center_title": "ctrTitle", "slide_number": "sldNum", "date": "dt", "footer": "ftr"}.get(ph, ph)
            master_style = self.title if ph in _TITLE_PH else self.other if ph in _OTHER_PH else self.body
            chain.append(master_style)
        chain.append(self.default)
        if not shape.is_placeholder:
            chain.append(self.other)
        return [c for c in chain if c is not None]

    def is_title(self, shape: BaseShape) -> bool:
        return shape.is_placeholder and str(shape._element.ph_type).split(" ")[0].lower() in ("title", "center_title")


def _font_name(typeface: str | None, theme: ThemeRaw, title: bool) -> str | None:
    if not typeface:
        return None
    if typeface.startswith("+mj"):
        return theme.major_font
    if typeface.startswith("+mn"):
        return theme.minor_font
    return typeface


def _layers(p: etree._Element, level: int, chain: list[etree._Element]) -> list[etree._Element]:
    """pPr абзаца и lvlNpPr источников, от высшего приоритета к низшему."""
    layers = []
    ppr = p.find("a:pPr", NS)
    if ppr is not None:
        layers.append(ppr)
    for lst in chain:
        lvl = lst.find(f"a:lvl{level + 1}pPr", NS)
        if lvl is not None:
            layers.append(lvl)
    return layers


def _first(values):
    return next((v for v in values if v is not None), None)


def _run_style(rprs: list[etree._Element], ctx: ColorContext, theme: ThemeRaw, title: bool,
               font_ref_color: tuple[str | None, str | None]) -> RunStyle:
    def attr(name):
        return _first(r.get(name) for r in rprs)

    color, ref = None, None
    for r in rprs:
        fill = r.find("a:solidFill", NS)
        if fill is not None:
            color, ref = resolve(color_element(fill), ctx)
            break
    if color is None:
        color, ref = font_ref_color
    latin = _first((r.find("a:latin", NS).get("typeface") if r.find("a:latin", NS) is not None else None)
                   for r in rprs)
    font = _font_name(latin, theme, title) or (theme.major_font if title else theme.minor_font)
    sz = attr("sz")
    return RunStyle(
        font=font,
        size_pt=int(sz) / 100 if sz else None,
        bold=attr("b") in ("1", "true"),
        italic=attr("i") in ("1", "true"),
        color=color,
        color_ref=ref,
    )


def _paragraphs(shape: BaseShape, styles: TextStyles, ctx: ColorContext) -> list[Paragraph]:
    tx = shape._element.find("p:txBody", NS)
    if tx is None:
        return []
    chain = styles.chain(shape)
    title = styles.is_title(shape)
    font_ref = shape._element.find("p:style/a:fontRef", NS)
    font_ref_color = resolve(color_element(font_ref), ctx) if font_ref is not None else (None, None)

    out = []
    for p in tx.findall("a:p", NS):
        text = "".join(t.text or "" for t in p.iter(f"{{{A}}}t"))
        ppr = p.find("a:pPr", NS)
        level = int(ppr.get("lvl", "0")) if ppr is not None else 0
        layers = _layers(p, level, chain)
        defs = [d for d in (layer.find("a:defRPr", NS) for layer in layers) if d is not None]

        bullet = False
        for layer in layers:
            if layer.find("a:buNone", NS) is not None:
                break
            if layer.find("a:buChar", NS) is not None or layer.find("a:buAutoNum", NS) is not None:
                bullet = True
                break

        run_styles = [
            _run_style([r.find("a:rPr", NS), *defs] if r.find("a:rPr", NS) is not None else defs,
                       ctx, styles.theme, title, font_ref_color)
            for r in p.findall("a:r", NS) if (r.findtext("a:t", namespaces=NS) or "").strip()
        ]
        end = p.find("a:endParaRPr", NS)
        base = run_styles[0] if run_styles else _run_style(
            [end, *defs] if end is not None else defs, ctx, styles.theme, title, font_ref_color)
        out.append(Paragraph(
            text=text, level=level, bullet=bullet and bool(text.strip()),
            align=_first(layer.get("algn") for layer in layers),
            style=base, run_styles=run_styles,
        ))
    # хвостовые пустые абзацы не несут информации — но один оставляем: у пустого плейсхолдера
    # это единственный носитель унаследованного стиля (кегль, цвет), по нему считается подгонка
    while len(out) > 1 and not out[-1].text.strip():
        out.pop()
    return out


# ── фигуры ─────────────────────────────────────────────────────────────────


def _kind(el: etree._Element, has_text: bool) -> tuple[ShapeKind, str | None]:
    tag = etree.QName(el).localname
    if tag == "grpSp":
        return ShapeKind.group, None
    if tag == "pic":
        return ShapeKind.picture, None
    if tag == "cxnSp":
        return ShapeKind.line, None
    if tag == "graphicFrame":
        gd = el.find(".//a:graphicData", NS)
        uri = gd.get("uri", "") if gd is not None else ""
        if uri.endswith("/table"):
            return ShapeKind.table, uri
        if uri.endswith("/chart"):
            return ShapeKind.chart, uri
        if "diagram" in uri:
            return ShapeKind.smartart, uri
        return ShapeKind.other, uri
    if tag == "sp":
        return (ShapeKind.text if has_text else ShapeKind.shape), None
    return ShapeKind.other, None


_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


DETAIL = 16


def _detail_grid(im: Image.Image) -> list[list[float]]:
    """Разброс яркости по клеткам сетки DETAIL×DETAIL (строки сверху вниз)."""
    cell = 12
    gray = im.convert("L").resize((DETAIL * cell, DETAIL * cell))
    px = gray.load()
    grid = []
    for gy in range(DETAIL):
        row = []
        for gx in range(DETAIL):
            vals = [px[gx * cell + x, gy * cell + y] for y in range(cell) for x in range(cell)]
            mean = sum(vals) / len(vals)
            row.append(round((sum((v - mean) ** 2 for v in vals) / len(vals)) ** 0.5 / 255, 3))
        grid.append(row)
    return grid


def _image_ref(el: etree._Element, part, cache: dict[str, ImageRef]) -> ImageRef | None:
    """Картинка из pic или blipFill фигуры/фона: хеш, размер, средний цвет, доля непрозрачного."""
    blip = el.find(".//a:blip", NS)
    rid = blip.get(_R_EMBED) if blip is not None else None
    if not rid:
        return None
    try:
        img_part = part.related_part(rid)
    except KeyError:
        return None
    blob = img_part.blob
    sha1 = hashlib.sha1(blob).hexdigest()
    if sha1 in cache:
        return cache[sha1]
    ref = ImageRef(sha1=sha1, content_type=img_part.content_type)
    try:
        with Image.open(io.BytesIO(blob)) as im:
            ref.px_w, ref.px_h = im.size
            small = im.convert("RGBA").resize((32, 32))
            data = small.tobytes()
            opaque = [data[i:i + 3] for i in range(0, len(data), 4) if data[i + 3] > 127]
            ref.opaque_ratio = round(len(opaque) / 1024, 3)
            if opaque:
                ref.avg_color = "".join(f"{sum(c[i] for c in opaque) // len(opaque):02X}" for i in range(3))
            if im.size[0] >= 600:
                ref.detail = _detail_grid(im)
    except (OSError, ValueError):  # EMF/WMF и прочее, что Pillow не открывает
        pass
    cache[sha1] = ref
    return ref


class Extractor:
    def __init__(self, prs):
        self.prs = prs
        self.images: dict[str, ImageRef] = {}
        pres_el = prs.part._element
        self.default_text_style = pres_el.find("p:defaultTextStyle", NS)

    def walk(self, shapes, styles: TextStyles, ctx: ColorContext, xf: Xform = IDENTITY,
             parent: int | None = None, z: list[int] | None = None) -> Iterator[Shape]:
        z = z if z is not None else [0]
        for sh in shapes:
            el = sh._element
            z[0] += 1
            paragraphs = _paragraphs(sh, styles, ctx) if sh.has_text_frame else []
            prompt = False
            if sh.is_placeholder and not any(p.text.strip() for p in paragraphs):
                base = getattr(sh, "_base_placeholder", None)
                if base is not None and base.has_text_frame:
                    base_paragraphs = _paragraphs(base, styles, ctx)
                    if any(p.text.strip() for p in base_paragraphs):
                        paragraphs, prompt = base_paragraphs, True
            has_text = any(p.text.strip() for p in paragraphs) and not prompt
            # пустой текстовый плейсхолдер (заголовок, подзаголовок, тело) — тоже место под текст
            ph_kind = str(el.ph_type).split(" ")[0].lower() if sh.is_placeholder else ""
            empty_text_ph = sh.is_placeholder and sh.has_text_frame and ph_kind not in (
                "picture", "chart", "table", "media_clip", "slide_number", "date", "footer", "slide_image")
            kind, _ = _kind(el, has_text or prompt or empty_text_ph)

            xfrm = _xfrm(el)
            if xfrm is not None and xfrm.find("a:off", NS) is not None and parent is not None:
                off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
                bbox = xf.apply(int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
            else:  # верхний уровень или плейсхолдер без xfrm — python-pptx учитывает наследование
                bbox = xf.apply(sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)

            sppr = el.find("p:spPr", NS) if el.find("p:spPr", NS) is not None else el.find("p:grpSpPr", NS)
            fill_kind, fill = resolve_fill(sppr, ctx)
            base = getattr(sh, "_base_placeholder", None) if sh.is_placeholder else None
            while fill_kind is None and base is not None:  # заливка плейсхолдера наследуется от макета/мастера
                fill_kind, fill = resolve_fill(base._element.find("p:spPr", NS), ctx)
                base = getattr(base, "_base_placeholder", None)
            if fill_kind is None:
                fill_ref = el.find("p:style/a:fillRef", NS)
                if fill_ref is not None and fill_ref.get("idx", "0") != "0":
                    fill_kind, fill = "solid", resolve(color_element(fill_ref), ctx)[0]
            if kind == ShapeKind.picture:
                fill_kind = "picture"
            ln = sppr.find("a:ln", NS) if sppr is not None else None
            line = resolve(color_element(ln.find("a:solidFill", NS)), ctx)[0] \
                if ln is not None and ln.find("a:solidFill", NS) is not None else None
            geom = sppr.find("a:prstGeom", NS) if sppr is not None else None
            body_pr = el.find("p:txBody/a:bodyPr", NS)
            norm = body_pr.find("a:normAutofit", NS) if body_pr is not None else None
            autofit = None
            if body_pr is not None:
                autofit = ("norm" if norm is not None else "shape" if body_pr.find("a:spAutoFit", NS) is not None
                           else "none" if body_pr.find("a:noAutofit", NS) is not None else None)
            insets = None
            if body_pr is not None:  # умолчания DrawingML: 0.1" по бокам, 0.05" сверху/снизу
                insets = tuple(int(body_pr.get(k, d)) for k, d in
                               (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720)))

            item = Shape(
                id=sh.shape_id, name=sh.name, kind=kind, bbox=bbox,
                rotation=round(float(xfrm.get("rot", "0")) / 60000, 2) if xfrm is not None else 0.0,
                parent_group=parent, z=z[0],
                geometry=geom.get("prst") if geom is not None else None,
                fill=fill, fill_kind=fill_kind, line=line,
                placeholder=PlaceholderRef(type=el.ph_type and str(el.ph_type).split(" ")[0].lower(),
                                           idx=el.ph_idx) if sh.is_placeholder else None,
                paragraphs=paragraphs, prompt_text=prompt, autofit=autofit, insets=insets,
                autofit_scale=int(norm.get("fontScale", "100000")) / 100000 if norm is not None else 1.0,
                image=_image_ref(el, sh.part, self.images) if kind == ShapeKind.picture or fill_kind == "picture" else None,
            )
            if kind == ShapeKind.table:
                tbl = sh.table
                item.table_size = (len(tbl.rows), len(tbl.columns))
                # рамка graphicFrame часто не совпадает с таблицей: реальный размер — сумма колонок и строк
                grid_w = sum(c.width for c in tbl.columns)
                grid_h = sum(r.height for r in tbl.rows)
                if grid_w and grid_h:
                    item.bbox = BBox(x=item.bbox.x, y=item.bbox.y, w=max(item.bbox.w, grid_w),
                                     h=max(item.bbox.h, grid_h))
            elif kind == ShapeKind.chart:
                chart = sh.chart
                item.chart_type = str(chart.chart_type).split(" ")[0]
                item.chart_series = sum(len(plot.series) for plot in chart.plots)
            yield item

            if isinstance(sh, GroupShape):
                yield from self.walk(sh.shapes, styles, ctx, xf.child(_xfrm(el)), sh.shape_id, z)

    def background(self, owner, ctx: ColorContext, source: str) -> Background | None:
        bg = owner._element.find("p:cSld/p:bg", NS)
        if bg is None:
            return None
        bg_pr = bg.find("p:bgPr", NS)
        if bg_pr is not None:
            kind, color = resolve_fill(bg_pr, ctx)
            image = _image_ref(bg_pr, owner.part, self.images) if kind == "picture" else None
            if image is not None:
                color = image.avg_color
            return Background(kind=kind or "none", color=color, image=image, source=source)
        bg_ref = bg.find("p:bgRef", NS)
        if bg_ref is not None:
            return Background(kind="solid", color=resolve(color_element(bg_ref), ctx)[0], source=source)
        return None


P15 = "http://schemas.microsoft.com/office/powerpoint/2012/main"
_GUIDE_UNIT_EMU = 12700 / 8  # pos направляющих — в 1/8 пункта


def _guides(el: etree._Element, source: str) -> list[Guide]:
    return [
        Guide(orient="h" if g.get("orient") == "horz" else "v", pos=round(int(g.get("pos", "0")) * _GUIDE_UNIT_EMU),
              source=source)
        for g in el.iter(f"{{{P15}}}guide")
    ]


def _theme(master: SlideMaster) -> ThemeRaw:
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            x = etree.fromstring(rel.target_part.blob)
            cs = x.find(".//a:clrScheme", NS)
            colors = {}
            for c in cs if cs is not None else []:
                hex_color, _ = resolve(c[0], ColorContext({}, {})) if len(c) else (None, None)
                if hex_color:
                    colors[etree.QName(c).localname] = hex_color
            major = x.find(".//a:fontScheme/a:majorFont/a:latin", NS)
            minor = x.find(".//a:fontScheme/a:minorFont/a:latin", NS)
            return ThemeRaw(
                name=cs.get("name") if cs is not None else None, colors=colors,
                major_font=major.get("typeface") if major is not None else None,
                minor_font=minor.get("typeface") if minor is not None else None,
            )
    return ThemeRaw(name=None, colors={}, major_font=None, minor_font=None)


def extract(path: str | Path) -> TemplateRaw:
    path = Path(path)
    data = path.read_bytes()
    prs = Presentation(io.BytesIO(data))
    ex = Extractor(prs)

    masters, layouts, slides = [], [], []
    guides = _guides(prs.part._element, "presentation")
    ctx_by_master: dict[int, tuple[ColorContext, TextStyles, Background]] = {}
    layout_info: dict[int, tuple[int, Background]] = {}

    for mi, m in enumerate(prs.slide_masters):
        theme = _theme(m)
        cmap_el = m._element.find("p:clrMap", NS)
        cmap = dict(cmap_el.attrib) if cmap_el is not None else {}
        ctx = ColorContext(theme.colors, cmap)
        styles = TextStyles(m, ex.default_text_style, theme)
        bg = ex.background(m, ctx, "master") or Background(kind="none", source="master")
        ctx_by_master[mi] = (ctx, styles, bg)
        guides += _guides(m._element, f"master:{mi}")
        masters.append(MasterRaw(index=mi, theme=theme, color_map=cmap, background=bg,
                                 shapes=list(ex.walk(m.shapes, styles, ctx))))
        for layout in m.slide_layouts:
            lbg = ex.background(layout, ctx, "layout") or bg
            layout_info[id(layout._element)] = (mi, lbg)
            guides += _guides(layout._element, f"layout:{layout.name}")
            layouts.append(LayoutRaw(name=layout.name, master=mi, background=lbg,
                                     shapes=list(ex.walk(layout.shapes, styles, ctx))))

    for i, s in enumerate(prs.slides, 1):
        layout: SlideLayout = s.slide_layout
        mi, lbg = layout_info[id(layout._element)]
        ctx, styles, _ = ctx_by_master[mi]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        slides.append(SlideRaw(
            index=i, layout=layout.name, master=mi,
            background=ex.background(s, ctx, "slide") or lbg,
            shapes=list(ex.walk(s.shapes, styles, ctx)), notes=notes.strip(),
        ))

    return TemplateRaw(
        source=path.name, sha256=hashlib.sha256(data).hexdigest(),
        slide_w=prs.slide_width, slide_h=prs.slide_height,
        masters=masters, layouts=layouts, slides=slides, guides=guides,
    )


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out")
    args = ap.parse_args()
    raw = extract(args.pptx)
    out = Path(args.out or Path(args.pptx).with_suffix(".raw.json"))
    out.write_text(raw.model_dump_json(indent=1, exclude_defaults=True), encoding="utf-8")
    n_shapes = sum(len(s.shapes) for s in raw.slides)
    print(f"{raw.source}: {len(raw.slides)} слайдов, {len(raw.layouts)} макетов, {n_shapes} фигур, "
          f"{len({s.image.sha1 for sl in raw.slides for s in sl.shapes if s.image})} уникальных картинок -> {out}")


if __name__ == "__main__":
    main()
