"""Сборка колоды по шаблону: образцы выбираются по типу, текст встаёт в слоты, образцы-заглушки не остаются."""

import io
from collections import Counter
from pathlib import Path

import pytest
from pptx import Presentation

from prism.layout.compose import choose_pattern, compose_deck
from prism.layout.fit import ParaSpec, fit, wrap
from prism.parsing.spec import parse_template_sync
from prism.planning.schemas import ChartData, ChartSeries, Item, SlideContent, SlideKind

TEMPLATES = Path(__file__).resolve().parent.parent / "data" / "templates"
SAMPLES = {"Название презентации", "Заголовок слайда", "Заголовок", "Текст карточки", "Как мы запускаем",
           "Описание этапа"}

CONTENT = [
    SlideContent(kind=SlideKind.title, title="Доставка еды в Казани"),
    SlideContent(kind=SlideKind.cards, title="Три причины начать сейчас",
                 items=[Item(heading="Рынок", text="Растёт на 21% в год"),
                        Item(heading="Комиссии", text="Конкуренты берут до 35%")]),
    SlideContent(kind=SlideKind.process, title="Запуск за три шага",
                 items=[Item(text="Рестораны"), Item(text="Курьеры"), Item(text="Маркетинг")]),
]


def _texts(prs) -> list[list[str]]:
    return [[sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            for slide in prs.slides]


def test_wrap_keeps_numbers_together():
    lines = wrap("выручка 1 150 ₽ за месяц", width_pt=60, family=None, size_pt=12, bold=False)
    assert any("1 150 ₽" in line for line in lines)


def test_fit_shrinks_to_scale_steps():
    paras = [ParaSpec(text="Очень длинный заголовок слайда, который не влезет в одну строку", font=None,
                      size_pt=36, bold=False)]
    r = fit(paras, box_w=12 * 360_000, box_h=3 * 360_000, steps=[1.0, 24 / 36, 20 / 36])
    assert r.fits and r.scale < 1.0


def test_compose_synthetic(deck):
    spec = parse_template_sync(deck, use_cache=False)
    data, reports = compose_deck(deck, spec, CONTENT)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(CONTENT)                    # исходные образцы удалены
    texts = _texts(prs)
    assert texts[0] == ["Доставка еды в Казани"]
    # карточки: две из трёх заполнены, третья (с образцом «Заголовок») удалена целиком
    assert "Три причины начать сейчас" in texts[1] and "Растёт на 21% в год" in texts[1]
    assert not SAMPLES & {t for slide in texts for t in slide}, texts
    # процесс: номера в стиле образца «01», «02», «03»
    assert {"01", "02", "03"} <= set(texts[2]) and "Рестораны" in texts[2]
    assert [r.kind for r in reports] == ["title", "cards", "process"]


def test_choose_pattern_prefers_own_kind(deck):
    spec = parse_template_sync(deck, use_cache=False)
    for content, expected in zip(CONTENT, ["title", "cards", "process"], strict=True):
        assert choose_pattern(spec, content, Counter()).kind.value == expected


def test_chart_is_native(deck):
    spec = parse_template_sync(deck, use_cache=False)
    chart = SlideContent(kind=SlideKind.chart, title="Рынок растёт",
                         chart=ChartData(categories=["2024", "2025"], series=[ChartSeries(name="млрд", values=[1, 2])]))
    data, _ = compose_deck(deck, spec, [chart])
    slide = Presentation(io.BytesIO(data)).slides[0]
    assert any(sh.has_chart for sh in slide.shapes)            # нативный график, не картинка
    assert not any(sh.shape_type == 13 for sh in slide.shapes)  # 13 = PICTURE


@pytest.mark.skipif(not TEMPLATES.exists(), reason="нет data/templates")
@pytest.mark.parametrize("name", ["vk_workspace", "vk_education", "vk_tech", "lct2026"])
def test_real_templates_compose_without_leftovers(name):
    path = TEMPLATES / f"{name}.pptx"
    if not path.exists():
        pytest.skip(f"нет {name}.pptx")
    spec = parse_template_sync(path)
    data, _ = compose_deck(path, spec, CONTENT)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(CONTENT)
    joined = " ".join(t for slide in _texts(prs) for t in slide)
    for placeholder in ("Lorem ipsum", "Образец текста", "Текст карточки"):
        assert placeholder not in joined, (name, placeholder)
