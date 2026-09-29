"""Аудит: детерминированные проверки находят дефекты по файлу, выбранные исправления их убирают."""

import asyncio

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from prism.audit.deterministic import audit_deck
from prism.audit.fixes import apply_fixes


def _box(slide, x, y, w, h, text, size=14, color=None, name=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.text_frame.word_wrap = True
    tb.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if name:
        tb.name = name
    return tb


def _deck(tmp_path):
    prs = Presentation()
    w = prs.slide_width
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.background.fill.solid()
    s1.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _box(s1, 400000, 300000, 6000000, 600000, "Выручка выросла на 20% за год и обогнала рынок")
    _box(s1, 400000, 1200000, 1400000, 300000, "Очень длинный текст, который не поместится в узкую и низкую рамку", 24)
    _box(s1, w - 800000, 2500000, 2000000, 400000, "Блок за краем слайда")
    _box(s1, 400000, 3500000, 3000000, 400000, "Бледный текст на белом фоне", 12, "EEEEEE")
    _box(s1, 400000, 4300000, 3000000, 400000, "[Что изменится]")
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s2, 400000, 300000, 6000000, 600000, "Один заголовок и больше ничего")
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def test_deterministic_checks_find_layout_and_integrity_defects(tmp_path):
    issues = audit_deck(_deck(tmp_path))
    found = {(i.slide, i.check) for i in issues}
    assert {(1, "overflow"), (1, "off_slide"), (1, "contrast"), (1, "sample_text"), (2, "empty_slide")} <= found
    assert all(i.mode == "deterministic" and i.group for i in issues)
    off = next(i for i in issues if i.check == "off_slide")
    assert off.box and off.fixes[0].id == "clamp"
    # детерминированность: тот же файл — тот же ответ
    assert [i.id for i in audit_deck(tmp_path / "deck.pptx")] == [i.id for i in issues]


def test_selected_fixes_remove_the_defects(tmp_path):
    path = _deck(tmp_path)
    issues = audit_deck(path)
    pick = {"off_slide": "clamp", "contrast": "recolor", "sample_text": "delete", "empty_slide": "delete_slide"}
    choices = [(i, pick[i.check]) for i in issues if i.check in pick]
    report = asyncio.run(apply_fixes(path, choices))
    assert not report.failed and report.deleted_slides == [2]
    left = {i.check for i in audit_deck(path)}
    assert not left & set(pick)
    assert len(Presentation(str(path)).slides) == 1


def test_model_fixes_fail_cleanly_without_model(tmp_path):
    path = _deck(tmp_path)
    overflow = next(i for i in audit_deck(path) if i.check == "overflow")
    report = asyncio.run(apply_fixes(path, [(overflow, "shorten")]))
    assert overflow.id in report.failed and not report.applied
