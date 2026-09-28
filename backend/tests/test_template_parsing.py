"""Токены, паттерны и классификация: синтетические слайды с известной структурой + реальные шаблоны."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from prism.color import contrast_ratio, delta_e
from prism.parsing.classify import Usage, heuristic_classification, heuristic_usage
from prism.parsing.extract import extract
from prism.parsing.patterns import SlotKind, build_pattern
from prism.parsing.tokens import build_tokens
from prism.planning.schemas import SlideKind

TEMPLATES = Path(__file__).resolve().parent.parent / "data" / "templates"
CM = 360_000


def _text(slide, x, y, w, h, text, size, color="000000", bold=False, fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if fill else MSO_SHAPE.RECTANGLE,
                                   Emu(x), Emu(y), Emu(w), Emu(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    shape.line.fill.background()
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Play"
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


@pytest.fixture
def deck(tmp_path) -> Path:
    """Слайд 1: титул. Слайд 2: 3 карточки (подложка + заголовок + текст), средняя — выделенная.
    Слайд 3: процесс из 4 шагов (номер + подпись). Слайд 4: служебный гайдлайн."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(33 * CM), Emu(19 * CM)
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    _text(s1, 2 * CM, 7 * CM, 20 * CM, 3 * CM, "Название презентации", 40, bold=True)

    s2 = prs.slides.add_slide(blank)
    _text(s2, 2 * CM, 1 * CM, 25 * CM, 2 * CM, "Заголовок слайда", 32)
    for i in range(3):
        x = (2 + i * 10) * CM
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(5 * CM), Emu(9 * CM), Emu(8 * CM))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string("0077FF" if i == 1 else "EEF2F7")  # акцентная карточка
        card.line.fill.background()
        _text(s2, x + CM // 2, 6 * CM, 8 * CM, CM, "Заголовок", 18, "0077FF", bold=True)
        _text(s2, x + CM // 2, 7 * CM + CM // 2, 8 * CM, 3 * CM, "Текст карточки", 14, "333333")

    s3 = prs.slides.add_slide(blank)
    _text(s3, 2 * CM, 1 * CM, 25 * CM, 2 * CM, "Как мы запускаем", 32)
    for i in range(4):
        x = (2 + i * 7) * CM
        _text(s3, x, 6 * CM, 2 * CM, 2 * CM, f"0{i + 1}", 28, "0077FF", bold=True)
        _text(s3, x, 9 * CM, 6 * CM, 2 * CM, "Описание этапа", 14)

    s4 = prs.slides.add_slide(blank)
    _text(s4, 2 * CM, 1 * CM, 25 * CM, 2 * CM, "Оформление", 32)
    _text(s4, 2 * CM, 4 * CM, 25 * CM, 2 * CM, "Шрифт: Play. Цвета: #0077FF, #EEF2F7", 14)
    _text(s4, 2 * CM, 7 * CM, 25 * CM, 2 * CM, "Иконки можно брать тут: https://example.com/icons", 14)

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_color_math():
    assert contrast_ratio("FFFFFF", "000000") == pytest.approx(21.0)
    assert contrast_ratio("0077FF", "FFFFFF") < 4.5          # фирменная пара VK ниже нормы для мелкого текста
    assert delta_e("0077FF", "0078FF") < 1 and delta_e("0077FF", "FF0053") > 50


def test_tokens_from_synthetic(deck):
    t = build_tokens(extract(deck))
    assert t.fonts.heading == t.fonts.body == "Play"
    roles = {s.size_pt: s.role for s in t.type_scale}
    assert roles[14.0] == "body" and roles[40.0] in ("title", "display")
    accent = t.color("accent")
    assert accent is not None and delta_e(accent, "0077FF") < 5
    assert abs(t.margins.left - 2 * CM) < CM // 2


def test_cards_group_includes_highlighted_card(deck):
    raw = extract(deck)
    p = build_pattern(raw, raw.slides[1], build_tokens(raw))
    assert p.title and p.title.sample == "Заголовок слайда"
    g = p.groups[0]
    assert (g.count, g.arrangement) == (3, "row")
    for item in g.items:
        kinds = [s.kind for s in item]
        assert SlotKind.surface in kinds and SlotKind.heading in kinds and SlotKind.text in kinds, kinds


def test_process_group(deck):
    raw = extract(deck)
    p = build_pattern(raw, raw.slides[2], build_tokens(raw))
    g = p.groups[0]
    assert g.count == 4 and g.arrangement == "row"
    assert [s.sample for it in g.items for s in it if s.kind == SlotKind.number] == ["01", "02", "03", "04"]


def test_heuristic_classification(deck):
    raw = extract(deck)
    tokens = build_tokens(raw)
    cls = heuristic_classification(raw.slides, {s.index: build_pattern(raw, s, tokens) for s in raw.slides})
    by = {c.index: c for c in cls.slides}
    assert by[1].kind == SlideKind.title
    assert by[2].kind == SlideKind.cards
    assert by[3].kind == SlideKind.process
    assert by[4].usage == Usage.service


# ── реальные шаблоны ────────────────────────────────────────────────────────

needs = pytest.mark.skipif(not TEMPLATES.exists(), reason="нет data/templates")


@needs
@pytest.mark.skipif(not (TEMPLATES / "vk_workspace.pptx").exists(), reason="нет vk_workspace.pptx")
def test_vk_workspace_known_patterns():
    raw = extract(TEMPLATES / "vk_workspace.pptx")
    tokens = build_tokens(raw)
    expected = {7: (6, "grid"), 8: (6, "grid"), 9: (8, "grid"), 16: (3, "row"), 26: (4, "row"), 27: (4, "row")}
    for idx, (count, arrangement) in expected.items():
        g = build_pattern(raw, raw.slides[idx - 1], tokens).groups[0]
        assert (g.count, g.arrangement) == (count, arrangement), idx


@needs
@pytest.mark.skipif(not (TEMPLATES / "lct2026.pptx").exists(), reason="нет lct2026.pptx")
def test_lct_holdout_service_and_assets():
    raw = extract(TEMPLATES / "lct2026.pptx")
    usage = {s.index: heuristic_usage(s) for s in raw.slides}
    assert all(usage[i] == Usage.asset_library for i in range(31, 38))
    assert usage[2] == Usage.service                       # «Привет, участник хакатона»
    assert usage[21] == Usage.pattern                      # слайд с нативным графиком
