"""Поиск для короткого брифа и MCP-серверы: кому искать, какие страницы читать, куда инструментам нельзя."""

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from prism.audit.deterministic import audit_deck
from prism.config import ResearchConfig
from prism.generation.checks import normalize
from prism.generation.research import _pick_pages, needs_research
from prism.layout.compose import _fill_in_slots, _short_title, _widen_into_empty
from prism.mcp import pptx_server, web_server
from prism.parsing.model import BBox
from prism.parsing.patterns import Slot, SlotKind
from prism.planning.schemas import Item, SlideContent, SlideKind


def test_short_brief_needs_research_detailed_does_not():
    cfg = ResearchConfig()
    assert needs_research("компания Nestle", cfg)
    detailed = "Выручка 58, 61, 66, 71 млн ₽ по месяцам; 185 000 установок; 41% заказов; окупаемость 13 месяцев."
    assert not needs_research(detailed, cfg)


def test_pick_pages_one_per_site_and_skips_video():
    a = [{"url": "https://ru.wikipedia.org/a"}, {"url": "https://ru.wikipedia.org/b"}, {"url": "https://rbc.ru/x"}]
    b = [{"url": "https://www.youtube.com/watch?v=1"}, {"url": "https://nestle.com/ir"}]
    urls = [p["url"] for p in _pick_pages([a, b], limit=5)]
    assert urls == ["https://ru.wikipedia.org/a", "https://nestle.com/ir", "https://rbc.ru/x"]


@pytest.mark.parametrize("url", ["http://127.0.0.1:8000/api/health", "http://localhost/", "http://10.0.0.5/",
                                 "http://169.254.169.254/latest/meta-data", "file:///etc/passwd"])
def test_fetch_refuses_internal_addresses(url):
    with pytest.raises(web_server.BlockedURL):
        asyncio.run(web_server._check_public(url))


def test_pptx_tools_stay_inside_data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(pptx_server, "DATA_DIR", tmp_path)
    with pytest.raises(ValueError):
        pptx_server._path("/etc/passwd")
    with pytest.raises(ValueError):
        pptx_server._path("../../secret.pptx", must_exist=False)
    assert pptx_server._path("decks/a.pptx", must_exist=False) == tmp_path / "decks" / "a.pptx"


def _deck(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Emu(500000), Emu(500000), Emu(1500000), Emu(300000))
    tb.text_frame.word_wrap = True
    tb.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Очень длинный текст, который никак не поместится в узкую и низкую рамку на слайде"
    run.font.size = Pt(24)
    tb2 = slide.shapes.add_textbox(Emu(500000), Emu(2000000), Emu(3000000), Emu(500000))
    tb2.text_frame.text = "[Что изменится]"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def test_audit_finds_overflow_and_leftover_sample(tmp_path):
    kinds = {(i.slide, i.check) for i in audit_deck(_deck(tmp_path))}
    assert (1, "overflow") in kinds and (1, "sample_text") in kinds


def test_set_shape_text_keeps_file_inside_data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(pptx_server, "DATA_DIR", tmp_path)
    path = _deck(tmp_path)
    prs = Presentation(str(path))
    sid = prs.slides[0].shapes[1].shape_id
    out = pptx_server.set_shape_text("deck.pptx", 1, sid, ["Готово"], out_path="fixed.pptx")
    assert out["saved"] == "fixed.pptx"
    assert "Готово" in [s.text_frame.text for s in Presentation(str(tmp_path / "fixed.pptx")).slides[0].shapes]


def _slot(sid, kind, sample, x, y, w=100, h=20):
    return Slot(shape_id=sid, kind=kind, bbox=BBox(x=x, y=y, w=w, h=h), sample=sample, styles=[])


def test_fill_in_slots_skip_template_labels():
    slots = [_slot(1, SlotKind.heading, "[Приоритет 1]", 0, 0), _slot(2, SlotKind.text, "Результат декабря", 0, 30),
             _slot(3, SlotKind.heading, "[Что изменится]", 0, 60)]
    assert [s.shape_id for s in _fill_in_slots(slots)] == [1, 3]


def test_text_widens_into_empty_column_to_the_right():
    row = [_slot(1, SlotKind.heading, "[Риск]", 0, 0), _slot(2, SlotKind.heading, "[Сигнал]", 120, 0),
           _slot(3, SlotKind.heading, "[Действие]", 240, 0)]
    wider = _widen_into_empty(row, filled={1, 2})
    assert 1 not in wider and wider[2].right == 340


def test_footer_does_not_end_on_preposition():
    assert _short_title("Запуск приложения окупается за 13 месяцев и уже сократил зависимость") \
        == "Запуск приложения окупается"
    assert _short_title("Приложение «Зерно» окупается за 13 месяцев и снижает комиссии") == "Приложение «Зерно»"


def test_normalize_drops_duplicate_items():
    c = SlideContent(kind=SlideKind.kpi, title="t", items=[Item(value="70%", text="через агрегаторы")] * 2
                     + [Item(value="30%", text="комиссия")])
    assert len(normalize(c).items) == 2
