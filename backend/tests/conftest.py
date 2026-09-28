"""Общие фикстуры: синтетическая колода с известной структурой (без закрытых шаблонов из data/)."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

CM = 360_000


def _text(slide, x, y, w, h, text, size, color="000000", bold=False, fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if fill else MSO_SHAPE.RECTANGLE,
                                   Emu(x), Emu(y), Emu(w), Emu(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    shape.line.fill.background()
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Play"
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


@pytest.fixture
def deck(tmp_path) -> Path:
    """Слайд 1: титул. Слайд 2: 3 карточки (подложка + заголовок + текст), средняя — выделенная.
    Слайд 3: процесс из 4 шагов (номер + подпись). Слайд 4: служебный гайдлайн."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(33 * CM), Emu(19 * CM)
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    _text(s1, 2 * CM, 7 * CM, 20 * CM, 3 * CM, "Название презентации", 40, bold=True)

    s2 = prs.slides.add_slide(blank)
    _text(s2, 2 * CM, 1 * CM, 25 * CM, 2 * CM, "Заголовок слайда", 32)
    for i in range(3):
        x = (2 + i * 10) * CM
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(5 * CM), Emu(9 * CM), Emu(8 * CM))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string("0077FF" if i == 1 else "EEF2F7")  # акцентная карточка
        card.line.fill.background()
        _text(s2, x + CM // 2, 6 * CM, 8 * CM, CM, "Заголовок", 18, "0077FF", bold=True)
        _text(s2, x + CM // 2, 7 * CM + CM // 2, 8 * CM, 3 * CM, "Текст карточки", 14, "333333")

    s3 = prs.slides.add_slide(blank)
    _text(s3, 2 * CM, 1 * CM, 25 * CM, 2 * CM, "Как мы запускаем", 32)
    for i in range(4):
        x = (2 + i * 7) * CM
        _text(s3, x, 6 * CM, 2 * CM, 2 * CM, f"0{i + 1}", 28, "0077FF", bold=True)
        _text(s3, x, 9 * CM, 6 * CM, 2 * CM, "Описание этапа", 14)

    s4 = prs.slides.add_slide(blank)
    _text(s4, 2 * CM, 1 * CM, 25 * CM, 2 * CM, "Оформление", 32)
    _text(s4, 2 * CM, 4 * CM, 25 * CM, 2 * CM, "Шрифт: Play. Цвета: #0077FF, #EEF2F7", 14)
    _text(s4, 2 * CM, 7 * CM, 25 * CM, 2 * CM, "Иконки можно брать тут: https://example.com/icons", 14)

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


