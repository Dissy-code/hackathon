"""Извлечение шаблона: синтетические pptx с известным ответом + сверка на реальных шаблонах (если лежат локально)."""

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from prism.parsing.colors import A, ColorContext, resolve
from prism.parsing.extract import extract
from prism.parsing.model import ShapeKind

TEMPLATES = Path(__file__).resolve().parent.parent / "data" / "templates"


def _color(xml: str) -> str | None:
    el = etree.fromstring(f'<a:x xmlns:a="{A}">{xml}</a:x>')[0]
    ctx = ColorContext({"dk1": "000000", "lt1": "FFFFFF", "accent1": "0077FF"}, {"tx1": "dk1", "bg1": "lt1"})
    return resolve(el, ctx)[0]


@pytest.mark.parametrize(
    ("xml", "expected"),
    [
        ('<a:srgbClr val="0077ff"/>', "0077FF"),
        ('<a:schemeClr val="accent1"/>', "0077FF"),
        ('<a:schemeClr val="tx1"/>', "000000"),          # через clrMap
        ('<a:schemeClr val="bg1"><a:lumMod val="50000"/></a:schemeClr>', "808080"),
        ('<a:srgbClr val="000000"><a:lumMod val="100000"/><a:lumOff val="100000"/></a:srgbClr>', "FFFFFF"),
        ('<a:srgbClr val="FF0000"><a:shade val="50000"/></a:srgbClr>', "800000"),
        ('<a:sysClr val="windowText" lastClr="111111"/>', "111111"),
    ],
)
def test_color_resolution(xml, expected):
    assert _color(xml) == expected


@pytest.fixture
def synthetic(tmp_path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "Заголовок-вывод"

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(1_000_000), Emu(2_000_000),
                                  Emu(3_000_000), Emu(1_500_000))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0x21, 0x21, 0x21)
    run = card.text_frame.paragraphs[0].add_run()
    run.text = "Карточка"
    run.font.size, run.font.name, run.font.bold = Pt(18), "Play", True
    run.font.color.rgb = RGBColor(0x00, 0x77, 0xFF)

    # группа, масштабированная в 2 раза: дочерняя фигура 100x100 по chOff (0,0) -> 200x200 со сдвигом группы
    grp = slide.shapes.add_group_shape()
    inner = grp.shapes.add_shape(MSO_SHAPE.OVAL, Emu(0), Emu(0), Emu(100_000), Emu(100_000))
    xfrm = grp._element.find(f"{{http://schemas.openxmlformats.org/presentationml/2006/main}}grpSpPr/{{{A}}}xfrm")
    for tag, attrs in (("off", {"x": "500000", "y": "600000"}), ("ext", {"cx": "200000", "cy": "200000"}),
                       ("chOff", {"x": "0", "y": "0"}), ("chExt", {"cx": "100000", "cy": "100000"})):
        xfrm.find(f"{{{A}}}{tag}").attrib.update(attrs)
    inner.name = "inner"

    path = tmp_path / "synthetic.pptx"
    prs.save(path)
    return path


def test_synthetic_card_style(synthetic):
    raw = extract(synthetic)
    shapes = raw.slides[0].shapes
    card = next(s for s in shapes if s.text == "Карточка")
    assert card.kind == ShapeKind.text and card.geometry == "roundRect" and card.fill == "212121"
    st = card.paragraphs[0].style
    assert (st.font, st.size_pt, st.bold, st.color) == ("Play", 18.0, True, "0077FF")


def test_title_inherits_from_master(synthetic):
    raw = extract(synthetic)
    title = next(s for s in raw.slides[0].shapes if s.placeholder and s.placeholder.type == "title")
    st = title.paragraphs[0].style
    assert st.size_pt == 44.0                      # titleStyle мастера дефолтного шаблона
    assert st.font == raw.masters[0].theme.major_font
    assert title.bbox.w > 0                        # позиция унаследована от макета


def test_group_transform(synthetic):
    raw = extract(synthetic)
    inner = next(s for s in raw.slides[0].shapes if s.name == "inner")
    assert inner.parent_group is not None
    assert (inner.bbox.x, inner.bbox.y, inner.bbox.w, inner.bbox.h) == (500_000, 600_000, 200_000, 200_000)


# ── реальные шаблоны (data/ не в git: тесты пропускаются, если файлов нет) ──

REAL = sorted(TEMPLATES.glob("*.pptx")) if TEMPLATES.exists() else []


@pytest.mark.skipif(not REAL, reason="нет шаблонов в data/templates")
@pytest.mark.parametrize("path", REAL, ids=lambda p: p.stem)
def test_real_template_invariants(path):
    raw = extract(path)
    assert raw.slides and raw.masters
    for slide in raw.slides:
        by_id = {s.id: s for s in slide.shapes}
        for s in slide.shapes:
            if s.parent_group is not None:
                g, b = by_id[s.parent_group].bbox, s.bbox
                assert b.x >= g.x - 20000 and b.right <= g.right + 20000, (slide.index, s.name)
            for p in s.paragraphs:
                assert p.style.font, (slide.index, s.name, p.text)
                assert not (p.style.font or "").startswith("+"), "ссылка на шрифт темы не разрешена"


@pytest.mark.skipif(not (TEMPLATES / "vk_workspace.pptx").exists(), reason="нет vk_workspace.pptx")
def test_vk_workspace_known_slide():
    s7 = extract(TEMPLATES / "vk_workspace.pptx").slides[6]
    cards = [s for s in s7.shapes if s.kind == ShapeKind.text and s.geometry == "roundRect"]
    assert len(cards) == 6
    assert {c.fill for c in cards} == {"212121"}
    head, body = cards[0].paragraphs[:2]
    assert (head.style.font, head.style.size_pt, head.style.color) == ("Play", 18.0, "0077FF")
    assert (body.style.size_pt, body.style.color) == (14.0, "E4E7EA")
    assert s7.background.color == "000000"
